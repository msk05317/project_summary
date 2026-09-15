# 하바플레이트 W37 시트 읽기
#   python3 backend/tests/test_hrva_w37.py
#
# 이 시트의 함정
#   1) 모델 목록 네 덩어리가 먼저 와서 '계획/실적' 줄이 87행이다 (24행까지만 봤다)
#   2) 12월~8월은 '실적' 열만 있고 9월만 계획/실적 짝이다
#   3) 그 월 열이 해를 넘어간다 (12월 → 1월 → … → 9월)
#   4) 합계 줄 '아래'에 개발 그룹 총계가 한 줄 더 있다
#   5) 판가에 소수점이 있다 (3150.4 · 1461.66)
import datetime as _dt
import pathlib, sys
import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import plan_matrix_import as pm
import hrva_status_import as hs

TODAY = _dt.date(2026, 9, 15)


def build():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'W37'

    # 위에 오는 모델 목록 덩어리들 — 표를 아래로 밀어낸다
    ws['B10'] = '2. 양산모델'
    ws['B11'] = 'No'; ws['C11'] = '모델'; ws['E11'] = '판가'
    ws['G11'] = 'PO 수량'; ws['H11'] = '출하수량'; ws['I11'] = 'PO 잔량'
    for i, (mid, price, po, sh) in enumerate([
            ('713-312133-006', 3150.4, 434, 407),
            ('713-311273-004', 2765.6, 1204, 1196)]):
        r = 12 + i
        ws.cell(r, 2).value = i + 1; ws.cell(r, 3).value = mid
        ws.cell(r, 5).value = price
        ws.cell(r, 7).value = po; ws.cell(r, 8).value = sh
        ws.cell(r, 9).value = po - sh
    ws['B14'] = 'Total'

    ws['B40'] = '4. 개발품 (LAM)'
    ws['B41'] = 'No'; ws['C41'] = '모델'; ws['E41'] = '판가'
    ws['F41'] = '개발단계'; ws['H41'] = '개발종류'
    ws['J41'] = 'PO 수량'; ws['K41'] = '출하수량'; ws['L41'] = 'PO 잔량'
    ws['B42'] = 1; ws['C42'] = '15-453089-00'; ws['E42'] = 1461.66
    ws['F42'] = 'LAIR'; ws['H42'] = 'HVM'
    ws['J42'] = 5; ws['K42'] = 1; ws['L42'] = 4

    # 주차 표 — 85행부터
    ws['B85'] = '양산'; ws['C85'] = '파트번호'
    ws['D85'] = 'PO 수량'; ws['E85'] = '출하실적'; ws['F85'] = 'PO 잔량'
    for col, mon in (('G', '12월'), ('H', '1월'), ('I', '8월')):
        ws[f'{col}85'] = mon
        ws[f'{col}87'] = '실적'
    ws['J85'] = '9월'; ws.merge_cells('J85:Q85')
    for col, wk in (('J', 'W36'), ('L', 'W37'), ('N', 'W38'), ('P', 'W39')):
        ws[f'{col}86'] = wk
        ws.merge_cells(f'{col}86:{chr(ord(col)+1)}86')
        ws[f'{col}87'] = '계획'
        ws[f'{chr(ord(col)+1)}87'] = '실적'
    ws['R85'] = '9월 Total'; ws['R87'] = '계획'; ws['S87'] = '실적'

    rows = [
        ('713-312133-006', 434, 407, 27, (20, 10, 119),
         {'J': 10, 'K': 7, 'L': 10, 'M': 23, 'N': 20}),
        ('713-C21755-001', 64, 63, 1, (None, None, 0), {'N': 1}),
    ]
    r = 88
    for mid, po, sh, rem, mons, cells in rows:
        ws.cell(r, 3).value = mid
        ws.cell(r, 4).value = po; ws.cell(r, 5).value = sh; ws.cell(r, 6).value = rem
        for col, v in zip(('G', 'H', 'I'), mons):
            if v is not None:
                ws[f'{col}{r}'] = v
        for col, v in cells.items():
            ws[f'{col}{r}'] = v
        r += 1
    ws.cell(r, 3).value = 'total'
    ws.cell(r, 4).value = 498; ws.cell(r, 5).value = 470
    ws['J%d' % r] = 10; ws['N%d' % r] = 21
    r += 1
    # 합계 '아래'에 개발 그룹 총계
    ws.cell(r, 2).value = '개발'
    ws.cell(r, 3).value = '개발 (41종 진행중 / 146종 완료)'
    ws.cell(r, 4).value = 1073; ws.cell(r, 5).value = 787
    ws['J%d' % r] = 7; ws['K%d' % r] = 11; ws['N%d' % r] = 55
    return wb


wb = build()
out = pm.parse_plan_matrix(wb, sheet_name='W37', today=TODAY)
rows = {r['label']: r for r in out['rows']}
ok = 0

assert out['sheet'] == 'W37', '표를 못 찾았다 (계획/실적 줄이 87행)'
ok += 1
assert list(rows) == ['713-312133-006', '713-C21755-001',
                      '개발 (41종 진행중 / 146종 완료)'], list(rows)
ok += 1

a = rows['713-312133-006']
assert (a['po_qty'], a['shipped_qty']) == (434, 407), a
w = a['weeks']['2026-09']
assert w['W36'] == {'plan': 10, 'actual': 7}, w
assert w['W37'] == {'plan': 10, 'actual': 23}, w
assert w['W38'] == {'plan': 20, 'actual': 0}, w
ok += 1

# 월 실적만 있는 열도 읽는다
assert a['months']['2026-08'] == {'plan': 0, 'actual': 119}, a['months']
ok += 1
# 해를 넘어간다 — 12월은 작년이다
assert '2025-12' in a['months'], f"12월이 작년으로 안 잡혔다: {sorted(a['months'])}"
assert a['months']['2025-12']['actual'] == 20
assert a['months']['2026-01']['actual'] == 10
ok += 1
# 9월은 주차가 있으니 월 합계 열을 안 쓴다 (이중 계산)
assert '2026-09' not in a['months'], a['months']
ok += 1

# 합계 아래의 개발 그룹 총계
dev = rows['개발 (41종 진행중 / 146종 완료)']
assert dev['is_aggregate'] is True
assert (dev['po_qty'], dev['shipped_qty']) == (1073, 787), dev
assert dev['weeks']['2026-09']['W36'] == {'plan': 7, 'actual': 11}, dev['weeks']
ok += 1
assert out['total'] and out['total']['po_qty'] == 498, out['total']
ok += 1

# 판가 소수점
models = {m['id']: m for m in hs.parse_models(wb['W37'])}
assert models['713-312133-006']['price'] == 3150.4, models['713-312133-006']
assert models['713-311273-004']['price'] == 2765.6
assert models['15-453089-00']['price'] == 1461.66, models['15-453089-00']
ok += 1
# 정수 판가는 정수 그대로
assert hs._money('3400') == 3400 and isinstance(hs._money('3400'), int)
assert hs._money('') is None and hs._money('N/A') is None
ok += 1
# 개발 섹션은 개발로
assert models['15-453089-00']['group'] == '개발'
assert models['713-312133-006']['group'] == '양산'
ok += 1

print(f'전부 통과 ({ok}/11)')

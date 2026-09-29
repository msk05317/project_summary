"""잔특근 취합 엑셀 읽기.

    python3 backend/tests/test_overtime_import.py

이 파일이 지키는 것:
  - 병합된 두 층 머리줄을 펴서 읽는다. 안 펴면 두 번째 칸부터 빈칸이다.
  - '야간(가용인원)' 과 '야간비율(%)' 을 구별한다. 첫 줄만 보면 둘 다
    '야간' 이라 비율이 인원을 덮어쓴다 — 실제로 그랬다.
  - 'X' 를 0 으로 바꾸지 않는다.
  - 요약 시트(전체_주간현황)의 비고·일치확인을 사업부에 붙인다.
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook  # noqa: E402

import overtime_import as OI  # noqa: E402

ok = 0


def t(name, cond):
    global ok
    print(("  ok  " if cond else "  FAIL") + "  " + name)
    if not cond:
        raise SystemExit(1)
    ok += 1


# 실제 파일과 같은 머리줄 — 한국어 밑에 영어/베트남어가 줄바꿈으로 붙는다
TOP = ["사업부\nDivision", "요일\nThứ",
       "총 보유인원\nTổng nhân sự", None, None,
       "지원인력\nNhân sự hỗ trợ", None,
       "실제 가용인원\nNhân sự khả dụng", None, None, None,
       "직접 가용 인원\nNhân sự trực tiếp", None, None, None,
       "간접 가용 인원\nNhân sự gián tiếp", None, None, None,
       "잔업 인원\nSố người tăng ca", None, None,
       "잔업률\nTỷ lệ tăng ca", None, None,
       "특근 인원\nSố người làm thêm", None, None,
       "특근률\nTỷ lệ làm thêm", None, None,
       "비고\nGhi chú", "일치확인\nKiểm tra khớp"]
SUB = [None, None,
       "전체\nTổng", "주간\nCa ngày", "야간\nCa đêm",
       "지원\n인력\n(+)\nNhân sự được hỗ trợ", "지원\n인력\n(-)\nNhân sự điều đi",
       "전체\nTổng", "주간\n(가용인원)\nCa ngày", "야간\n(가용인원)\nCa đêm",
       "야간\n비율\n(%)\n(가용 기준)\nTỷ lệ ca đêm",
       "전체\nTổng", "주간(명)\nCa ngày (người)", "야간(명)\nCa đêm (người)",
       "야간\n비율(%)\nTỷ lệ ca đêm (%)",
       "전체\nTổng", "주간(명)\nCa ngày (người)", "야간(명)\nCa đêm (người)",
       "야간\n비율(%)\nTỷ lệ ca đêm (%)",
       "전체\nTổng", "직접\nTrực tiếp", "간접\nGián tiếp",
       "전체\nTổng", "직접\nTrực tiếp", "간접\nGián tiếp",
       "전체\nTổng", "직접\nTrực tiếp", "간접\nGián tiếp",
       "전체\nTổng", "직접\nTrực tiếp", "간접\nGián tiếp",
       None, None]

MERGES = ["C2:E2", "F2:G2", "H2:K2", "L2:O2", "P2:S2", "T2:V2",
          "W2:Y2", "Z2:AB2", "AC2:AE2", "AF2:AF3", "AG2:AG3",
          "A2:A3", "B2:B3"]

# PCB W39 실제 값 — 요일별 잔업 인원, 일요일만 특근
OTP = [(79, 78, 1), (98, 93, 5), (105, 104, 1),
       (127, 122, 5), (117, 112, 5), (119, 117, 2)]
WD = ["월\nThứ 2", "화\nThứ 3", "수\nThứ 4", "목\nThứ 5",
      "금\nThứ 6", "토\nThứ 7", "일\nChủ nhật"]


def div_sheet(wb, name, title):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    for j, v in enumerate(TOP):
        ws.cell(row=2, column=j + 1, value=v)
    for j, v in enumerate(SUB):
        ws.cell(row=3, column=j + 1, value=v)
    for m in MERGES:
        ws.merge_cells(m)
    for i, wd in enumerate(WD):
        r = 4 + i
        sun = (i == 6)
        row = [name if i == 0 else None, wd,
               441, 300, 141,          # 총보유
               None, None,             # 지원
               441, 300, 141, 0.32,    # 가용 + 야간비율
               378, 242, 136, 0.36,    # 직접
               63, 58, 5, 0.079,       # 간접
               # 잔업 인원 / 잔업률
               *(("X", "X", "X", 0, 0, 0) if sun
                 else (OTP[i][0], OTP[i][1], OTP[i][2],
                       round(OTP[i][0] / 441, 3), 0, 0)),
               # 특근 인원 / 특근률
               *((88, 86, 2, 0.2, 0.228, 0.032) if sun
                 else ("X", "X", "X", 0, 0, 0)),
               None, "OK"]
        for j, v in enumerate(row):
            ws.cell(row=r, column=j + 1, value=v)
    ws.cell(row=11, column=1, value="합계/평균\nTotal")
    ws.cell(row=11, column=3, value=441)
    ws.cell(row=11, column=20, value=108)   # 여기를 요일로 읽으면 안 된다
    return ws


def book():
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("W39_주간현황")
    ws["A1"] = "W39 사업부별 주간 잔특근 현황"
    ws["A2"] = "W39: 9/21~9/27  |  인원 = 주간 평균"

    div_sheet(wb, "PCB", "[PCB] W39 (9/21~9/27) 주간 잔특근 현황")
    div_sheet(wb, "네트워크", "[네트워크] W39 (9/21~9/27) 주간 잔특근 현황")

    # 요약 시트 — 사업부별 비고 · 일치확인이 여기 있다
    s = wb.create_sheet("전체_주간현황")
    s["A1"] = "[전사 사업부별] W39 (9/21~9/27) 주간 잔특근 현황"
    for j, v in enumerate(TOP):
        s.cell(row=2, column=j + 1, value=v)
    for j, v in enumerate(SUB):
        s.cell(row=3, column=j + 1, value=v)
    for m in MERGES:
        s.merge_cells(m)
    s.cell(row=4, column=1, value="PCB")
    s.cell(row=4, column=2, value="주간 평균\nTB tuần")
    s.cell(row=4, column=33, value="OK")
    s.cell(row=5, column=1, value="네트워크")
    s.cell(row=5, column=2, value="주간 평균\nTB tuần")
    s.cell(row=5, column=32, value="직접+간접 인원이 가용인원과 불일치")
    s.cell(row=5, column=33, value="확인필요")
    s.cell(row=6, column=1, value="전체 합계/평균")

    c = wb.create_sheet("확인사항")
    c["A1"] = "W39 전사 잔특근 취합 확인사항"
    c["A3"], c["B3"], c["C3"] = "사업부", "내용", "취합 처리"
    c["A4"], c["B4"], c["C4"] = "PRESS", "파일 2개 제출", "수정본 반영"

    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def main():
    p = OI.parse_workbook(book(), year=2026)

    t("주차를 찾았다", p["week"] == "2026-W39")
    t("기간을 찾았다", p["range"] == "9/21~9/27")
    t("사업부 2개 (요약 시트는 빼고)", sorted(p["divisions"]) == ["pcb", "네트워크"])
    t("못 읽은 줄 없음", p["warnings"] == [])

    d = p["divisions"]["pcb"]
    t("요일 7개", sorted(d["days"]) == sorted(
        ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]))
    t("'합계/평균' 줄에서 멈춘다", len(d["days"]) == 7)

    mon = d["days"]["mon"]
    t("총보유 441/300/141",
      mon["headcount"] == {"total": 441.0, "day": 300.0, "night": 141.0})
    # 야간비율(0.32)이 야간 인원(141)을 덮어쓰면 여기서 깨진다
    t("야간비율이 야간 인원을 안 덮는다", mon["available"]["night"] == 141.0)
    t("직접 주간(명)도 읽는다", mon["direct"]["day"] == 242.0)
    t("간접 야간(명)도 읽는다", mon["indirect"]["night"] == 5.0)
    t("월요일 잔업 79/78/1",
      mon["overtime"]["people"] == {"total": 79.0, "direct": 78.0, "indirect": 1.0})
    t("월요일 특근은 X → None",
      mon["special"]["people"]["total"] is None)

    sun = d["days"]["sun"]
    t("일요일 특근 88/86/2",
      sun["special"]["people"] == {"total": 88.0, "direct": 86.0, "indirect": 2.0})
    t("일요일 잔업은 X → None",
      sun["overtime"]["people"]["total"] is None)

    t("비고는 요약 시트에서 온다",
      p["divisions"]["네트워크"]["note"].startswith("직접+간접"))
    t("일치확인도 온다", p["divisions"]["네트워크"]["check"] == "확인필요")
    t("PCB 는 OK", p["divisions"]["pcb"]["check"] == "OK")
    t("확인사항 시트를 읽는다",
      len(p["issues"]) == 1 and p["issues"][0].startswith("PRESS"))

    # 저장 구조로 넣었을 때도 숫자가 맞는지 (여기서 틀리면 앱이 틀린다)
    import overtime as OT
    data = OT.put_week({}, p["week"], {"range": p["range"],
                                       "divisions": p["divisions"],
                                       "order": p["order"]})
    s = OT.for_division(data, "pcb")["summary"]
    t("잔업률 24.4%", s["overtime"]["rate"]["total"] == 0.2438)
    t("특근률 20.0%", s["special"]["rate"]["total"] == 0.1995)
    t("잔업 인원 108", s["overtime"]["people"]["total"] == 108)

    # 사업부 시트가 하나도 없으면 조용히 빈 걸 넣지 말고 실패시킨다
    from openpyxl import Workbook as WB
    wb = WB()
    wb.active["A1"] = "W39 아무것도 없음"
    b = io.BytesIO()
    wb.save(b)
    try:
        OI.parse_workbook(b.getvalue(), year=2026)
        t("사업부 시트 없으면 예외", False)
    except ValueError:
        t("사업부 시트 없으면 예외", True)

    print("test_overtime_import: 전부 통과 (%d개)" % ok)


if __name__ == "__main__":
    main()

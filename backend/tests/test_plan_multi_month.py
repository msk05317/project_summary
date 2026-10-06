# -*- coding: utf-8 -*-
"""한 파일에 든 달을 전부 읽는가.

엔클로저 1년치 계획(시트 13장)을 올렸더니 아무것도 안 들어갔다. 화면이
보고 있는 달(2026-10)이 파일에 없어서 _pbx_pick_sheet 가 멈췄고, '주차표
없는 파일' 되돌림이 그걸 받아 모델 목록으로 읽어 버렸다.

여기서 지키는 것
  · 화면의 달이 파일에 없어도 파일에 든 달을 전부 읽는다
  · 화면의 달이 파일에 있으면 그 달이 맨 앞 (미리보기 표가 그 달이다)
  · 달이 겹쳐 들어오지 않는다
  · 주차표가 아예 없는 파일은 **그대로 멈춘다** — 멈춰야 위에서 받아
    '모델 값만' 이나 '개발 프로세스' 로 넘긴다
"""
import ast
import sys
from io import BytesIO
from pathlib import Path

BACK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACK))

from openpyxl import Workbook, load_workbook   # noqa: E402

import pbx_plan_import as pbx                  # noqa: E402
import week_calendar as wcal                   # noqa: E402
import weekly_template as wt                   # noqa: E402

FAIL = []


def ok(c, m):
    if not c:
        FAIL.append(m)


class HTTPException(Exception):
    def __init__(self, status_code=400, detail=""):
        self.status_code, self.detail = status_code, detail
        super().__init__(detail)


# main.py 에서 함수만 꺼내 쓴다 (fastapi 없이 돌리려고)
NS = {"HTTPException": HTTPException, "pbx_plan_import": pbx}
_tree = ast.parse((BACK / "main.py").read_text(encoding="utf-8"))
_want = {"_pbx_pick_sheet", "_pbx_pick_months"}
for _node in _tree.body:
    if isinstance(_node, ast.FunctionDef) and _node.name in _want:
        exec(compile(ast.Module([_node], []), "main.py", "exec"), NS)
for _w in _want:
    ok(_w in NS, "main.py 에서 %s 를 못 꺼냈다" % _w)
pick = NS.get("_pbx_pick_months")


MONTHS = ["2026-11", "2026-12", "2027-01", "2027-02", "2027-03"]


def book(months, filled=True):
    ms = [{"id": "A", "name": "A-001", "group": "양산",
           "weekly_plan": {mo: {wcal.get_month_weeks(mo)[0]: {"plan": 10, "actual": 7}}
                           for mo in months}}]
    sheets = [{"month": mo, "weeks": wcal.get_month_weeks(mo)} for mo in months]
    data = wt.build_weekly_template("T", sheets=sheets, models=ms, filled=filled)
    return load_workbook(BytesIO(data), data_only=True)


# ── 1) 화면의 달이 파일에 없다 → 전부, 파일 순서대로 ────────────────
wb = book(MONTHS)
got = pick(wb, "2026-10")
ok([p["month"] for p in got] == MONTHS,
   "1 읽은 달: %r" % [p["month"] for p in got])
ok(all(p["rows"] for p in got), "1 줄이 빈 달이 있다")


# ── 2) 화면의 달이 파일에 있다 → 그 달이 맨 앞 ──────────────────────
got = pick(wb, "2027-02")
ok(got[0]["month"] == "2027-02", "2 맨 앞: %r" % got[0]["month"])
ok(sorted(p["month"] for p in got) == sorted(MONTHS),
   "2 달이 빠졌다: %r" % [p["month"] for p in got])
ok(len({p["month"] for p in got}) == len(got),
   "2 같은 달이 두 번: %r" % [p["month"] for p in got])


# ── 3) 한 달짜리 파일 ───────────────────────────────────────────────
got = pick(book(["2026-11"]), "")
ok([p["month"] for p in got] == ["2026-11"], "3 한 달: %r" % [p["month"] for p in got])


# ── 4) 주차표가 아예 없으면 멈춘다 ──────────────────────────────────
#
# 멈춰야 한다. 위에서 그 멈춤을 받아 '개발 프로세스' 나 '모델 값만' 으로
# 넘긴다 — 여기서 빈 목록을 돌려주면 그 길이 막힌다.
plain = Workbook()
ws = plain.active
ws.append(["모델명", "판가($)", "PO수량"])
ws.append(["A-001", 100, 5])
stopped = False
try:
    pick(plain, "2026-10")
except HTTPException:
    stopped = True
ok(stopped, "4 주차표 없는 파일에서 안 멈췄다")


# ── 5) 빈 양식(숫자 없음)도 멈춘다 ──────────────────────────────────
#
# 줄이 통째로 비면 파서가 건너뛰니 '읽을 달' 이 하나도 없다. 이것도
# 멈춰야 '모델 값만' 으로 넘어가서 판가·PO 라도 들어간다.
stopped = False
try:
    pick(book(MONTHS, filled=False), "")
except HTTPException:
    stopped = True
ok(stopped, "5 빈 양식에서 안 멈췄다")


if FAIL:
    print("FAIL %d" % len(FAIL))
    for m in FAIL:
        print(" -", m)
    sys.exit(1)
print("ok  여러 달 한 파일 — 5묶음")

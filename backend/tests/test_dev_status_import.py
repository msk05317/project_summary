# 개발현황 엑셀(하바플레이트)을 모델에 반영하는 길.
#   python3 backend/tests/test_dev_status_import.py
#
# 이 파일은 열 이름이 모델 등록 양식과 겹친다 ('모델', '판가', '비고').
# 일반 파서에 맡기면 두 가지가 어긋난다.
#   1. 'PO 수량' 과 'PO 잔량' 이 둘 다 po 로 시작해서 잔량이 수량을 덮는다
#   2. '구분' 열이 없어서 개발품이 양산으로 들어간다
import ast, datetime, pathlib, sys

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import dev_status_import as di          # noqa: E402

SRC = (ROOT / 'main.py').read_text(encoding='utf-8')
tree = ast.parse(SRC)
ok = 0

# ── 1) 파서 ─────────────────────────────────────────────
import openpyxl                          # noqa: E402
from openpyxl import Workbook            # noqa: E402

def sheet(rows, head=('No', '모델', '소재', '판가　', '개발종류', '연간수량',
                      'PO 수량', '출하수량', 'PO 잔량', '고객요청일',
                      '가공 완료', '비고')):
    wb = Workbook()
    ws = wb.active
    for c, h in enumerate(head, start=2):
        ws.cell(row=2, column=c, value=h)
    for i, r in enumerate(rows):
        for c, v in enumerate(r, start=2):
            ws.cell(row=3 + i, column=c, value=v)
    return wb

D = datetime.datetime
W = sheet([
    [1, '713-C08296-001', '다이넥스', 3400, 'RPM', 'N/A', 34, 27, 7, 'PO 취소', None, '→ 캔슬비용 청구 협의중'],
    [2, '713-E28195-001', '다이넥스', 3400, 'RPM', 'N/A', 2, 1, 1, D(2026, 9, 26), D(2026, 9, 21), None],
    [3, '713-D23686-001', '다이넥스', 3400, 'RPM', 'N/A', 6, 0, 6, D(2026, 10, 3), '확인 중', '고객 도면 확인 중'],
    ['Total', None, None, None, None, None, 42, 28, 14, None, None, None],
])

assert di.sniff(W), '개발현황 엑셀을 못 알아본다'
ok += 1

res = di.parse_workbook(W)
rows = res['rows']
assert len(rows) == 3, f'Total 줄까지 세거나 빠뜨렸다: {len(rows)}'
ok += 1

r0, r1, r2 = rows
# PO 수량이 잔량에 덮이면 안 된다 (34 ≠ 7)
assert r0['po_qty'] == 34 and r0['shipped_qty'] == 27, r0
ok += 1
assert r0['cancelled'] is True and r0['request_text'] == 'PO 취소', r0
ok += 1
# 날짜 칸의 '-' 를 '미정' 으로 세면 멀쩡한 줄이 전부 확인 중이 된다
assert r1['machining_date'] == '2026-09-21' and r1['pending'] is False, r1
ok += 1
assert r2['machining_date'] == '' and r2['machining_text'] == '확인 중', r2
assert r2['pending'] is True, r2
ok += 1
assert r1['dev_type'] == 'RPM' and r1['price'] == 3400, r1
ok += 1

# 모델 등록 양식은 이 파서가 건드리면 안 된다 ('개발종류' 열이 없다)
tpl = sheet([], head=('파트넘버', '모델명', '구분', '유형', '판가($)',
                      '재료비($)', 'PO수량', '실적수량', '비고'))
assert not di.sniff(tpl), '모델 등록 양식까지 개발현황으로 본다'
ok += 1

# ── 2) main.py 연결부 ───────────────────────────────────
imp = SRC[SRC.index('async def admin_models_import'):]
imp = imp[:imp.index('\n@app.')]
assert 'dev_status_import' in imp and '_apply_dev_status' in imp, \
    '업로드 창이 개발현황을 안 본다'
assert imp.index('dev_status_import') < imp.index("header_row = None"), \
    '일반 파서가 먼저라 개발현황이 잘못 읽힌다'
ok += 1

# 'PO 잔량' 이 'PO 수량' 을 덮던 버그
assert 'elif k.startswith("po"):' not in imp, \
    "'PO 잔량' 이 아직 PO 수량 자리를 덮는다"
assert 'k in ("po", "po수량"' in imp, 'PO 수량 열을 정확히 안 짚는다'
ok += 1

# ── 3) 반영 함수 ────────────────────────────────────────
fn = {n.name: ast.get_source_segment(SRC, n) for n in tree.body
      if isinstance(n, ast.FunctionDef) and n.name == '_apply_dev_status'}
assert fn, '_apply_dev_status 가 없다'
body = fn['_apply_dev_status']
assert "m[\"group\"] = \"개발\"" in body, '개발품으로 안 넣는다'
assert '"machining"' in body and 'st["expected"]' in body, \
    "'가공 완료' 를 가공 (조립) 계획일로 안 넣는다"
assert '"드롭예정"' in body, 'PO 취소를 드롭예정으로 안 바꾼다'
assert 'part_number' in body, '파트넘버로 안 맞춘다'
ok += 1

print(f'전부 통과 · {ok}개 항목')

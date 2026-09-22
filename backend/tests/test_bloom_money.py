# 블룸 '금액 실적' 시트 → 보드의 money.
#   python3 backend/tests/test_bloom_money.py
#
# 계획 대비 실적 카드 맨 위 '매출' 칸이 이걸 읽는다.
# '출하계획' · '매출계획' 이 두 번 나온다 (영업 계획 / 예상실적) —
# 계획으로는 예상실적(뒤쪽)을 쓴다. 시트의 달성율도 그걸 분모로 쓴다.
import ast, pathlib, sys

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import bloom_daily_import as bd          # noqa: E402
from openpyxl import Workbook            # noqa: E402

ok = 0
wb = Workbook()
ws = wb.active
ws.title = '금액 실적'
ws['B2'] = '2026년도 BLOOM 출하 계획'
H1 = ['제품명', 'CODE', '출하 국가', '단가\nUSD', '영업 계획', None, None, '예상실적']
H2 = [None, None, None, None, '09월\n출하계획', '09월\n매출계획\nUSD', None,
      '09월\n출하계획', '09월\n매출계획\nUSD', '09월 출하\n완료', '09월\n매출완료\nUSD', '잔량']
for i, v in enumerate(H1):
    ws.cell(3, 2 + i, v)
for i, v in enumerate(H2):
    ws.cell(4, 2 + i, v)
ROWS = [
    ['SL7', 711720, '인도', 5500, 700, 3850000, None, 708, 3894000, 208, 1144000, 500],
    ['WDM', 126391, '한국', 7990, 42, 335580, None, 64, 511360, None, 0, 64],
    ['합계', None, None, None, 742, 4185580, None, 772, 9999999, 208, 1144000, 564],
]
for r, row in enumerate(ROWS, start=5):
    for i, v in enumerate(row):
        ws.cell(r, 2 + i, v)

m = bd.parse_money(wb)
assert m and m['sheet'] == '금액 실적', m
assert m['month'] == '9', m['month']
ok += 1

it = {x['item']: x for x in m['items']}
assert set(it) == {'SL7', 'WDM'}, '합계 줄을 품목으로 셌다'
sl7 = it['SL7']
# 계획은 뒤쪽(예상실적), 영업 계획은 sales_*
assert sl7['plan_qty'] == 708 and sl7['plan_usd'] == 3894000, sl7
assert sl7['sales_qty'] == 700 and sl7['sales_usd'] == 3850000, sl7
assert sl7['done_qty'] == 208 and sl7['done_usd'] == 1144000, sl7
assert sl7['country'] == '인도' and sl7['price'] == 5500
ok += 1

# 합계는 우리가 더한다. 시트 합계 줄이 틀려 있어도 (9,999,999) 따라가지 않는다
assert m['total']['plan_usd'] == 3894000 + 511360, m['total']
assert m['total']['done_usd'] == 1144000
assert m['stated_total']['plan_usd'] == 9999999
ok += 1

# 금액 시트가 없는 파일이면 None (일 보고만 올린 경우)
wb2 = Workbook()
assert bd.parse_money(wb2) is None
ok += 1

# 연결부
SRC = (ROOT / 'main.py').read_text(encoding='utf-8')
up = SRC[SRC.index('def _bloom_parse_upload'):]
up = up[:up.index('\n@app.')]
assert 'parse_money(wb)' in up, '업로드할 때 금액 시트를 안 읽는다'
mg = SRC[SRC.index('def _bloom_merge'):]
mg = mg[:mg.index('\ndef ')]
assert '"money": parsed.get("money") or old.get("money")' in mg, '금액이 보드에 안 남는다'
sl = SRC[SRC.index('def _bloom_slice'):]
sl = sl[:sl.index('\ndef ')]
assert 'out.pop("money", None)' in sl, '품목 화면에 전체 금액이 붙는다'
ok += 1

print(f'전부 통과 · {ok}개 항목')

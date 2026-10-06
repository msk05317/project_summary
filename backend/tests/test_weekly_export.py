# -*- coding: utf-8 -*-
"""내보내기 — 채워 내려준 양식을 그대로 다시 읽는다.

'양식 받기' 가 지금 올라간 숫자를 채워서 내려준다. 그 파일에 이번 주만
더 적어 그대로 올리는 게 보통의 한 주다. 그래서 **채워 낸 것이 같은
자리로 돌아오는지**가 이 파일이 지키는 것이다.

여기서 어긋나면 눈에 안 보이는 사고가 난다 — 받은 파일을 손 하나 안 대고
그대로 올렸는데 지난 주 숫자가 다른 주차로 들어가거나 0 으로 깔린다.

### 파서가 빈 칸을 어떻게 보는가 (중요)

pbx_plan_import._rows_of 는 이렇게 읽는다.
  · 그 줄의 주차 칸이 하나도 안 적혀 있으면 → 그 줄을 건너뛴다
  · 숫자가 하나라도 있으면 → 같은 줄의 빈 칸은 0 으로 들어간다

그래서 '빈 칸은 그대로 둬라' 는 **줄 단위**로만 참이다. 내보내기가
이것을 실제로 막는다 — 받은 파일에는 지난 주차가 이미 적혀 있다.
아래 2번이 그걸 지킨다.

특히 보는 것
  · 채워 낸 계획·실적이 그 모델, 그 주차로 돌아온다
  · 손 안 댄 파일을 그대로 올려도 바뀌는 게 없다 (멱등)
  · 0 은 셀을 비워서 나간다 (0 을 찍어 두면 '안 정한 것' 과 구분이 안 된다)
  · 미달 사유도 같은 주차로 돌아오고, 모델 줄로 읽히지 않는다
  · 다른 달 숫자가 이 달 시트로 새지 않는다
  · filled=False 는 빈 양식이다 — 올려도 반영할 줄이 없다
"""
import sys
from io import BytesIO
from pathlib import Path

BACK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACK))

from openpyxl import load_workbook          # noqa: E402

import pbx_plan_import as pbx               # noqa: E402
import week_calendar as wcal                # noqa: E402
import weekly_template as wt                # noqa: E402

FAIL = []


def ok(c, m):
    if not c:
        FAIL.append(m)


def sheets_for(months, notes=None):
    out = []
    for m in months:
        s = {"month": m, "weeks": wcal.get_month_weeks(m)}
        if notes and m in notes:
            s["notes"] = notes[m]
        out.append(s)
    return out


def build(models, months=("2026-10",), filled=True, notes=None):
    return wt.build_weekly_template("T", sheets=sheets_for(list(months), notes),
                                    models=models, filled=filled)


def parse(data, month):
    wb = load_workbook(BytesIO(data))
    return pbx.parse_sheet(wb[month], month=month)


def live(parsed):
    """{라벨: {주차: (계획, 실적)}} — 0 뿐인 칸은 버린다.

    파서가 빈 칸을 0 으로 채워 주니, 실제로 값이 있는 칸만 봐야 '무엇이
    돌아왔는가' 를 볼 수 있다.
    """
    out = {}
    for r in parsed["rows"]:
        got = {}
        for w, v in (r["weeks"] or {}).items():
            p, a = v.get("plan") or 0, v.get("actual") or 0
            if p or a:
                got[w] = (p, a)
        out[r["label"]] = got
    return out


W = wcal.get_month_weeks("2026-10")          # W41..W44
W9 = wcal.get_month_weeks("2026-09")         # ..W40

MODELS = [
    {"id": "A-413", "name": "A-413", "part_number": "PN-413", "dev_type": "413",
     "group": "양산", "price": 1250.5, "po_qty": 240,
     "weekly_plan": {"2026-10": {W[0]: {"plan": 20, "actual": 12},
                                 W[2]: {"plan": 10, "actual": 0}},
                     "2026-09": {W9[-1]: {"plan": 99, "actual": 88}}}},
    {"id": "B-009", "name": "B-009", "part_number": "PN-009", "dev_type": "009",
     "group": "양산",
     "weekly_plan": {"2026-10": {W[1]: {"plan": 30, "actual": 30}}}},
    # 주차 숫자가 아예 없는 모델 — 빈 줄로 나가서 파서가 건너뛴다
    {"id": "C-new", "name": "C-new", "part_number": "PN-new", "dev_type": "413",
     "group": "개발"},
]


# ── 1) 채워 낸 숫자가 그 자리로 돌아온다 ────────────────────────────
p = parse(build(MODELS), "2026-10")
got = live(p)
ok(p["month"] == "2026-10", "1 달: %r" % p["month"])
ok(set(p["weeks"]) == set(W), "1 주차: %r" % p["weeks"])

a = got.get("PN-413", {})
ok(a == {W[0]: (20, 12), W[2]: (10, 0)}, "1 A 가 다르게 돌아왔다: %r" % a)
b = got.get("PN-009", {})
ok(b == {W[1]: (30, 30)}, "1 B 가 다르게 돌아왔다: %r" % b)
# 숫자가 없는 모델은 줄이 비어 있어서 파서가 건너뛴다 — 건드릴 게 없다
ok("PN-new" not in got, "1 빈 줄이 반영 대상으로 들어왔다: %r" % list(got))


# ── 2) 손 안 대고 그대로 올려도 바뀌는 게 없다 (멱등) ───────────────
#
# 받은 파일을 그대로 올리는 일은 자주 생긴다 ('열어 봤다가 저장만 했다').
# 그때 한 칸이라도 달라지면 모르는 사이에 데이터가 변한다. 파서가 빈 칸을
# 0 으로 채우는데도 이게 성립하는 이유가 내보내기의 요점이다 — 지난
# 주차가 이미 적혀 있으니 0 으로 덮일 칸이 없다.
by_label = {}
for m in MODELS:
    by_label[m["part_number"]] = m["id"]
    by_label[m["name"]] = m["id"]


def as_store(parsed):
    """파서 결과 → {모델id: {주차: {'plan','actual'}}} (저장되는 모양)"""
    out = {}
    for r in parsed["rows"]:
        key = by_label.get(r["label"])
        if not key:
            continue
        out[key] = {w: {"plan": v.get("plan") or 0, "actual": v.get("actual") or 0}
                    for w, v in (r["weeks"] or {}).items()}
    return out


round1 = as_store(p)
for m in MODELS:
    have = (m.get("weekly_plan") or {}).get("2026-10") or {}
    back = round1.get(m["id"])
    if not have:
        ok(back is None, "2 %s 는 반영 대상이 아니어야 한다: %r" % (m["id"], back))
        continue
    for w, v in have.items():
        want = (v.get("plan") or 0, v.get("actual") or 0)
        gotc = ((back or {}).get(w) or {})
        ok((gotc.get("plan") or 0, gotc.get("actual") or 0) == want,
           "2 %s %s: %r 이어야 하는데 %r" % (m["id"], w, want, gotc))
    # 적힌 적 없는 주차에 0 이 들어가도 손해는 없다 — 원래도 없던 값이다.
    # 다만 **원래 있던 값이 0 으로 덮이는 일**은 없어야 한다.
    for w in W:
        if w in have:
            continue
        gotc = ((back or {}).get(w) or {})
        ok((gotc.get("plan") or 0) == 0 and (gotc.get("actual") or 0) == 0,
           "2 %s %s 에 없던 숫자가 생겼다: %r" % (m["id"], w, gotc))

# 두 바퀴 — 돌려받은 것으로 다시 양식을 만들어도 같아야 한다
m2 = []
for m in MODELS:
    c = dict(m)
    c["weekly_plan"] = {"2026-10": round1.get(m["id"]) or {}}
    m2.append(c)
round2 = as_store(parse(build(m2), "2026-10"))
ok(round1 == round2, "2 두 바퀴에서 달라졌다:\n  1 %r\n  2 %r" % (round1, round2))


# ── 3) 0 은 셀을 비워서 나간다 ──────────────────────────────────────
#
# 0 을 찍어 두면 '아직 안 정한 것' 과 '정말 0' 이 구분되지 않는다. 적어도
# 엑셀을 눈으로 볼 때 0 이 깔린 표는 어디를 봐야 하는지 알 수 없다.
z = [{"id": "Z", "name": "Z", "group": "양산",
      "weekly_plan": {"2026-10": {W[0]: {"plan": 0, "actual": 0},
                                  W[1]: {"plan": 5, "actual": 0}}}}]
wz = load_workbook(BytesIO(build(z)))["2026-10"]
wrow, wcols = pbx._find_week_header(wz)
srow, pairs, _nc = pbx._find_sub_header(wz, wrow, wcols)
first = min(p_ for p_, _a in pairs.values())
mcol = pbx._find_model_col(wz, srow, first)
rz = next(r for r in range(srow + 1, wz.max_row + 1)
          if str(wz.cell(r, mcol).value or "").strip() == "Z")
byw = {"W%02d" % n: pc for n, (pc, _a) in pairs.items()}
ok(wz.cell(rz, byw[W[0]]).value is None,
   "3 0 이 셀에 찍혔다: %r" % wz.cell(rz, byw[W[0]]).value)
ok(wz.cell(rz, byw[W[1]]).value == 5,
   "3 5 가 안 나갔다: %r" % wz.cell(rz, byw[W[1]]).value)
gz = live(parse(build(z), "2026-10")).get("Z") or {}
ok(gz == {W[1]: (5, 0)}, "3 돌아온 값: %r" % gz)


# ── 4) 미달 사유도 돌아온다 ─────────────────────────────────────────
NOTES = {W[0]: "자재 입고 지연 — 10/9 입고 예정", W[2]: "고객 요청으로 출하 보류"}
data = build(MODELS, notes={"2026-10": NOTES})
ws = load_workbook(BytesIO(data))["2026-10"]
wrow, wcols = pbx._find_week_header(ws)
srow, pairs, _nc = pbx._find_sub_header(ws, wrow, wcols)
mcol = pbx._find_model_col(ws, srow, min(p_ for p_, _a in pairs.values()))
back = pbx._week_notes(ws, srow, mcol, pairs)
ok(back.get(W[0]) == NOTES[W[0]], "4 %s 사유: %r" % (W[0], back.get(W[0])))
ok(back.get(W[2]) == NOTES[W[2]], "4 %s 사유: %r" % (W[2], back.get(W[2])))
ok(W[1] not in back and W[3] not in back,
   "4 안 적은 주에 사유가 생겼다: %r" % back)
# 사유 줄·합계 줄이 모델 줄로 읽히면 안 된다
labels = [r["label"] for r in parse(data, "2026-10")["rows"]]
ok(all("미달" not in x and "합계" not in x for x in labels),
   "4 사유·합계 줄이 모델로 읽혔다: %r" % labels)
# parse_sheet 가 돌려주는 week_notes 로도 같이 와야 한다 (화면이 그걸 쓴다)
ok((parse(data, "2026-10").get("week_notes") or {}).get(W[0]) == NOTES[W[0]],
   "4 parse_sheet 의 week_notes 가 비었다")


# ── 5) 달이 섞이지 않는다 ───────────────────────────────────────────
#
# A 는 2026-09 에도 숫자가 있다. 10월 시트로 새어 나오면 9월 실적이
# 10월로 옮겨 적힌다.
two = build(MODELS, months=("2026-09", "2026-10"))
a9 = live(parse(two, "2026-09")).get("PN-413") or {}
a10 = live(parse(two, "2026-10")).get("PN-413") or {}
ok(a9 == {W9[-1]: (99, 88)}, "5 9월: %r" % a9)
ok(a10 == {W[0]: (20, 12), W[2]: (10, 0)}, "5 10월: %r" % a10)


# ── 6) filled=False 는 빈 양식 ──────────────────────────────────────
#
# 지금 화면은 늘 채워 보내지만 함수는 빈 양식도 만들 수 있어야 한다.
# 빈 양식을 그대로 올리면 **반영할 줄이 없다** — 그게 맞는 동작이다.
# (줄이 통째로 비어 있으면 파서가 건너뛴다)
pe = parse(build(MODELS, filled=False), "2026-10")
ok(pe["rows"] == [], "6 빈 양식에서 줄이 읽혔다: %r" % [r["label"] for r in pe["rows"]])
# 모델 칸은 그대로 채워져 있어야 한다 (사람이 숫자만 적게)
we = load_workbook(BytesIO(build(MODELS, filled=False)))["2026-10"]
names = [str(we.cell(r, c).value or "")
         for r in range(wt.FIRST_ROW, wt.FIRST_ROW + len(MODELS))
         for c in (1, 2)]
ok(any("PN-413" in x or "A-413" in x for x in names),
   "6 빈 양식에 모델 칸이 안 채워졌다: %r" % names[:6])


if FAIL:
    print("FAIL %d" % len(FAIL))
    for m in FAIL:
        print(" -", m)
    sys.exit(1)
print("ok  내보내기 — 6묶음")

# -*- coding: utf-8 -*-
"""모델 목록 엑셀 — 양식대로 채워 올리면 적은 것만 들어가야 한다.

공식 양식(model_template.py)을 받아 채워 올리는 게 정해진 길이다.
그런데 그 길에 두 가지가 있었다.

  · '재료비율' 이 '재료비' 자리를 빼앗았다. 둘 다 '재료비' 로 시작하는데
    재료비율이 뒤에 있어서 덮어썼다. 그 칸은 수식이라 읽으면 None 이고,
    None → 0 이라 올릴 때마다 재료비가 전부 0 이 됐다.
  · 빈 칸이 기존 값을 0 으로 지웠다. PO·출하만 막혀 있었다.

임포터 본문은 main.py 안에 있고 main.py 는 여기서 import 할 수 없다
(fastapi 등이 필요하다). 그래서 열을 찾는 부분과 값을 쓰는 부분을
main.py 에서 그대로 떼어 와 돌린다. 떼어 온 코드가 원본과 어긋나면
알아채지 못하므로, 원본에 그 구절이 남아 있는지도 같이 본다.
"""
import re
import sys
from io import BytesIO
from pathlib import Path

HERE = Path(__file__).resolve().parent
BACK = HERE.parent
sys.path.insert(0, str(BACK))

from openpyxl import load_workbook          # noqa: E402
import model_template as mt                 # noqa: E402

FAIL = []


def ok(cond, msg):
    if not cond:
        FAIL.append(msg)


MAIN = (BACK / "main.py").read_text(encoding="utf-8")


# ── 떼어 온 코드가 원본과 같은지 ────────────────────────────────────
for frag, why in [
    ('elif "비율" in k:', "재료비율을 걸러내는 줄이 main.py 에 없다"),
    ('elif k in ("유형", "개발유형", "type", "devtype"):', "유형 정규화가 main.py 에 없다"),
    ("def _filled(col):", "_filled 헬퍼가 main.py 에 없다"),
    ('if _filled(col_map.get("price")):', "판가가 빈 칸 보호를 안 받는다"),
    ('if _filled(col_map.get("material_cost")):', "재료비가 빈 칸 보호를 안 받는다"),
    ('if _filled(col_map.get("group")):', "구분이 빈 칸 보호를 안 받는다"),
    ('if _filled(col_map.get("po_qty")):', "PO 가 빈 칸 보호를 안 받는다"),
]:
    ok(frag in MAIN, why)


# ── main.py 의 열 찾기 로직 (그대로 옮김) ───────────────────────────
def find_columns(ws):
    col_map, header_row = {}, None
    for r in range(1, min(ws.max_row, 12) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                continue
            t = str(v).strip()
            k = t.replace(" ", "").replace(".", "").lower()
            if k in ("모델", "모델명"):
                header_row = r
                col_map["name"] = c
            elif k == "구분":
                col_map["group"] = c
            elif k in ("유형", "개발유형", "type", "devtype"):
                col_map["dev_type"] = c
            elif "비율" in k:
                continue
            elif k.startswith("판가"):
                col_map["price"] = c
            elif k.startswith("재료비"):
                col_map["material_cost"] = c
            elif k in ("파트넘버", "파트번호", "품번", "partno", "partnumber"):
                col_map["part_number"] = c
            elif k in ("po", "po수량", "poqty", "poq'ty", "발주수량"):
                col_map["po_qty"] = c
            elif k in ("실적수량", "출하수량", "출하실적", "실적", "출하"):
                col_map["shipped_qty"] = c
            elif k == "비고":
                col_map["note"] = c
        if header_row:
            break
    return header_row, col_map


def apply_rows(ws, header_row, col_map, models):
    """main.py 의 '기존 모델 갱신' 부분만 떼어 왔다."""
    by_name = {str(m.get("name", "")).strip().lower(): m for m in models}
    by_pn = {}
    for m in models:
        p = str(m.get("part_number") or "").strip().lower()
        if p:
            by_pn[p] = m
    added = updated = 0
    for r in range(header_row + 1, ws.max_row + 1):
        name_v = ws.cell(row=r, column=col_map["name"]).value
        pn_v = (ws.cell(row=r, column=col_map["part_number"]).value
                if col_map.get("part_number") else None)
        pn = str(pn_v).strip() if pn_v is not None else ""
        if (name_v is None or not str(name_v).strip()) and not pn:
            continue
        name = str(name_v).strip() if name_v is not None else pn
        group = (str(ws.cell(row=r, column=col_map.get("group", 0)).value or "").strip()
                 if col_map.get("group") else "")
        group = "개발" if group == "개발" else "양산"
        dev_type = (str(ws.cell(row=r, column=col_map.get("dev_type", 0)).value or "").strip()
                    if col_map.get("dev_type") else "")

        def _num(col):
            if not col:
                return 0
            v = ws.cell(row=r, column=col).value
            try:
                return float(str(v).replace(",", "").replace("$", "").strip() or 0)
            except Exception:
                return 0

        def _filled(col):
            if not col:
                return False
            v = ws.cell(row=r, column=col).value
            return v is not None and str(v).strip() != ""

        price = _num(col_map.get("price"))
        mcost = _num(col_map.get("material_cost"))
        note = (str(ws.cell(row=r, column=col_map["note"]).value or "").strip()
                if col_map.get("note") else "")
        existing = by_pn.get(pn.lower()) if pn else by_name.get(name.lower())
        if existing is not None:
            if _filled(col_map.get("group")):
                existing["group"] = group
            if dev_type:
                existing["dev_type"] = dev_type
            if _filled(col_map.get("price")):
                existing["price"] = price
            if _filled(col_map.get("material_cost")):
                existing["material_cost"] = mcost
            if pn:
                existing["part_number"] = pn
                existing["name"] = name
            if note:
                existing["note"] = note
            updated += 1
        else:
            existing = {"id": pn or name, "name": name, "group": group,
                        "dev_type": dev_type, "price": price,
                        "material_cost": mcost, "status": "정상", "progress": 0}
            models.append(existing)
            added += 1
        if _filled(col_map.get("po_qty")):
            existing["po_qty"] = int(_num(col_map["po_qty"]))
        if _filled(col_map.get("shipped_qty")):
            existing["shipped_qty"] = int(_num(col_map["shipped_qty"]))
    return added, updated


# ── 실제 모양의 모델들 (엔클로저에서 본 그대로) ─────────────────────
def sample():
    return [
        {"id": "853-800575-413", "name": "853-800575-413", "group": "양산",
         "dev_type": "413", "price": 9100, "material_cost": 2306.75,
         "po_qty": 220, "shipped_qty": 163, "status": "정상", "progress": 0},
        {"id": "853-800575-301", "name": "853-800575-301", "group": "양산",
         "dev_type": "009", "price": 0, "material_cost": 0,
         "po_qty": 0, "shipped_qty": 0, "status": "정상", "progress": 0},
        {"id": "853-800575-611", "name": "853-800575-611", "group": "개발",
         "dev_type": "009", "price": 10798.95, "material_cost": 4383.18,
         "po_qty": 0, "shipped_qty": 0, "status": "정상", "progress": 0,
         "note": "9월 승인 대기"},
        {"id": "714-025898-009", "name": "자빌 009", "group": "양산",
         "dev_type": "009", "price": 2931.15, "material_cost": 468.15,
         "part_number": "714-025898-009", "po_qty": 1117, "shipped_qty": 1090,
         "status": "정상", "progress": 0},
    ]


WATCH = ("group", "dev_type", "price", "material_cost", "po_qty", "shipped_qty",
         "name", "part_number", "note", "status", "progress")

# COLS 순서: 파트넘버 모델명 구분 유형 판가 재료비 재료비율 PO수량 실적수량 비고
_COL = {"part_number": 1, "name": 2, "group": 3, "dev_type": 4, "price": 5,
        "material_cost": 6, "po_qty": 8, "shipped_qty": 9, "note": 10}


def filled_sheet(models, blank=()):
    """공식 빈 양식을 받아 사람이 채운 것처럼 값을 써 넣는다.

    blank 에 든 열은 비워 둔다 — '안 적은 칸' 을 흉내 내려는 것이다.
    """
    wb = load_workbook(BytesIO(mt.build_model_template("엔클로저", ["413", "009"])))
    ws = wb.active
    for i, m in enumerate(models):
        r = 3 + i
        for f, c in _COL.items():
            if f in blank:
                continue
            v = m.get(f)
            if v in (None, "", 0) and f in ("price", "material_cost", "po_qty",
                                            "shipped_qty"):
                continue          # 0 은 빈 칸으로 — 사람이 그렇게 적는다
            if v not in (None, ""):
                ws.cell(r, c).value = v
    buf = BytesIO(); wb.save(buf)
    return load_workbook(BytesIO(buf.getvalue()), data_only=True).active


def snap(models):
    return {m["id"]: {k: m.get(k) for k in WATCH} for m in models}


# ── 1. 내보내고 그대로 되올리면 아무것도 안 바뀐다 ──────────────────
before = sample()
want = snap(before)
ws = filled_sheet(before)
hr, cm = find_columns(ws)

ok(hr == 2, "머리글 행을 못 찾았다: %r" % hr)
ok(cm.get("material_cost") == 6,
   "재료비가 %r열로 잡혔다 (6이어야 한다 — 7은 재료비율)" % cm.get("material_cost"))
ok(cm.get("dev_type") == 4, "유형 열을 못 찾았다: %r" % cm.get("dev_type"))
ok(set(cm) == {"name", "group", "dev_type", "price", "material_cost",
               "part_number", "po_qty", "shipped_qty", "note"},
   "찾은 열이 다르다: %s" % sorted(cm))

after = sample()
added, updated = apply_rows(ws, hr, cm, after)
ok(added == 0, "양식대로 올렸는데 %d개가 새로 생겼다" % added)
ok(updated == len(before), "갱신 수가 안 맞다: %d" % updated)

got = snap(after)
for mid, w in want.items():
    g = got.get(mid)
    ok(g is not None, "%s 가 사라졌다" % mid)
    if g is None:
        continue
    for k in WATCH:
        a, b = w.get(k), g.get(k)
        if k in ("price", "material_cost"):
            same = abs(float(a or 0) - float(b or 0)) < 0.005
        else:
            same = (a or "") == (b or "") or a == b
        ok(same, "%s · %s 가 %r → %r 로 바뀌었다" % (mid, k, a, b))

# ── 2. 빈 칸은 지금 값을 지킨다 ─────────────────────────────────────
# 사람이 모델명과 PO 만 적고 판가·재료비·구분은 안 적은 양식
m2 = sample()
m2[0]["po_qty"] = 999
ws2b = filled_sheet(m2, blank=("price", "material_cost", "group"))
hr2b, cm2b = find_columns(ws2b)
after2 = sample()
apply_rows(ws2b, hr2b, cm2b, after2)
a2 = snap(after2)
ok(abs(a2["853-800575-413"]["price"] - 9100) < 0.01,
   "빈 칸이 판가를 %r 로 지웠다" % a2["853-800575-413"]["price"])
ok(abs(a2["853-800575-611"]["material_cost"] - 4383.18) < 0.01,
   "빈 칸이 재료비를 %r 로 지웠다" % a2["853-800575-611"]["material_cost"])
ok(a2["853-800575-611"]["group"] == "개발",
   "빈 칸이 구분을 %r 로 바꿨다 (개발이어야 한다)" % a2["853-800575-611"]["group"])
ok(a2["853-800575-413"]["po_qty"] == 999,
   "고친 PO 가 안 들어갔다: %r" % a2["853-800575-413"]["po_qty"])
ok(a2["853-800575-301"]["po_qty"] == 0,
   "안 고친 PO 가 바뀌었다: %r" % a2["853-800575-301"]["po_qty"])

# ── 3. 0 을 적으면 0 이 된다 ────────────────────────────────────────
m3 = sample()
m3[0]["price"] = "0"          # 사람이 0 을 적었다
ws3b = filled_sheet(m3)
h, c = find_columns(ws3b)
after3 = sample()
apply_rows(ws3b, h, c, after3)
ok(snap(after3)["853-800575-413"]["price"] == 0,
   "0 을 적었는데 판가가 %r 이다" % snap(after3)["853-800575-413"]["price"])

# ── 4. 머리글 변형도 받는다 ─────────────────────────────────────────
for hdr, field in [("개발유형", "dev_type"), ("개발 유형", "dev_type"),
                   (" 유형 ", "dev_type"), ("Type", "dev_type"),
                   ("재료비", "material_cost"), ("재료비($)", "material_cost"),
                   ("판가", "price"), ("판 가", "price"),
                   ("품번", "part_number"), ("Part No.", "part_number")]:
    wb4 = load_workbook(BytesIO(mt.build_model_template("", [])))
    ws4 = wb4.active
    _, c4 = find_columns(ws4)
    col = {"dev_type": 4, "material_cost": 6, "price": 5, "part_number": 1}[field]
    ws4.cell(2, col).value = hdr
    _, c4b = find_columns(ws4)
    ok(c4b.get(field) == col, "머리글 %r 을 %s 로 못 읽었다 (→ %r)" % (hdr, field, c4b.get(field)))

# ── 5. 재료비율 머리글은 절대 재료비가 되면 안 된다 ─────────────────
for hdr in ("재료비율", "재료 비율", "재료비율(%)", "비율"):
    wb5 = load_workbook(BytesIO(mt.build_model_template("", [])))
    ws5 = wb5.active
    ws5.cell(2, 7).value = hdr
    _, c5 = find_columns(ws5)
    ok(c5.get("material_cost") == 6,
       "머리글 %r 이 재료비 자리를 빼앗았다 (→ %r열)" % (hdr, c5.get("material_cost")))

# ── 6. 양식과 임포터가 같은 열을 말하는가 ───────────────────────────
#
# 양식의 머리글을 바꾸면 임포터가 못 읽는다. 둘이 한 몸이라는 걸
# 여기서 못 박아 둔다 — 한쪽만 고치면 이 검사가 터진다.
wsy = load_workbook(BytesIO(mt.build_model_template("엔클로저", ["413"]))).active
ok([wsy.cell(2, i).value for i in range(1, 11)] ==
   ["파트넘버", "모델명", "구분", "유형", "판가($)", "재료비($)",
    "재료비율", "PO수량", "실적수량", "비고"], "양식 머리글이 바뀌었다")
_, cy = find_columns(wsy)
ok(cy == {"part_number": 1, "name": 2, "group": 3, "dev_type": 4, "price": 5,
          "material_cost": 6, "po_qty": 8, "shipped_qty": 9, "note": 10},
   "양식 열과 임포터가 읽는 열이 어긋난다: %s" % sorted(cy.items()))
ok(str(wsy.cell(3, 7).value).startswith("="), "재료비율이 수식이 아니다")
ok(wsy.cell(3, 2).value in (None, ""), "빈 양식에 값이 들어 있다")

if FAIL:
    print("실패 %d건" % len(FAIL))
    for f in FAIL:
        print("  -", f)
    raise SystemExit(1)
print("전부 통과 · 검사 6묶음")

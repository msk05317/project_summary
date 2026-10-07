# -*- coding: utf-8 -*-
"""작성용 양식 — 한 해 한 장, 고른 달만 주차가 열린다.

보고용(plan_export)과 모양이 비슷해 보여도 하는 일이 반대다. 보고용은
붙이는 것이고 이건 **다시 올라오는** 것이다. 그래서 여기서 지키는 건
거의 다 '올렸을 때 제대로 읽히는가' 다.

  · 시트 이름이 그 달이다 — 올리는 쪽이 이걸로 달을 정한다
  · 고른 달에만 W 라벨이 있다 — 지난 달에 적어도 안 읽히니 칸을 안 준다
  · 중간에 '합계' 줄이 없다 — 읽는 쪽이 거기서 표가 끝난 줄 안다
  · 0 은 빈 칸이다 — '아직 안 정함' 과 '정말 0' 이 구분돼야 한다
  · 받아서 그대로 다시 읽으면 숫자가 그대로다
"""
import sys
from io import BytesIO
from pathlib import Path

BACK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACK))

from openpyxl import load_workbook          # noqa: E402

import pbx_plan_import as pbx               # noqa: E402
import plan_form as pf                      # noqa: E402
import week_calendar as wcal                # noqa: E402
import weekly_template as wt                # noqa: E402

FAIL = []


def ok(c, m):
    if not c:
        FAIL.append(m)


YEAR = 2026
MON = "2026-10"
W10 = wcal.get_month_weeks(MON)
W9 = wcal.get_month_weeks("2026-09")
W8 = wcal.get_month_weeks("2026-08")

MODELS = [
    {"id": "A-413", "name": "A-413", "dev_type": "413", "group": "양산",
     "price": 2931.15, "material_cost": 1204, "po_qty": 3226, "shipped_qty": 2670,
     "weekly_plan": {
         MON: {W10[0]: {"plan": 82, "actual": 0},
               W10[2]: {"plan": 123, "actual": 5}},
         "2026-09": {W9[0]: {"plan": 200, "actual": 180},
                     W9[1]: {"plan": 238, "actual": 265}},
         "2026-08": {W8[0]: {"plan": 0, "actual": 0}},
     }},
    {"id": "B-009", "name": "B-009", "dev_type": "009", "group": "양산",
     "po_qty": 1117, "shipped_qty": 1117},
    {"id": "C-301", "name": "C-301", "dev_type": "009", "group": "개발"},
]

data = pf.build_year_form("엔클로저", YEAR, MON, MODELS)
wb = load_workbook(BytesIO(data))


# ── 1. 시트 ─────────────────────────────────────────────────────────
ok(MON in wb.sheetnames, "시트 이름이 그 달이 아니다: %s" % wb.sheetnames)
ok("작성 안내" in wb.sheetnames, "작성 안내 시트가 없다")
ws = wb[MON]


# ── 2. 열두 달 · 고른 달만 주차 ─────────────────────────────────────
tops, weeks_seen = {}, []
for c in range(1, ws.max_column + 1):
    v = str(ws.cell(pf.HDR_TOP, c).value or "")
    if "월" in v and v.split("월")[0].strip().isdigit():
        tops[int(v.split("월")[0].strip())] = c
    w = str(ws.cell(pf.HDR_WEEK, c).value or "").upper()
    if w.startswith("W") and w[1:].isdigit():
        weeks_seen.append(w)
ok(sorted(tops) == list(range(1, 13)),
   "열두 달이 다 있어야 한다: %s" % sorted(tops))
ok(weeks_seen == list(W10),
   "고른 달의 주차만 있어야 한다\n  나온 것 %s\n  있어야 할 것 %s"
   % (weeks_seen, list(W10)))

# 월 합계는 열두 달 모두 있다
n_tot = sum(1 for c in range(1, ws.max_column + 1)
            if str(ws.cell(pf.HDR_WEEK, c).value or "") == "월 합계")
ok(n_tot == 12, "월 합계가 달마다 하나씩 있어야 한다: %d개" % n_tot)


# ── 3. 줄 ───────────────────────────────────────────────────────────
rows = {}
r_sum = None
for r in range(pf.FIRST_ROW, ws.max_row + 1):
    nm = str(ws.cell(r, 1).value or "").strip()
    if nm == "합계":
        r_sum = r
        break
    if nm:
        rows[nm] = r
ok(list(rows) == ["A-413", "B-009", "C-301"] or set(rows) ==
   {"A-413", "B-009", "C-301"}, "모델 줄이 다 안 나왔다: %s" % list(rows))
ok(r_sum is not None, "합계 줄이 없다")

# 중간에 묶음 합계가 있으면 읽는 쪽이 거기서 끊긴다
labels = [str(ws.cell(r, 1).value or "").strip()
          for r in range(pf.FIRST_ROW, ws.max_row + 1)]
ok(labels.count("합계") == 1,
   "'합계' 줄이 하나여야 한다 (묶음 합계를 두면 읽기가 거기서 끊긴다): %d개"
   % labels.count("합계"))


# ── 4. 값 ───────────────────────────────────────────────────────────
def mcol(mon_i, off=0):
    """그 달 월 합계 열 (+off)"""
    c = tops[mon_i]
    while str(ws.cell(pf.HDR_WEEK, c).value or "") != "월 합계":
        c += 1
    return c + off


r413 = rows["A-413"]
c10 = tops[10]
ok(ws.cell(r413, c10).value == 82, "이번 달 계획이 안 들어갔다: %r"
   % ws.cell(r413, c10).value)
ok(ws.cell(r413, c10 + 1).value is None,
   "0 은 빈 칸이어야 한다 ('아직 안 정함' 과 구분): %r" % ws.cell(r413, c10 + 1).value)
ok(ws.cell(r413, c10 + 2).value is None, "안 적은 주차에 숫자가 생겼다")

# 지난 달 — 월 합계만, 값은 주차를 더한 것
ok(ws.cell(r413, mcol(9)).value == 438,
   "지난 달 월 계획이 틀리다: %r" % ws.cell(r413, mcol(9)).value)
ok(ws.cell(r413, mcol(9, 1)).value == 445,
   "지난 달 월 실적이 틀리다: %r" % ws.cell(r413, mcol(9, 1)).value)
# 전부 0 인 달은 '-' 다 — 0 으로 적으면 '그 달에 하나도 안 나갔다' 가 된다
ok(ws.cell(r413, mcol(8)).value == pf.DASH,
   "주차가 전부 0 인 달은 '-' 여야 한다: %r" % ws.cell(r413, mcol(8)).value)
r009 = rows["B-009"]
ok(ws.cell(r009, mcol(9)).value == pf.DASH,
   "자료 없는 달에 0 이 생겼다: %r" % ws.cell(r009, mcol(9)).value)

# 판가·재료비·PO·실적은 적는 칸이라 값이 그대로 있어야 한다
heads = {str(ws.cell(pf.HDR_TOP, c).value or ""): c for c in range(1, 12)}
for h, want in (("판가($)", 2931.15), ("PO수량", 3226), ("실적수량", 2670)):
    ok(ws.cell(r413, heads[h]).value == want,
       "%s 가 안 채워졌다: %r" % (h, ws.cell(r413, heads[h]).value))


# ── 5. 수식이 '-' 를 만나도 안 깨진다 ───────────────────────────────
#
# 잔량이 =PO-실적 이면 한쪽이 비어 있을 때 #VALUE! 가 난다.
f = ws.cell(rows["C-301"], heads["잔량"]).value
ok(isinstance(f, str) and f.startswith("=") and "SUM(" in f,
   "잔량이 SUM 으로 안 빠진다 — 빈 칸에서 #VALUE! 가 난다: %r" % f)
fs = ws.cell(r_sum, mcol(10)).value
ok(isinstance(fs, str) and "COUNT(" in fs,
   "합계가 COUNT 로 안 감싸졌다: %r" % fs)


# ── 6. 되읽기 — 올렸을 때 그대로 들어가는가 ─────────────────────────
got = pbx.parse_sheet(ws)
ok(got is not None, "올리는 쪽이 이 시트를 못 읽는다")
if got:
    ok(got["month"] == MON, "달을 잘못 읽는다: %s" % got["month"])
    ok(got["weeks"] == list(W10), "주차를 잘못 읽는다: %s" % got["weeks"])
    by = {r["label"]: r for r in got["rows"]}
    ok("A-413" in by, "숫자가 있는 줄을 못 읽었다: %s" % list(by))
    if "A-413" in by:
        w = by["A-413"]["weeks"]
        ok((w[W10[0]]["plan"], w[W10[0]]["actual"]) == (82, 0),
           "되읽은 값이 다르다: %r" % w[W10[0]])
        ok((w[W10[2]]["plan"], w[W10[2]]["actual"]) == (123, 5),
           "되읽은 값이 다르다: %r" % w[W10[2]])
    # 숫자를 하나도 안 적은 줄은 안 읽혀야 한다 — 읽히면 그 달이 0 으로 덮인다
    ok("B-009" not in by,
       "빈 줄이 읽혔다 — 올리면 그 달이 통째로 0 이 된다")


# ── 7. 미달 사유 ────────────────────────────────────────────────────
d2 = pf.build_year_form("엔클로저", YEAR, MON, MODELS,
                        notes={W10[0]: "자재 입고 지연"})
ws2 = load_workbook(BytesIO(d2))[MON]
g2 = pbx.parse_sheet(ws2)
ok((g2 or {}).get("week_notes", {}).get(W10[0]) == "자재 입고 지연",
   "미달 사유가 되읽히지 않는다: %r" % (g2 or {}).get("week_notes"))


# ── 8. 라우트 · 화면 ────────────────────────────────────────────────
SRC = (BACK / "main.py").read_text(encoding="utf-8")
ok('@app.get("/admin/projects/{project_key}/plan-form")' in SRC, "라우트가 없다")
_b = SRC[SRC.index("def admin_plan_form("):]
_b = _b[:_b.index("\n@app.")] if "\n@app." in _b else _b
ok("get_admin_session" in _b, "라우트에 관리자 확인이 없다")
ok("build_year_form" in _b, "라우트가 작성용을 안 만든다")

AV2 = (BACK / "admin_v2.html").read_text(encoding="utf-8")
ok("'/plan-form'" in AV2, "화면이 작성용을 안 부른다")
ok("_dlItem('form'" in AV2 and "_dlItem('report'" in AV2,
   "받기 메뉴에 작성용·보고용이 다 있어야 한다")
ok("wp-year" not in AV2, "없어진 연간 버튼을 아직 찾고 있다")


# ── 9. 안 되는 달 ───────────────────────────────────────────────────
for bad, why in (("2025-10", "다른 해를 골랐는데 안 막는다"),
                 ("2026-13", "없는 달을 안 막는다")):
    try:
        pf.build_year_form("x", YEAR, bad, MODELS)
        FAIL.append(why)
    except Exception:
        pass


if FAIL:
    print("실패 %d건" % len(FAIL))
    for x in FAIL:
        print("  -", x)
    raise SystemExit(1)
print("전부 통과 · 검사 9묶음 · 달 12개 · 열린 주차 %d개" % len(W10))

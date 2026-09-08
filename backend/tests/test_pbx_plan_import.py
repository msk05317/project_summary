"""출하계획 엑셀(PBX 'Plan and Actual') 파서 회귀 테스트.

이 파서는 남이 만든 파일을 읽는다. 열 위치도, 시트 이름도, 수식도 바뀔 수 있다.
그래서 '무엇을 믿지 않기로 했는지'를 여기에 못 박아 둔다.

    cd backend && python3 tests/test_pbx_plan_import.py
"""
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

from openpyxl import Workbook          # noqa: E402

import pbx_plan_import as pbx          # noqa: E402


# ── 실제 파일과 같은 모양의 시트를 만든다 ────────────────────────
def _sheet(wb, title, weeks, rows, total_row=True, bad_total=False):
    ws = wb.create_sheet(title)
    first = 6                                   # F열부터 주차
    ws.cell(4, 5, "Model")                      # 사실은 5행이 머리글이지만
    ws.cell(5, 3, "Customer")
    ws.cell(5, 5, "Model")
    for i, w in enumerate(weeks):
        c = first + i * 2
        ws.cell(4, c, w)
        ws.cell(5, c, "Plan")
        ws.cell(5, c + 1, "Actual")
    tot_c = first + len(weeks) * 2
    ws.cell(4, tot_c, "TOTAL")
    ws.cell(5, tot_c, "Plan")
    ws.cell(5, tot_c + 1, "Actual")

    r = 6
    for label, vals in rows:
        ws.cell(r, 5, label)
        for i, (p, a) in enumerate(vals):
            ws.cell(r, first + i * 2, p)
            ws.cell(r, first + i * 2 + 1, a)
        # 실제 파일처럼 TOTAL 열이 깨져 있을 수 있다
        ws.cell(r, tot_c, 0 if bad_total else sum(p for p, _ in vals))
        r += 1
    if total_row:
        ws.cell(r, 2, "Total")
        r += 1
        ws.cell(r, first, 999)                  # 합계행 아래 찌꺼기
    return ws


def _wb(weeks=("W36", "W37", "W38", "W39", "W40"), **kw):
    wb = Workbook()
    wb.remove(wb.active)
    rows = [
        ("AETHER GDX 925-800083-394", [(40, 41), (50, 0), (50, 0), (50, 0), (50, 0)]),
        ("KIYO GX 575-B68653-XXXX", [(2, 2), (7, 0), (7, 0), (5, 0), (4, 0)]),
    ]
    rows = [(lb, v[:len(weeks)]) for lb, v in rows]
    _sheet(wb, "HVM_Plan and Actual_Sep", list(weeks), rows, **kw)
    return wb


MODELS = [
    {"id": "925-800083-394", "name": "AETHER GDX", "group": "양산"},
    # 품번 칸에 이름이 섞여 등록된 실제 사례
    {"id": "KIYO GX 575-B68653-XXXX", "name": "KIYO GX", "group": "양산"},
    {"id": "853-290073-014", "name": "VXT-AHM", "group": "양산"},
    {"id": "853-151282-200", "name": "VXT-AHM", "group": "개발"},
    {"id": "SUPREMA AC 3007952", "name": "SUPREMA AC 3007952", "group": "양산"},
]


def test_월을_주차번호로_정한다():
    """시트 이름('..._Sep')이 아니라 주차 번호로 달을 판단해야 한다."""
    g = pbx.parse(_wb())
    assert g["month"] == "2026-09", g["month"]
    assert g["weeks"] == ["W36", "W37", "W38", "W39", "W40"], g["weeks"]


def test_합계행에서_멈춘다():
    """'Total' 아래의 찌꺼기 행을 모델로 읽으면 합계가 부풀어 오른다."""
    g = pbx.parse(_wb())
    assert len(g["rows"]) == 2, [r["label"] for r in g["rows"]]


def test_남의_달_주차는_버린다():
    """W41 은 2026년 10월 것이다. 9월 시트에 섞여 있어도 담지 않는다."""
    g = pbx.parse(_wb(weeks=("W36", "W37", "W41")))
    assert "W41" not in g["weeks"], g["weeks"]


def test_TOTAL열이_깨져도_주차합은_맞는다():
    """실제 파일에서 이 열 수식이 깨져 계획이 400 대신 333 으로 나왔다."""
    g = pbx.parse(_wb(bad_total=True))
    plan = sum(w["plan"] for r in g["rows"] for w in r["weeks"].values())
    assert plan == 240 + 25, plan


def test_이름과_품번이_섞인_라벨을_맞춘다():
    m, how = pbx.Matcher(MODELS).match("AETHER GDX 925-800083-394")
    assert m and m["id"] == "925-800083-394", (m, how)


def test_계열품번이_실제_출하품번과_맞는다():
    """등록은 575-B68653-XXXX, 출하는 575-B68653-R7031 로 나간다."""
    m, how = pbx.Matcher(MODELS).match("575-B68653-R7031")
    assert m and m["name"] == "KIYO GX", (m, how)


def test_숫자만_있는_품번도_맞춘다():
    m, _ = pbx.Matcher(MODELS).match("SUPREMA AC 3007952")
    assert m and m["name"] == "SUPREMA AC 3007952", m


def test_이름이_겹치면_구분으로_가른다():
    M = pbx.Matcher(MODELS)
    assert M.match("VXT-AHM", "개발")[0]["id"] == "853-151282-200"
    assert M.match("VXT-AHM", "양산")[0]["id"] == "853-290073-014"
    # 구분을 모르면 임의로 고르지 말고 미매칭으로 남겨야 한다
    m, why = M.match("VXT-AHM")
    assert m is None and "중복" in why, (m, why)


def test_등록_안된_품번은_조용히_넘어가지_않는다():
    m, why = pbx.Matcher(MODELS).match("853-800083-326")
    assert m is None and why == "등록 안 됨", (m, why)


if __name__ == "__main__":
    fails = 0
    for name, fn in sorted(globals().items()):
        if not name.startswith("test_"):
            continue
        try:
            fn()
            print(f"  PASS  {name}")
        except AssertionError as e:
            fails += 1
            print(f"  FAIL  {name}\n        {e}")
    print(f"\n{'실패 ' + str(fails) + '건' if fails else '전부 통과'}")
    sys.exit(1 if fails else 0)

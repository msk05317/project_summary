# -*- coding: utf-8 -*-
"""블룸 양식 ↔ 파서 — 둘이 같은 모양을 보고 있는지.

양식(bloom_template)과 파서(bloom_daily_import)는 머리글 글자로 맞물린다.
'공정⏎구분' 을 '공정' 으로 바꾸면 양식은 멀쩡해 보이는데 파서가 열을 못 찾아
품목이 0개가 되고, 화면에는 "시트를 찾지 못했습니다" 로 뜬다. 실제로 그렇게
한 번 터졌다. 그래서 여기서 둘을 묶는다 — 양식을 만들어서 파서로 읽는다.

검사:
  1. 양식이 만들어지고 모양이 기준 파일과 같다 (머리글 네 줄, 고정 H5)
  2. 파서가 빈 양식의 열을 다 찾는다
  3. 사람이 채운 것처럼 숫자를 넣으면 그대로 읽힌다
  4. 터졌던 두 가지가 다시 안 터진다
     - '공정구분' 을 공정 열로 읽는다
     - '공정재고' 를 공정 열로 잘못 읽지 않는다
  5. 한 파일에 두 달이 있으면 늦은 달을 읽는다
"""
import datetime as dt
import sys
from io import BytesIO
from pathlib import Path

BACK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACK))

from openpyxl import Workbook, load_workbook   # noqa: E402

import bloom_daily_import as bd                # noqa: E402
from bloom_template import build_bloom_template, month_days   # noqa: E402

FAIL = []
MONTH = "2026-10"
ITEMS = ["YFP", "Corva KPE", "SL7", "SS8", "TC", "WDM", "BOP 기구", "BOP Assy"]


def ok(c, m):
    if not c:
        FAIL.append(m)


def load(blob):
    return load_workbook(BytesIO(blob))


# ── 1. 양식의 모양 ───────────────────────────────────────────────────
blob = build_bloom_template(MONTH, ITEMS)
ok(len(blob) > 5000, "양식이 너무 작다: %d bytes" % len(blob))
wb = load(blob)
ws = wb["10월 보고자료"]
ok("작성 안내" in wb.sheetnames, "'작성 안내' 시트가 없다 — 처음 받는 사람이 막힌다")
ok(ws.freeze_panes == "H5", "고정 틀이 H5 가 아니다: %r" % ws.freeze_panes)

g = bd._grid(ws)


def flat(r, c):
    return bd._s(g.get((r, c))).replace("\n", "").replace(" ", "")


for (r, c, want) in [
    (2, 1, "품목"), (2, 2, "재고"), (3, 2, "출하대기"), (3, 3, "구매품대기"),
    (2, 4, "공정구분"), (2, 5, "공정재고"), (2, 6, "10월합계"),
    (4, 6, "계획"), (4, 7, "실적"), (4, 8, "계획"), (4, 9, "실적"),
]:
    ok(flat(r, c) == want,
       "머리글 %d행%d열이 %r 이어야 하는데 %r — 파서가 이 글자로 열을 찾는다"
       % (r, c, want, flat(r, c)))

days = month_days(MONTH)
ok(len(days) == 31, "10월이 31일이 아니다: %d" % len(days))
ok(bd._as_date(g.get((3, 8))) == dt.date(2026, 10, 1), "첫 날짜가 10/1 이 아니다")
ok(bd._as_date(g.get((3, 68))) == dt.date(2026, 10, 31), "끝 날짜가 10/31 이 아니다")


# ── 2. 파서가 빈 양식을 읽는가 ───────────────────────────────────────
ok(bd.find_sheet(wb) is ws, "find_sheet 가 '10월 보고자료' 를 못 집었다")
lay = bd.find_layout(ws, g)
ok(lay is not None, "find_layout 이 빈손이다 — 양식과 파서가 어긋났다")
if lay:
    sub, cols, pairs = lay
    ok(sub == 4, "하위 머리글 줄이 4행이 아니다: %r" % sub)
    for key, want in (("item", 1), ("wait_ship", 2), ("wait_part", 3),
                      ("step", 4), ("wip", 5),
                      ("total_plan", 6), ("total_actual", 7)):
        ok(cols.get(key) == want,
           "%s 열이 %d 이어야 하는데 %r" % (key, want, cols.get(key)))
    ok("note" in cols, "비고 열을 못 찾았다")
    ok(len(pairs) == 31, "날짜쌍이 31개가 아니다: %d" % len(pairs))
    ok(pairs[0][0] == dt.date(2026, 10, 1), "첫 날짜쌍이 10/1 이 아니다")
    # '공정재고' 가 공정 열로 새지 않았는지 — 이게 샜으면 step 이 5열이 된다
    ok(cols.get("step") != cols.get("wip"), "공정 열과 공정재고 열이 같다")

got = bd.parse_daily(wb)
ok([i["item"] for i in got["items"]] == ITEMS,
   "빈 양식에서 품목이 그대로 안 나온다: %s" % [i["item"] for i in got["items"]])
for it in got["items"]:
    steps = [s["step"] for s in it["steps"]]
    ok(steps == ["NCT", "조립", "출하"],
       "%s 의 공정이 NCT·조립·출하 가 아니다: %s" % (it["item"], steps))


# ── 3. 사람이 채운 것처럼 넣으면 그대로 읽히는가 ──────────────────────
wb2 = load(blob)
ws2 = wb2["10월 보고자료"]
# YFP(5~7행) · 출하(7행) 에 10/1 계획 120 / 실적 100, 10월 합계 3000/2800
ws2.cell(5, 2, 40)           # 출하대기
ws2.cell(5, 3, 7)            # 구매품대기
ws2.cell(7, 5, 12)           # 공정재고
ws2.cell(7, 6, 3000)         # 10월 계획
ws2.cell(7, 7, 2800)         # 10월 실적
ws2.cell(7, 8, 120)          # 10/1 계획
ws2.cell(7, 9, 100)          # 10/1 실적
ws2.cell(7, 10, 130)         # 10/2 계획
ws2.cell(7, 70, "부자재 10/5 입고")   # 비고 — 그 공정 줄에 붙는 메모
# 표 아래 ◆ 줄 — 화면의 '확인이 필요한 것' 으로 올라가는 자리
ws2.cell(ws2.max_row + 2, 1, "◆ YFP 부자재 부족 ETA 10/5")
buf = BytesIO()
wb2.save(buf)
r2 = bd.parse_daily(load(buf.getvalue()))
y = [i for i in r2["items"] if i["item"] == "YFP"]
ok(len(y) == 1, "채운 파일에서 YFP 를 못 찾았다")
if y:
    y = y[0]
    ok(y["wait_ship"] == 40, "출하대기 40 이 아니다: %r" % y["wait_ship"])
    ok(y["wait_part"] == 7, "구매품대기 7 이 아니다: %r" % y["wait_part"])
    sh = [s for s in y["steps"] if s["step"] == "출하"]
    ok(len(sh) == 1, "YFP 에 출하 줄이 없다")
    if sh:
        sh = sh[0]
        ok(sh["wip"] == 12, "공정재고 12 가 아니다: %r" % sh["wip"])
        ok(sh["month_plan"] == 3000, "월 계획 3000 이 아니다: %r" % sh["month_plan"])
        ok(sh["month_actual"] == 2800, "월 실적 2800 이 아니다: %r" % sh["month_actual"])
        d1 = sh["days"].get("2026-10-01") or {}
        ok(d1.get("plan") == 120, "10/1 계획 120 이 아니다: %r" % d1.get("plan"))
        ok(d1.get("actual") == 100, "10/1 실적 100 이 아니다: %r" % d1.get("actual"))
        d2 = sh["days"].get("2026-10-02") or {}
        ok(d2.get("plan") == 130, "10/2 계획 130 이 아니다: %r" % d2.get("plan"))
        ok("부자재" in (sh.get("note") or ""),
           "비고 칸이 그 공정 줄의 메모로 안 들어갔다: %r" % sh.get("note"))
# 표 아래 '◆ ...' 줄 → 화면의 '확인이 필요한 것'
nt = r2.get("notes") or []
ok(any("부자재 부족" in (n.get("text") or "") for n in nt),
   "표 아래 ◆ 줄이 '확인이 필요한 것' 으로 안 올라간다: %r" % (nt,))
ok(any(n.get("eta") == "10/5" for n in nt),
   "ETA 10/5 를 못 잡았다: %r" % (nt,))


# ── 4. 터졌던 것 — 머리글 글자 ───────────────────────────────────────
SRC = (BACK / "bloom_daily_import.py").read_text(encoding="utf-8")
ok('flat.startswith("공정")' in SRC and '"재고" not in flat' in SRC,
   "파서가 '공정구분' 을 다시 못 읽는다 — 품목이 0개가 되는 그 버그다")
ok("_sheet_last_date" in SRC and "def find_sheet" in SRC,
   "find_sheet 가 가장 늦은 달을 고르지 않는다")

# '공정' 머리글만 있는 옛 파일도 그대로 읽혀야 한다 (뒤로 호환)
wb3 = load(blob)
ws3 = wb3["10월 보고자료"]
ws3.unmerge_cells(start_row=2, start_column=4, end_row=4, end_column=4)
ws3.cell(2, 4, "공정")
buf = BytesIO()
wb3.save(buf)
l3 = load(buf.getvalue())
s3 = bd.find_sheet(l3)
lay3 = bd.find_layout(s3, bd._grid(s3))
ok(lay3 and lay3[1].get("step") == 4,
   "머리글이 '공정' 뿐인 옛 파일을 못 읽는다 — 뒤로 호환이 깨졌다")


# ── 5. 두 달이 든 파일 — 늦은 달을 읽는다 ────────────────────────────
sep = build_bloom_template("2026-09", ITEMS)
two = load(sep)
two.remove(two["작성 안내"])
oct_ws = load(blob)["10월 보고자료"]
# 9월 시트 뒤에 10월 시트를 붙인다 (값만 옮겨도 모양 판별에는 충분하다)
new = two.create_sheet("10월 보고자료")
for row in oct_ws.iter_rows():
    for cell in row:
        if cell.value is not None:
            new.cell(cell.row, cell.column, cell.value)
buf = BytesIO()
two.save(buf)
picked = bd.find_sheet(load(buf.getvalue()))
ok(picked is not None and picked.title == "10월 보고자료",
   "두 달이 든 파일에서 %r 를 골랐다 — 늦은 달이어야 한다"
   % (picked.title if picked is not None else None))


if FAIL:
    print("실패 %d건" % len(FAIL))
    for f in FAIL:
        print("  -", f)
    raise SystemExit(1)
print("전부 통과 · 검사 5묶음 · 품목 %d · 날짜 %d" % (len(ITEMS), len(days)))

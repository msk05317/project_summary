# -*- coding: utf-8 -*-
"""블룸 보드 — 달이 넘어가도 지난 달이 남아 있는가.

10월 1일에 10월 파일을 올리는 순간 9월 계획 대비 실적이 사라졌다. 파일에는
'9월 보고자료' 와 '10월 보고자료' 가 같이 들어 있는데 한 장만 읽었기
때문이다. 지금은 달을 전부 읽어 달별로 둔다.

저장 자리:
  projects.<key>.daily_board          ← 이번 달 (예전 그대로. 챗봇·매출·RAG 가 본다)
  projects.<key>.daily_board_months   ← 지난 달만

달을 가릴 때는 **날짜**를 쓴다. 9월 시트의 제목도 '...보고(10/01)' 이라
적혀 있어서, 제목으로 달을 잡으면 9월 자료가 10월 자리로 들어간다.
"""
import datetime as dt
import sys
from io import BytesIO
from pathlib import Path

BACK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACK))

from openpyxl import load_workbook                 # noqa: E402

import bloom_daily_import as bd                    # noqa: E402
from bloom_template import build_bloom_template    # noqa: E402

FAIL = []
ITEMS = ["YFP", "Corva KPE", "SL7"]
KEEP = 5          # main.py 의 _BLOOM_KEEP_MONTHS


def ok(c, m):
    if not c:
        FAIL.append(m)


MAIN = (BACK / "main.py").read_text(encoding="utf-8")


# ── 두 달이 든 파일 하나 만들기 ──────────────────────────────────────
def two_month_book():
    """'9월 보고자료' + '10월 보고자료'. 제목은 둘 다 10/02 로 적는다 —
    실제 파일이 그렇다. 제목을 믿으면 안 된다는 걸 여기서 묶어 둔다."""
    sep = load_workbook(BytesIO(build_bloom_template("2026-09", ITEMS)))
    octo = load_workbook(BytesIO(build_bloom_template("2026-10", ITEMS)))
    del sep["작성 안내"]
    s9 = sep["9월 보고자료"]
    s9.cell(1, 1, "블룸 계획 대비 실적 보고(10/02)")
    s9.cell(7, 6, 1000)        # YFP 출하 9월 계획
    s9.cell(7, 7, 980)         # YFP 출하 9월 실적
    s9.cell(7, 8, 50)          # 9/1 계획
    s9.cell(7, 9, 48)          # 9/1 실적

    o = octo["10월 보고자료"]
    o.cell(1, 1, "블룸 계획 대비 실적 보고(10/02)")
    o.cell(7, 6, 1200)         # YFP 출하 10월 계획
    o.cell(7, 7, 0)
    o.cell(7, 8, 60)

    new = sep.create_sheet("10월 보고자료")
    for row in o.iter_rows():
        for c in row:
            if c.value is not None:
                new.cell(c.row, c.column, c.value)
    # 병합도 같이 옮긴다. 품목 이름은 A5:A7 로 묶여 있어서, 병합을 빼먹으면
    # 출하 줄(7행)의 품목이 비어 파서가 그 줄을 통째로 버린다.
    for rng in list(o.merged_cells.ranges):
        new.merge_cells(str(rng))
    buf = BytesIO()
    sep.save(buf)
    return load_workbook(BytesIO(buf.getvalue()))


WB = two_month_book()


# ── 1. 시트를 전부 찾는가 ───────────────────────────────────────────
ss = bd.find_sheets(WB)
ok(len(ss) == 2, "시트를 %d개 찾았다 — 9월·10월 둘 다여야 한다" % len(ss))
ok([w.title for w, _ in ss] == ["9월 보고자료", "10월 보고자료"],
   "find_sheets 가 이른 달부터 주지 않는다: %s" % [w.title for w, _ in ss])
ok(ss and ss[-1][1] == dt.date(2026, 10, 31), "마지막 시트의 끝 날짜가 10/31 이 아니다")
ok(bd.find_sheet(WB).title == "10월 보고자료",
   "find_sheet 가 늦은 달을 안 고른다: %r" % bd.find_sheet(WB).title)


# ── 2. 달은 날짜로 가린다 (제목이 아니라) ───────────────────────────
def month_of(board):
    ds = (board or {}).get("dates") or []
    return str(ds[0])[:7] if ds else ""


parsed = {}
for ws, _ in ss:
    p = bd.parse_daily(WB, sheet_name=ws.title)
    parsed[month_of(p)] = p

ok(sorted(parsed) == ["2026-09", "2026-10"],
   "달을 잘못 가렸다: %s — 제목(10/02)을 보고 있는 것 아닌가" % sorted(parsed))
for mon, want_title in (("2026-09", "10/02"), ("2026-10", "10/02")):
    ok(want_title in (parsed[mon].get("title") or ""),
       "%s 시트 제목이 바뀌었다 — 이 테스트의 전제가 깨졌다" % mon)

# 9월 숫자가 9월에, 10월 숫자가 10월에
def ship(p, name="YFP"):
    for it in p["items"]:
        if it["item"] != name:
            continue
        for s in it["steps"]:
            if (s.get("group") or s["step"]) == "출하":
                return s.get("month_plan"), s.get("month_actual")
    return None, None


ok(ship(parsed["2026-09"]) == (1000, 980),
   "9월 출하가 1000/980 이 아니다: %r" % (ship(parsed["2026-09"]),))
ok(ship(parsed["2026-10"]) == (1200, 0),
   "10월 출하가 1200/0 이 아니다: %r" % (ship(parsed["2026-10"]),))
ok(parsed["2026-09"]["dates"][0] == "2026-09-01", "9월 보드의 첫 날짜가 9/1 이 아니다")
ok(parsed["2026-10"]["dates"][0] == "2026-10-01", "10월 보드의 첫 날짜가 10/1 이 아니다")


# ── 3. 어디에 담기는가 — 이번 달은 제자리, 지난 달만 옆에 ──────────
def store_months(proj):
    """main.py 의 _bloom_store_months 와 같은 규칙."""
    out = {}
    prev = proj.get("daily_board_months")
    if isinstance(prev, dict):
        for m, b in prev.items():
            if isinstance(b, dict) and b.get("items"):
                out[str(m)] = b
    cur = proj.get("daily_board")
    if isinstance(cur, dict) and cur.get("items"):
        m = month_of(cur)
        if m:
            out[m] = cur
    return out


def split(store):
    """저장할 때 나누는 규칙 — 최신 달은 daily_board, 나머지는 옆으로."""
    keep = sorted(store)[-(KEEP + 1):]
    store = {m: store[m] for m in keep}
    newest = keep[-1]
    return store[newest], {m: b for m, b in store.items() if m != newest}, newest


proj = {}
cur, older, newest = split({m: dict(p) for m, p in parsed.items()})
proj["daily_board"] = cur
proj["daily_board_months"] = older
ok(newest == "2026-10", "최신 달이 10월이 아니다: %r" % newest)
ok(month_of(proj["daily_board"]) == "2026-10",
   "daily_board 에 10월이 안 들어갔다 — 챗봇·매출이 지난 달을 보게 된다")
ok(list(proj["daily_board_months"]) == ["2026-09"],
   "지난 달 자리가 9월 하나가 아니다: %s" % list(proj["daily_board_months"]))
ok(sorted(store_months(proj)) == ["2026-09", "2026-10"],
   "두 자리를 합쳐도 달이 안 나온다: %s" % sorted(store_months(proj)))

# 한 달만 든 파일을 올려도 지난 달이 살아 있어야 한다
only10 = {"2026-10": dict(parsed["2026-10"])}
st = store_months(proj)
st.update(only10)
cur2, older2, _ = split(st)
ok(list(older2) == ["2026-09"],
   "10월만 든 파일을 올리니 9월이 사라졌다: %s" % list(older2))

# 너무 오래된 달은 버린다
many = {"2026-%02d" % m: {"items": [1], "dates": ["2026-%02d-01" % m]}
        for m in range(1, 13)}
_c, _o, _n = split(many)
ok(len(_o) + 1 == KEEP + 1, "달을 %d개 들고 있다 — %d개여야 한다" % (len(_o) + 1, KEEP + 1))
ok(_n == "2026-12", "가장 최근 달을 못 잡았다: %r" % _n)
ok(sorted(_o)[0] == "2026-07", "버리는 기준이 틀렸다: %s" % sorted(_o)[:2])


# ── 4. 라우트가 그 규칙을 쓰는가 ────────────────────────────────────
for frag, why in [
    ('def _bloom_parse_months(', "달별 읽기 함수가 없다"),
    ('def _bloom_apply_months(', "달별 합치기 함수가 없다"),
    ('def _bloom_store_months(', "달별 읽어오기 함수가 없다"),
    ('_BLOOM_KEEP_MONTHS = %d' % KEEP, "지난 달 보관 개수가 %d 가 아니다" % KEEP),
    ('def get_daily_board(project_key: str, date: str = "", month: str = "")',
     "보드 라우트에 month 가 없다 — 앱이 지난 달을 못 부른다"),
    ('out["months"] = months', "응답에 months 가 없다 — 앱이 어느 달이 있는지 모른다"),
    ('out["month"] = pick', "응답에 month 가 없다"),
    ('proj["daily_board"] = board', "이번 달이 daily_board 에 안 들어간다"),
    ('proj["daily_board_months"] = older', "지난 달이 따로 안 쌓인다"),
    ('return str(ds[0])[:7] if ds else ""',
     "달을 날짜로 안 가린다 — 제목을 보면 9월이 10월로 들어간다"),
]:
    ok(frag in MAIN, why)

ok(MAIN.count("_bloom_parse_upload(await file.read())") == 0,
   "아직 한 달만 읽는 옛 경로가 남아 있다")

# 품목 화면은 지금도 자기 몫만 본다
ok("board = _bloom_slice(board, _item)" in MAIN,
   "품목 프로젝트가 자기 몫만 보는 규칙이 사라졌다")



# ── 5. '월 합계' 가 없는 달도 월간 보기가 비지 않는가 ───────────────
#
# 파일마다 'N월 합계' 열이 있기도 하고 없기도 하다. 없는 달은 month_plan /
# month_actual 이 통째로 None 이라, 날짜별 계획·실적이 21일치 들어 있어도
# 월간 보기가 텅 비어 보였다. 응답을 만들 때만 날짜 칸을 더해 채운다.
def fill(board):
    """main.py 의 _bloom_with_month_totals 와 같은 규칙."""
    items = []
    for it in (board.get("items") or []):
        steps = []
        for st in (it.get("steps") or []):
            if st.get("month_plan") is None or st.get("month_actual") is None:
                st = dict(st)
                days = st.get("days") or {}
                for fld, key, prior in (("month_plan", "plan", "prior_plan"),
                                        ("month_actual", "actual", "prior_actual")):
                    if st.get(fld) is not None:
                        continue
                    tot, found = 0, False
                    base = st.get(prior)
                    if base is not None:
                        tot, found = int(base), True
                    for v in days.values():
                        x = (v or {}).get(key)
                        if x is None:
                            continue
                        tot += int(x)
                        found = True
                    if found:
                        st[fld] = tot
            steps.append(st)
        it = dict(it)
        it["steps"] = steps
        items.append(it)
    out = dict(board)
    out["items"] = items
    return out


raw = {"items": [{"item": "YFP", "steps": [
    # 월 합계가 비었고 날짜만 있는 줄
    {"step": "출하", "month_plan": None, "month_actual": None,
     "days": {"2026-09-01": {"plan": 10, "actual": 8},
              "2026-09-02": {"plan": 20, "actual": None},
              "2026-09-03": {"plan": None, "actual": 5}}},
    # 월 합계가 적힌 줄 — 손대지 않는다
    {"step": "조립", "month_plan": 999, "month_actual": 111,
     "days": {"2026-09-01": {"plan": 1, "actual": 1}}},
    # 날짜 칸 앞의 '이전 데이터' 도 더한다
    {"step": "NCT", "month_plan": None, "month_actual": None,
     "prior_plan": 100, "prior_actual": 90,
     "days": {"2026-09-01": {"plan": 5, "actual": 4}}},
    # 아무것도 없는 줄은 None 그대로 — 0 으로 만들면 '계획 0' 처럼 읽힌다
    {"step": "검사", "month_plan": None, "month_actual": None, "days": {}},
]}]}
got = {s["step"]: (s.get("month_plan"), s.get("month_actual"))
       for s in fill(raw)["items"][0]["steps"]}
ok(got["출하"] == (30, 13), "날짜 칸을 안 더한다: %r" % (got["출하"],))
ok(got["조립"] == (999, 111), "적혀 있는 월 합계를 덮어썼다: %r" % (got["조립"],))
ok(got["NCT"] == (105, 94), "'이전 데이터' 를 안 더한다: %r" % (got["NCT"],))
ok(got["검사"] == (None, None),
   "빈 줄을 0 으로 만들었다 — '계획 0' 과 '안 적힘' 은 다르다: %r" % (got["검사"],))

# 원본은 안 바뀐다 (models.json 에 없는 숫자를 적어 두면 안 된다)
ok(raw["items"][0]["steps"][0]["month_plan"] is None,
   "저장된 보드를 건드렸다 — 채우는 건 응답에서만이어야 한다")

for frag, why in [
    ("def _bloom_with_month_totals(", "월 합계 채우는 함수가 없다"),
    ("board = _bloom_with_month_totals(board)", "보드 응답에 안 쓰인다"),
]:
    ok(frag in MAIN, why)

# admin 화면에도 달 고르는 줄이 있는가
AV2 = (BACK / "admin_v2.html").read_text(encoding="utf-8")
for frag, why in [
    ("window._bdMonthBar =", "admin 에 달 고르는 줄이 없다"),
    ("'/daily-board' + _q", "admin 이 달을 지정해 부르지 않는다"),
    ("window._bdMonth = '';", "프로젝트를 바꿔도 달이 초기화되지 않는다"),
]:
    ok(frag in AV2, why)


if FAIL:
    print("실패 %d건" % len(FAIL))
    for f in FAIL:
        print("  -", f)
    raise SystemExit(1)
print("전부 통과 · 검사 5묶음 · 달 2개 · 보관 %d개월" % KEEP)

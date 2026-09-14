# CURIE 계획&실적 시트 읽기
#   python3 backend/tests/test_curie_plan_import.py
#
# 이 시트는 기존 것들과 두 군데가 다르다.
#   1) 위에 요약 블록이 있어서 '계획/실적' 줄이 13행이다 (12행까지만 보던 문제)
#   2) '모델명' 머리글이 세로로 병합돼 있다 ('모델명모델명모델명' 이 되던 문제)
import datetime as _dt
import pathlib, sys
import openpyxl
from openpyxl.utils import get_column_letter as _L

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import plan_matrix_import as pm

TODAY = _dt.date(2026, 9, 14)


def build():
    wb = openpyxl.Workbook()
    ws = wb.active
    # 위쪽 요약 블록 — 여기 때문에 표가 아래로 밀린다
    ws['B2'], ws['C2'], ws['E2'] = '큐리', 'PO수량', '출하수량'
    ws['B3'], ws['C3'], ws['E3'] = '키스톤', 2, 0
    ws['B8'], ws['C8'], ws['E8'] = '합산', 6531, 787
    ws['B10'] = '<CURIE 계획&실적>'

    ws['B11'], ws['C11'], ws['D11'], ws['E11'] = '모델명', 'PO 수량', '출하실적', 'PO 잔량'
    ws.merge_cells('B11:B13')          # 세로 병합된 머리글
    ws['F11'] = '7월 Total'            # 주차 열이 없는 달 — 합계가 유일한 값
    ws['H11'] = '8월'
    ws['P11'] = '8월 Total'            # 주차 열이 있는 달 — 버려야 한다
    ws['R11'] = '9월'
    ws['Z11'] = '9월 Total'
    ws.merge_cells('H11:O11')
    ws.merge_cells('R11:Y11')
    for col, w in (('H', 'W32'), ('J', 'W33'), ('L', 'W34'), ('N', 'W35'),
                   ('R', 'W36'), ('T', 'W37'), ('V', 'W38'), ('X', 'W39')):
        ws[f'{col}12'] = w
    for c in range(6, 27):             # F~Z 계획/실적 쌍
        ws.cell(13, c).value = '계획' if (c - 6) % 2 == 0 else '실적'

    # A열은 프로그램명 — 머리글이 없어서 라벨 열로 오인하기 쉽다
    rows = [
        ('스타쉽', '리프트 캐리지 (66C)', 24, 12, 12,
         {'F': 4, 'G': 4, 'H': 4, 'K': 4, 'M': 4, 'N': 4, 'V': 4}),
        ('게이트웨이', '도기도어 (76C)', 3200, 379, 2821,
         {'F': 247, 'G': 247, 'H': 70, 'I': 70, 'J': 62, 'K': 62}),
        ('게이트웨이', '버스바/시트메탈류(11종)', 1740, 0, 1740, {'R': 1380}),
    ]
    r = 14
    for prog, name, po, sh, rem, cells in rows:
        ws.cell(r, 1).value = prog
        ws.cell(r, 2).value = name
        ws.cell(r, 3).value = po
        ws.cell(r, 4).value = sh
        ws.cell(r, 5).value = rem
        for col, v in cells.items():
            ws[f'{col}{r}'] = v
        r += 1
    ws.cell(r, 2).value = 'Total'
    ws.cell(r, 3).value = 4964
    ws.cell(r, 4).value = 391
    return wb


out = pm.parse_plan_matrix(build(), today=TODAY)
rows = {r['label']: r for r in out['rows']}
ok = 0

assert out['sheet'], '표를 아예 못 찾았다 (계획/실적 줄이 13행)'
ok += 1
assert list(rows) == ['리프트 캐리지 (66C)', '도기도어 (76C)', '버스바/시트메탈류(11종)'], \
    f'라벨 열을 잘못 잡았다(A열 프로그램명?): {list(rows)}'
ok += 1

lift = rows['리프트 캐리지 (66C)']
assert (lift['po_qty'], lift['shipped_qty'], lift['remaining']) == (24, 12, 12), lift
ok += 1

w8 = lift['weeks']['2026-08']
assert w8['W32'] == {'plan': 4, 'actual': 0}, w8['W32']
assert w8['W33'] == {'plan': 0, 'actual': 4}, w8['W33']
assert w8['W34'] == {'plan': 0, 'actual': 4}, w8['W34']
assert w8['W35'] == {'plan': 4, 'actual': 0}, w8['W35']
ok += 1
assert lift['weeks']['2026-09']['W38'] == {'plan': 4, 'actual': 0}
ok += 1

# 주차 열이 있는 달의 'N월 Total' 은 버린다 (주차와 이중 계산 방지)
assert '2026-08' not in lift['months'] and '2026-09' not in lift['months'], lift['months']
ok += 1
# 주차 열이 없는 7월은 합계가 유일한 값이라 살린다
assert lift['months']['2026-07'] == {'plan': 4, 'actual': 4}, lift['months']
assert rows['도기도어 (76C)']['months']['2026-07'] == {'plan': 247, 'actual': 247}
ok += 1

# W40 처럼 앱의 월 소유 규칙과 엑셀 머리글이 어긋나는 주차는 앱 규칙이 정본
assert rows['버스바/시트메탈류(11종)']['weeks']['2026-09']['W36']['plan'] == 1380
ok += 1

assert rows['버스바/시트메탈류(11종)']['is_aggregate'] is True, '11종 묶음 행을 못 알아봤다'
assert rows['리프트 캐리지 (66C)']['is_aggregate'] is False
ok += 1

assert out['total'] and out['total']['po_qty'] == 4964, out['total']
ok += 1

print(f'전부 통과 ({ok}/10)')

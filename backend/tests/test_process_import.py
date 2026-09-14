# 개발 프로세스 일정표 — 계획·실적 열을 머리글로 찾는지 검증
#   python3 backend/tests/test_process_import.py
#
# 자리를 번호로 세면 'Model' 열이 있는 시트와 없는 시트가 한 칸씩 밀린다.
# 그러면 실적이 계획 칸으로 들어가고, 다 끝난 단계가 미시작으로 보인다.
import pathlib, sys
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
import process_import as pi

STEPS = ['01 FA PO', '02 Material Order', '03 Material Receiving',
         '04 CB', '05 BV1', '06 BV2']


def _sheet(with_model):
    wb = Workbook(); ws = wb.active
    head = ['No.', 'Part Number'] + (['Model'] if with_model else [])
    for i, h in enumerate(head, start=1):
        ws.cell(1, i, h)
    c = len(head) + 1
    for s in STEPS:
        ws.cell(1, c, s)
        ws.cell(2, c, 'Planned'); ws.cell(2, c + 1, 'Actual')
        c += 2
    row = ['1', '853-151282-200'] + (['VXT-AHM'] if with_model else [])
    # 01·02 는 '완료' 글자, 03 은 실적 날짜, 04 는 계획만
    row += ['', '완료', '', '완료', '', '2026-06-12', '2026-08-28', '', '', '', '', '']
    for i, v in enumerate(row, start=1):
        ws.cell(3, i, v)
    return ws


def _read(ws):
    head_row, pn_col, steps, _ = pi.find_layout(ws)
    assert head_row == 1 and len(steps) == 6, (head_row, len(steps))
    out = []
    for no, en, pc, ac in steps:
        p = ws.cell(3, pc).value or ''
        a = ws.cell(3, ac).value or ''
        out.append((no, en, str(p), str(a)))
    return pn_col, out


ok = 0

# 1. Model 열이 있는 시트 (PBX·MajorModule 모양)
pn_col, got = _read(_sheet(True))
assert pn_col == 2, pn_col
assert got[0] == (1, 'FA PO', '', '완료'), got[0]
assert got[2] == (3, 'Material Receiving', '', '2026-06-12'), got[2]
assert got[3] == (4, 'CB', '2026-08-28', ''), got[3]
ok += 1

# 2. Model 열이 없는 시트 (EMA 모양) — 같은 값이 같은 자리에 들어가야 한다
pn_col2, got2 = _read(_sheet(False))
assert pn_col2 == 2, pn_col2
assert got2 == got, list(zip(got, got2))
ok += 1

# 3. 'Model' 을 단계로 잘못 읽지 않는다
assert all(en != 'Model' for _, en, _, _ in got), got
ok += 1

# 4. 단계 수는 엑셀이 정한다 — 코드에 박아 두지 않는다
wb = Workbook(); ws = wb.active
ws.cell(1, 1, 'No.'); ws.cell(1, 2, 'Part Number')
for i in range(15):
    c = 3 + i * 2
    ws.cell(1, c, f'{i+1:02d} Step{i+1}')
    ws.cell(2, c, 'Planned'); ws.cell(2, c + 1, 'Actual')
ws.cell(3, 1, '1'); ws.cell(3, 2, 'PN-1')
assert len(pi.find_layout(ws)[2]) == 15, len(pi.find_layout(ws)[2])
ok += 1

# 5. 두 자리 단계 번호도 읽는다
assert pi.find_layout(ws)[2][-1][0] == 15
ok += 1

print(f'전부 통과 ({ok}/5)')

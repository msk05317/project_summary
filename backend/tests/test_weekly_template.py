# -*- coding: utf-8 -*-
"""주차별 계획·실적 양식 — 어느 프로젝트든 같은 모양이어야 한다.

엔클로저에서만 되면 쓸모가 없다. 모델이 0 종인 프로젝트, 80 종인
프로젝트, 품번이 아예 없는 프로젝트, 유형이 아예 없는 프로젝트가 다
섞여 있고 전부 같은 양식으로 나와야 한다.

가장 중요한 검사는 '주차가 보드와 같은가' 다. 2026-10 의 주차는
W41~W44 이고 W40 은 9월 소유다. 양식이 제 나름대로 주차를 세면 올려도
칸이 안 맞는다 — 그래서 week_calendar 한 군데서만 받는다.
"""
import sys
from io import BytesIO
from pathlib import Path

BACK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACK))

from openpyxl import load_workbook          # noqa: E402
import week_calendar as wcal                # noqa: E402
import weekly_template as wt                # noqa: E402

FAIL = []


def ok(cond, msg):
    if not cond:
        FAIL.append(msg)


def sheets_for(months):
    return [{"month": m, "weeks": wcal.get_month_weeks(m)} for m in months]


def head_cols(ws):
    """머리글 → 열 번호. 쓰지 않는 칸은 빠지므로 자리를 박아 두면 안 된다."""
    out = {}
    for c in range(1, 15):
        v = ws.cell(wt.HDR_TOP, c).value
        t = "" if v is None else str(v).strip()
        if not t or t == "월 합계" or t.startswith("W"):
            break
        out[t] = c
    return out


def n_fixed(ws):
    return len(head_cols(ws))


def names_in(ws):
    c = head_cols(ws).get("모델명")
    if not c:
        return []
    out = []
    for r in range(wt.FIRST_ROW, ws.max_row + 1):
        # 목록은 '합계' 줄에서 끝난다 (그 아래는 미달 사유 줄)
        head = str(ws.cell(r, 1).value or "").strip()
        if head in ("합계", "미달 사유"):
            break
        v = ws.cell(r, c).value
        if v not in (None, ""):
            out.append(str(v).strip())
    return out


# ── 프로젝트 모양들 — 실제로 있는 조합을 본떴다 ─────────────────────
def mk(n, dev_type=True, part=True, dev=0, prefix="M"):
    out = []
    for i in range(n):
        m = {"id": "%s-%04d" % (prefix, i), "name": "%s-%04d" % (prefix, i),
             "group": "개발" if i < dev else "양산"}
        if dev_type:
            m["dev_type"] = ["413", "009", "EMA"][i % 3]
        if part:
            m["part_number"] = "PN-%s-%04d" % (prefix, i)
        out.append(m)
    return out


SHAPES = {
    "빈 프로젝트": [],
    "한 종": mk(1, dev_type=False, part=False),
    "품번 없음 (큐리·하바 형)": mk(29, dev_type=True, part=False, dev=28),
    "유형 없음": mk(12, dev_type=False, part=True),
    "둘 다 없음": mk(8, dev_type=False, part=False),
    "큰 프로젝트 (파워박스 형)": mk(82, dev_type=True, part=True, dev=50),
    "이름만 id 에 있는 것": [{"id": "ONLY-ID-1", "group": "양산"},
                            {"id": "ONLY-ID-2", "group": "개발"}],
}

# ── 1. 모든 모양에서 만들어지고, 모델이 하나도 안 샌다 ──────────────
for label, ms in SHAPES.items():
    try:
        raw = wt.build_weekly_template(label, sheets=sheets_for(["2026-10"]), models=ms)
    except Exception as e:
        FAIL.append("%s · 양식이 안 만들어졌다: %s" % (label, e))
        continue
    ws = load_workbook(BytesIO(raw))["2026-10"]
    want = sorted(str(m.get("name") or m.get("id")).strip() for m in ms)
    got = sorted(names_in(ws))
    ok(got == want, "%s · 모델 %d종 중 %d종만 들어갔다" % (label, len(want), len(got)))
    ok(ws.max_row >= wt.FIRST_ROW + len(ms),
       "%s · 새 모델 적을 빈 줄이 없다" % label)

# ── 2. 주차가 보드와 같은 달력에서 나온다 ───────────────────────────
#
# 여기가 어긋나면 올려도 칸이 안 맞는다. 달이 바뀌는 경계를 특히 본다.
for mon in ("2026-01", "2026-02", "2026-09", "2026-10", "2026-12",
            "2027-01", "2027-03"):
    weeks = wcal.get_month_weeks(mon)
    raw = wt.build_weekly_template("T", sheets=[{"month": mon, "weeks": weeks}],
                                   models=mk(3))
    ws = load_workbook(BytesIO(raw))[mon]
    got = []
    c = n_fixed(ws) + 1
    while True:
        v = ws.cell(wt.HDR_TOP, c).value
        if v in (None, "", "월 합계"):
            break
        got.append(str(v).strip())
        c += wt.WEEK_N
    ok(got == weeks, "%s · 양식 주차 %s 가 달력 %s 와 다르다" % (mon, got, weeks))
    ok(ws.cell(wt.HDR_SUB, n_fixed(ws) + 1).value == "계획",
       "%s · 계획/실적 줄이 없다" % mon)

# 2026-10 은 W41 부터다 (W40 은 9월 소유). 이걸 못 박아 둔다.
ok(wcal.get_month_weeks("2026-10")[0] == "W41",
   "달력이 바뀌었다 — 2026-10 첫 주차가 %s" % wcal.get_month_weeks("2026-10")[0])
ok("W40" in wcal.get_month_weeks("2026-09"),
   "W40 이 9월 소유가 아니게 됐다")

# ── 3. 여러 달을 한 파일에 ──────────────────────────────────────────
mons = ["2026-10", "2026-11", "2026-12"]
raw = wt.build_weekly_template("T", sheets=sheets_for(mons), models=mk(5))
wb = load_workbook(BytesIO(raw))
ok(wb.sheetnames == mons + ["작성 안내"], "시트 구성이 다르다: %s" % wb.sheetnames)
for mon in mons:
    ok(len(names_in(wb[mon])) == 5, "%s 시트에 모델이 안 들어갔다" % mon)

# ── 4. 유형끼리 모이고, 유형 없는 것은 맨 아래 ──────────────────────
mix = [{"id": "C", "name": "C", "dev_type": "009"},
       {"id": "A", "name": "A", "dev_type": "413"},
       {"id": "Z", "name": "Z"},                      # 유형 없음
       {"id": "B", "name": "B", "dev_type": "009"},
       {"id": "D", "name": "D", "dev_type": "413"}]
ws = load_workbook(BytesIO(wt.build_weekly_template(
    "T", sheets=sheets_for(["2026-10"]), models=mix)))["2026-10"]
order = names_in(ws)
# 유형은 사전순이라 '009' 가 '413' 보다 앞이다. 유형 없는 것은 맨 아래.
ok(order == ["B", "C", "A", "D", "Z"], "정렬이 다르다: %s" % order)
ok(order[-1] == "Z", "유형 없는 모델이 맨 아래가 아니다")
ok(order[:2] == ["B", "C"], "같은 유형끼리 안 모였다: %s" % order)

# ── 5. 서로 구분 안 되는 줄을 짚어낸다 ──────────────────────────────
same_name = [{"id": "1", "name": "버스바"}, {"id": "2", "name": "버스바"},
             {"id": "3", "name": "시트메탈"}]
d = wt.duplicate_rows(same_name)
ok(len(d) == 1, "같은 이름 두 줄을 못 짚었다: %s" % d)

by_pn = [{"id": "1", "name": "버스바", "part_number": "P1"},
         {"id": "2", "name": "버스바", "part_number": "P2"}]
ok(wt.duplicate_rows(by_pn) == [], "파트넘버로 갈리는데 중복이라고 한다")

same_pn = [{"id": "1", "name": "가", "part_number": "P1"},
           {"id": "2", "name": "나", "part_number": "P1"}]
ok(len(wt.duplicate_rows(same_pn)) == 1, "같은 파트넘버 두 줄을 못 짚었다")

# ── 6. 모양이 흔들리지 않는다 ───────────────────────────────────────
ws = load_workbook(BytesIO(wt.build_weekly_template(
    "엔클로저", sheets=sheets_for(["2026-09"]), models=mk(4))))["2026-09"]
ok([ws.cell(wt.HDR_TOP, i).value for i in range(1, 5)] ==
   ["파트넘버", "모델명", "유형", "구분"], "왼쪽 네 칸이 바뀌었다")
from openpyxl.utils import get_column_letter as _L   # noqa: E402
_want_fz = "%s%d" % (_L(n_fixed(ws) + 1), wt.FIRST_ROW)
ok(ws.freeze_panes == _want_fz,
   "틀 고정이 %s 가 아니다: %s" % (_want_fz, ws.freeze_panes))
ok(ws.page_setup.orientation == "landscape", "가로 인쇄가 아니다")
ok(str(ws.print_title_rows).replace("$", "") == "2:3",
   "머리글 반복이 없다: %s" % ws.print_title_rows)
# 월 합계는 수식이어야 한다 (사람이 더하면 틀린다)
wk = wcal.get_month_weeks("2026-09")
tot_c = len(wt.FIXED) + len(wt.VALUE) + len(wk) * wt.WEEK_N + 1
ok(str(ws.cell(wt.FIRST_ROW, tot_c).value).startswith("="),
   "월 합계가 수식이 아니다: %r" % ws.cell(wt.FIRST_ROW, tot_c).value)
ok(ws.cell(wt.HDR_TOP, tot_c).value == "월 합계", "월 합계 칸이 없다")
# 숫자 칸은 비어 있어야 한다 — 양식이지 데이터가 아니다
for r in range(wt.FIRST_ROW, wt.FIRST_ROW + 4):
    for c in range(len(wt.FIXED) + 1, tot_c):
        ok(ws.cell(r, c).value is None,
           "숫자 칸에 값이 들어 있다 (r%d c%d)" % (r, c))

# ── 7. 달·주차가 없으면 거절한다 ────────────────────────────────────
for args, why in [({"sheets": []}, "달 없이"),
                  ({"sheets": [{"month": "2026-10", "weeks": []}]}, "주차 없이")]:
    try:
        wt.build_weekly_template("T", models=mk(2), **args)
        FAIL.append("%s 만들었는데 안 막혔다" % why)
    except ValueError:
        pass

# ── 8. 안 쓰는 칸은 빼고 내보낸다 ───────────────────────────────────
#
# 열 개 프로젝트로 뽑아 보니 파트넘버 칸이 거의 다 비어 있었다 (품번이
# 모델명 칸에 들어 있는 곳이 대부분이다). 빈 칸을 두면 사람들이 거기
# 뭔가 적고, 적힌 건 어느 모델에도 안 붙어 조용히 사라진다.
def cols_of(ms):
    ws = load_workbook(BytesIO(wt.build_weekly_template(
        "T", sheets=sheets_for(["2026-10"]), models=ms)))["2026-10"]
    return list(head_cols(ws))


V = ["판가($)", "재료비($)", "PO수량", "실적수량"]   # 값 칸은 늘 붙는다
ok(cols_of(mk(5, dev_type=True, part=True)) == ["파트넘버", "모델명", "유형", "구분"] + V,
   "둘 다 있는데 칸이 빠졌다: %s" % cols_of(mk(5, dev_type=True, part=True)))
ok(cols_of(mk(5, dev_type=True, part=False)) == ["모델명", "유형", "구분"] + V,
   "파트넘버가 하나도 없는데 칸이 남아 있다: %s" % cols_of(mk(5, dev_type=True, part=False)))
ok(cols_of(mk(5, dev_type=False, part=True)) == ["파트넘버", "모델명", "구분"] + V,
   "유형이 하나도 없는데 칸이 남아 있다: %s" % cols_of(mk(5, dev_type=False, part=True)))
ok(cols_of(mk(5, dev_type=False, part=False)) == ["모델명", "구분"] + V,
   "둘 다 없는데 칸이 남아 있다: %s" % cols_of(mk(5, dev_type=False, part=False)))
# 모델이 하나라도 가지고 있으면 칸을 둔다 — 나머지는 빈 채로 둬도 된다
one = mk(5, dev_type=False, part=False)
one[2]["part_number"] = "PN-1"
ok("파트넘버" in cols_of(one), "한 모델만 파트넘버가 있어도 칸은 있어야 한다")
# 빈 양식은 네 칸을 다 둔다 — 무엇을 쓸지 모르니까
ok(cols_of([]) == ["파트넘버", "모델명", "유형", "구분"] + V,
   "빈 양식에서 칸이 빠졌다: %s" % cols_of([]))

# 모델명·구분은 절대 안 빠진다 (줄을 알아보는 칸, 새 모델을 고르는 칸)
for shape in SHAPES.values():
    c = cols_of(shape)
    ok("모델명" in c and "구분" in c, "모델명·구분이 빠졌다: %s" % c)

# 구분 목록(양산/개발)이 구분 칸에 붙어 있는가 — 예전에는 'D' 로 박혀 있었다
ws8 = load_workbook(BytesIO(wt.build_weekly_template(
    "T", sheets=sheets_for(["2026-10"]), models=mk(4, dev_type=False, part=False))))["2026-10"]
gi = head_cols(ws8).get("구분")
rng = " ".join(str(r) for dv in ws8.data_validations.dataValidation for r in dv.sqref.ranges)
ok(gi is not None and rng.startswith("%s%d" % (chr(64 + gi), wt.FIRST_ROW)),
   "양산/개발 목록이 구분 칸에 안 붙었다: 구분 %r · 범위 %s" % (gi, rng))


if FAIL:
    print("실패 %d건" % len(FAIL))
    for f in FAIL:
        print("  -", f)
    raise SystemExit(1)
print("전부 통과 · 검사 8묶음 · 프로젝트 모양 %d가지" % len(SHAPES))

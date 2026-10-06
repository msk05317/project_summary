# -*- coding: utf-8 -*-
"""연간 내보내기 — 사람이 쓰던 '업데이트 N주차 계획' 모양.

받는 양식(weekly_template)과 모양이 다르다. 받는 쪽은 한 달이 한 시트고
세로로 길다. 이건 한 해가 한 시트고 가로로 길다. 둘을 헷갈리면 안 된다.

여기서 지키는 것
  · 열두 달이 1월부터 12월까지 빠짐없이 들어간다
  · 주차는 **OneView 달력**을 쓴다 (받아서 다시 올렸을 때 칸이 맞아야 한다)
  · 모델 순서가 등록된 순서 그대로다 (매번 줄이 움직이면 못 읽는다)
  · 유형으로 묶이고, 묶음마다 합계 · 맨 아래 총 합계
  · 값이 없는 칸은 '-' 다. 0 을 적으면 '그 주에 하나도 안 나갔다' 가 된다
  · 잔여수량 · 월 합계 · 합계는 전부 수식이고, '-' 가 섞여도 안 깨진다
  · 전년도 9~12월 실적은 작년 주차에서 더해 온다
"""
import sys
from io import BytesIO
from pathlib import Path

BACK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACK))

from openpyxl import load_workbook          # noqa: E402

import plan_export as pe                    # noqa: E402
import week_calendar as wcal                # noqa: E402

FAIL = []


def ok(c, m):
    if not c:
        FAIL.append(m)


YEAR = 2026
W9 = wcal.get_month_weeks("%d-09" % YEAR)
W10 = wcal.get_month_weeks("%d-10" % YEAR)
PW9 = wcal.get_month_weeks("%d-09" % (YEAR - 1))

# 등록 순서 그대로 — 유형이 섞여 있어도 묶음 안의 순서는 그대로여야 한다
MODELS = [
    {"id": "A-403", "name": "A-403", "dev_type": "413", "po_qty": 0, "shipped_qty": 0},
    {"id": "A-413", "name": "A-413", "dev_type": "413",
     "po_qty": 3226, "shipped_qty": 2670,
     "weekly_plan": {"%d-09" % YEAR: {W9[0]: {"plan": 20, "actual": 10},
                                      W9[2]: {"plan": 50, "actual": 30}},
                     "%d-10" % YEAR: {W10[0]: {"plan": 82, "actual": 0}},
                     "%d-09" % (YEAR - 1): {PW9[0]: {"plan": 5, "actual": 7}}}},
    {"id": "A-404", "name": "A-404", "dev_type": "413", "po_qty": 31, "shipped_qty": 15},
    {"id": "B-009", "name": "B-009", "dev_type": "009",
     "po_qty": 1117, "shipped_qty": 1117},
    {"id": "B-301", "name": "B-301", "dev_type": "009", "po_qty": 10},
    # 유형이 없는 줄 — 맨 아래 묶음으로 가야 한다
    {"id": "Z-000", "name": "Z-000", "po_qty": 7},
]

data = pe.build_year_export("엔클로저", YEAR, MODELS, sheet_name="ENCLOSURE")
wb = load_workbook(BytesIO(data))
ws = wb["ENCLOSURE"]


# ── 1. 열두 달이 다 있는가 · 주차가 OneView 달력인가 ────────────────
tops = {}
for c in range(pe.FIRST_MONTH_COL, ws.max_column + 1):
    v = ws.cell(pe.ROW_TOP, c).value
    if v and str(v).endswith("월"):
        tops[str(v)] = c
ok(len(tops) == 12, "달이 %d개다 (열두 개여야 한다): %s" % (len(tops), sorted(tops)))
ok([int(k[:-1]) for k in tops] == sorted(int(k[:-1]) for k in tops),
   "달이 1월부터 차례로 안 늘어선다: %s" % list(tops))

# 주차 라벨을 전부 거둬 OneView 달력과 맞춘다
seen = []
for c in range(pe.FIRST_MONTH_COL, ws.max_column + 1):
    v = ws.cell(pe.ROW_WEEK, c).value
    if v and str(v).upper().startswith("W"):
        seen.append(str(v).upper())
want = []
for m in range(1, 13):
    want += list(wcal.get_month_weeks("%d-%02d" % (YEAR, m)))
ok(seen == want, "주차가 OneView 달력과 다르다\n  나온 것 %s\n  있어야 할 것 %s"
   % (seen[:8], want[:8]))
# 6월이 안 비어야 한다 (사람이 쓰던 파일은 W23~W26 이 통째로 없었다)
ok(any(w in seen for w in wcal.get_month_weeks("%d-06" % YEAR)), "6월 주차가 비었다")


# ── 2. 줄 순서 · 묶음 ───────────────────────────────────────────────
rows, labels = [], []
for r in range(pe.FIRST_ROW, ws.max_row + 1):
    nm = str(ws.cell(r, pe.COL_NAME).value or "").strip()
    rows.append((r, nm, str(ws.cell(r, pe.COL_GROUP).value or "").strip()))
    if nm and nm != "합계":
        labels.append(nm)
ok(labels == ["A-403", "A-413", "A-404", "B-009", "B-301", "Z-000"],
   "등록 순서가 안 지켜졌다: %s" % labels)
ok([g for _r, _n, g in rows if g.startswith("-")] == ["-413", "-009"],
   "묶음 라벨이 다르다: %s" % [g for _r, _n, g in rows if g])
# '베이스' 는 묶음의 **둘째 줄**에 붙는다 — 한 줄짜리 묶음에는 안 붙는다
_want_base = sum(1 for _d, _ms in pe.group_models(MODELS) if len(_ms) >= 2)
ok([g for _r, _n, g in rows].count("베이스") == _want_base,
   "'베이스' 가 %d줄이어야 하는데 %d줄이다"
   % (_want_base, [g for _r, _n, g in rows].count("베이스")))
ok(rows[-1][2] == "총 합계" or str(ws.cell(ws.max_row, pe.COL_GROUP).value or "") == "총 합계",
   "맨 아래가 총 합계가 아니다: %r" % (rows[-1],))
ok(sum(1 for _r, nm, _g in rows if nm == "합계") == 3,
   "묶음 합계가 세 줄이어야 한다 (413 · 009 · 유형 없음)")


# ── 3. 값이 없는 칸은 '-' · 있는 칸은 숫자 ──────────────────────────
def cell_of(name, mon, wk, which):
    r = next(r for r, nm, _g in rows if nm == name)
    c0 = tops["%d월" % int(mon[5:7])]
    wks = wcal.get_month_weeks(mon)
    return ws.cell(r, c0 + wks.index(wk) * 2 + (0 if which == "plan" else 1)).value


ok(cell_of("A-413", "%d-09" % YEAR, W9[0], "plan") == 20,
   "적은 계획이 안 나왔다: %r" % cell_of("A-413", "%d-09" % YEAR, W9[0], "plan"))
ok(cell_of("A-413", "%d-09" % YEAR, W9[0], "actual") == 10, "적은 실적이 안 나왔다")
ok(cell_of("A-413", "%d-09" % YEAR, W9[1], "plan") == pe.DASH,
   "안 적은 칸이 '-' 가 아니다: %r" % cell_of("A-413", "%d-09" % YEAR, W9[1], "plan"))
ok(cell_of("A-404", "%d-09" % YEAR, W9[0], "plan") == pe.DASH,
   "주차 자료가 없는 모델에 숫자가 생겼다")

r403 = next(r for r, nm, _g in rows if nm == "A-403")
ok(ws.cell(r403, pe.COL_PO).value == 0, "0 이 '-' 로 바뀌었다 (0 은 진짜 0이다)")
rz = next(r for r, nm, _g in rows if nm == "Z-000")
ok(ws.cell(rz, pe.COL_SHIP).value == pe.DASH,
   "안 적은 출하실적이 0 으로 나왔다: %r" % ws.cell(rz, pe.COL_SHIP).value)


# ── 4. 전년도 9~12월 실적 ───────────────────────────────────────────
r413 = next(r for r, nm, _g in rows if nm == "A-413")
ok(str(ws.cell(pe.ROW_TOP, pe.COL_PREV0).value or "").startswith("%d년" % (YEAR - 1)),
   "전년도 머리글이 아니다: %r" % ws.cell(pe.ROW_TOP, pe.COL_PREV0).value)
ok(ws.cell(r413, pe.COL_PREV0).value == 7,
   "작년 9월 실적을 못 더했다: %r" % ws.cell(r413, pe.COL_PREV0).value)
ok(ws.cell(r413, pe.COL_PREV0 + 1).value == pe.DASH,
   "작년 자료가 없는 달이 0 으로 나왔다 — '그 달에 안 나갔다' 로 읽힌다")
ok(ws.cell(r403, pe.COL_PREV0).value == pe.DASH, "작년 자료 없는 모델에 숫자가 생겼다")


# ── 5. 수식이 '-' 를 만나도 안 깨지는가 ─────────────────────────────
#
# 사람이 쓰던 파일에서 =D-E 가 #VALUE! 를 냈다. 한쪽이 '-' 면 빼기가 안 된다.
f_left = ws.cell(rz, pe.COL_LEFT).value
ok(isinstance(f_left, str) and f_left.startswith("="), "잔여수량이 수식이 아니다")
ok("SUM(" in f_left,
   "잔여수량이 그냥 빼기다 — 한쪽이 '-' 면 #VALUE! 가 난다: %r" % f_left)
for r, nm, _g in rows:
    if nm == "합계":
        v = ws.cell(r, pe.COL_PO).value
        ok(isinstance(v, str) and v.startswith("=SUM") is False and "SUM(" in v,
           "묶음 합계가 수식이 아니다: %r" % v)
        break


# ── 6. 받는 양식과 섞이지 않았는가 ──────────────────────────────────
#
# 이건 보는 양식이다. 판가·재료비가 들어가면 받는 양식과 헷갈린다.
heads = []
for c in range(1, pe.FIRST_MONTH_COL):
    for rr in (pe.ROW_TOP, pe.ROW_SUB):
        v = ws.cell(rr, c).value
        if v:
            heads.append(str(v))
ok(not any("판가" in h or "재료비" in h for h in heads),
   "보는 양식에 판가·재료비가 들어 있다: %s" % heads)
ok(any("잔여수량" in h for h in heads), "잔여수량 칸이 없다: %s" % heads)


# ── 7. 안 쓰는 달은 접혀 있는가 ─────────────────────────────────────
#
# 쉰두 주를 다 늘어놓으면 눈이 어디를 봐야 할지 모른다. 숫자가 있는 달만
# 펼치고 나머지는 접는다 — **지우는 게 아니라 접는 것**이라, 엑셀에서
# [+] 를 누르면 그 자리에서 펼쳐진다.
def hidden_cols(w):
    out = set()
    for _k, d in w.column_dimensions.items():
        if d.hidden and d.min and d.max:
            out |= set(range(d.min, d.max + 1))
    return out


def week_cols_of(w, mon):
    c0 = tops["%d월" % int(mon[5:7])]
    return [c0 + i for i in range(len(wcal.get_month_weeks(mon)) * 2)]


def total_cols_of(w, mon):
    c0 = tops["%d월" % int(mon[5:7])]
    return [c0 + len(wcal.get_month_weeks(mon)) * 2 + j for j in (0, 1)]


hid = hidden_cols(ws)
has = pe.months_with_weeks(MODELS, YEAR)
ok(sorted(has) == sorted(["%d-09" % YEAR, "%d-10" % YEAR]),
   "숫자가 있는 달을 잘못 골랐다: %s" % has)
for m in range(1, 13):
    mon = "%d-%02d" % (YEAR, m)
    wk, tot = week_cols_of(ws, mon), total_cols_of(ws, mon)
    if mon in has:
        ok(not (set(wk) & hid), "%s 는 숫자가 있는데 접혔다" % mon)
    else:
        ok(set(wk) <= hid, "%s 가 안 접혔다" % mon)
    # 접혀도 월 합계는 늘 보인다 — 그 달 계획·실적은 읽을 수 있어야 한다
    ok(not (set(tot) & hid), "%s 의 월 합계까지 접혔다" % mon)

# 접은 것뿐이지 지운 게 아니다
ok(len(seen) == sum(len(wcal.get_month_weeks("%d-%02d" % (YEAR, m)))
                    for m in range(1, 13)),
   "접으면서 주차 칸이 사라졌다")
# 엑셀에서 펼칠 수 있게 묶음(outline)으로 접었는가
ok(any((d.outline_level or 0) > 0 for _k, d in ws.column_dimensions.items()),
   "그냥 숨겼다 — 엑셀에서 [+] 로 펼칠 수가 없다")

# 골라서 펼칠 수도 있어야 한다
d2 = pe.build_year_export("T", YEAR, MODELS, sheet_name="S",
                          expand=["%d-03" % YEAR])
w2 = load_workbook(BytesIO(d2))["S"]
h2 = hidden_cols(w2)
ok(not (set(week_cols_of(w2, "%d-03" % YEAR)) & h2), "고른 달이 안 펼쳐졌다")
ok(set(week_cols_of(w2, "%d-09" % YEAR)) <= h2,
   "고른 달만 펼쳐야 하는데 다른 달도 펼쳐졌다")


# ── 8. 라우트가 등록 순서를 안 흔드는가 ─────────────────────────────
SRC = (BACK / "main.py").read_text(encoding="utf-8")
ok('@app.get("/admin/projects/{project_key}/plan-export")' in SRC, "라우트가 없다")
_body = SRC[SRC.index("def admin_plan_export("):]
_body = _body[:_body.index("\n@app.")]
ok("order_models" not in _body,
   "라우트가 모델을 다시 세운다 — 등록 순서가 흐트러진다")
ok("get_admin_session" in _body, "라우트에 관리자 확인이 없다")

AV2 = (BACK / "admin_v2.html").read_text(encoding="utf-8")
ok("plan-export?year=" in AV2, "화면이 연간 내보내기를 안 부른다")
ok("months: str" in SRC, "라우트가 펼칠 달을 안 받는다")
ok("wp-year" in AV2, "연간 내보내기 버튼이 없다")


if FAIL:
    print("실패 %d건" % len(FAIL))
    for f in FAIL:
        print("  -", f)
    raise SystemExit(1)
print("전부 통과 · 검사 8묶음 · 달 12개 · 주차 %d개" % len(seen))

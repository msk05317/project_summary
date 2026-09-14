"""자동차사업부 제품 엑셀 파서 회귀 테스트.

이 엑셀은 한 제품이 두 줄이고(비용 / 비율), 연도별 값도 두 줄에 나뉘며
(물량 / 매출), 고객사는 세로로 병합돼 있다. 한 군데만 어긋나도 조용히
다른 제품 값이 섞인다. 그래서 실제 파일 모양 그대로 만들어 확인한다.

    cd backend && python3 tests/test_auto_import.py
"""
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

from openpyxl import Workbook          # noqa: E402

import auto_import as ai               # noqa: E402

HEAD = ["고객사", "제품명", "모델명", "최종고객사", "차종", "납품 위치",
        "제품중량\n(KG)", "주조톤수\n(TON)", "구분",
        "재료비", "부품비", "주조비", "가공비", "관리이윤", "감가상각", "포장", "물류",
        "공정비", "합계\n(원화)", "판가\n(원화)", "불량률 \n(%)", "현재단계", "SOP시점"]
YEARS = ["2026", "2027"]


def _wb(price_formula="=27.06*1350"):
    """실제 파일과 같은 모양: 병합된 고객사, 비용/비율 두 줄, 물량/매출 두 줄."""
    wb = Workbook()
    ws = wb.active
    ws.title = "양산주요제품"
    for i, h in enumerate(HEAD, start=2):
        ws.cell(4, i, h)
    ycol = 2 + len(HEAD) + 1
    ws.cell(4, ycol - 1, "총")
    for i, y in enumerate(YEARS):
        ws.cell(4, ycol + i, y)

    # 셰플러 SVCT — 비용 줄(5) + 비율 줄(6)
    ws.cell(5, 2, "셰플러")
    ws.merge_cells(start_row=5, start_column=2, end_row=8, end_column=2)   # 두 제품에 걸친 병합
    vals = ["SVCT", "", "기아자동차", "EV3", "한국\n/이천공장", 5.07, 3500, "비용",
            23488.446, 1200, 10768.11, 6347.12, 2847.974, 2951.024, 1142, 1806,
            27062.228, 50550.674, price_formula, 0.3, "양산", "2024년 4월"]
    for i, v in enumerate(vals, start=3):
        ws.cell(5, i, v)
    ws.cell(6, 10, "%")
    ws.cell(5, ycol, 116755); ws.cell(5, ycol + 1, 140650)
    ws.cell(6, ycol, 56.74876775); ws.cell(6, ycol + 1, 68.3629325)

    # 같은 고객사의 두 번째 제품 (고객사 칸은 병합돼 비어 있다)
    vals2 = ["OV1", "", "기아자동차", "EV4", "한국\n/이천공장", 5.991, 3500, "비용",
             24257, 1200, 11929.6, 10331.96, 3704.324, 2407, 1896, 1806,
             33274.884, 57531.884, "=30.73*1350", 0.24, "개발", "2025년 5월"]
    for i, v in enumerate(vals2, start=3):
        ws.cell(7, i, v)
    ws.cell(8, 10, "%")
    ws.cell(7, ycol, 68354)
    ws.cell(8, ycol, 37.60358602)

    ws.cell(9, 2, "총 매출")          # 여기서 멈춰야 한다
    ws.cell(9, ycol, 999999)
    return wb


def _parse():
    wb = _wb()
    return ai.parse(wb, wb, "양산주요제품")


def test_제품마다_두_줄을_한_줄로_읽는다():
    g = _parse()
    assert [r["product"] for r in g["rows"]] == ["SVCT", "OV1"], g["rows"]


def test_총매출_행에서_멈춘다():
    """합계 행을 제품으로 읽으면 물량이 부풀어 오른다."""
    g = _parse()
    assert len(g["rows"]) == 2, [r["product"] for r in g["rows"]]


def test_병합된_고객사가_아래줄까지_따라온다():
    g = _parse()
    assert [r["customer"] for r in g["rows"]] == ["셰플러", "셰플러"]


def test_원가를_다섯_덩어리로_묶는다():
    """재료비=재료비+부품비+포장, 공정비=주조비+가공비. 나머지는 그대로."""
    c = _parse()["rows"][0]["auto"]["cost"]
    assert round(c["material"]) == 25830, c      # 23488.446+1200+1142
    assert round(c["process"]) == 17115, c       # 10768.11+6347.12
    assert round(c["admin"]) == 2848, c
    assert round(c["depr"]) == 2951, c
    assert round(c["logi"]) == 1806, c
    # 엑셀의 합계 열(50,550.674)과 맞아야 한다
    assert round(sum(c.values())) == 50551, sum(c.values())


def test_판가는_외화단가와_환율로_나뉜다():
    a = _parse()["rows"][0]["auto"]
    assert a["price_fx"] == 27.06 and a["fx_rate"] == 1350, a
    assert round(a["price_fx"] * a["fx_rate"]) == 36531


def _wb_slim():
    """공정비가 한 칸으로만 있는 시트 (실제 '양산주요제품 (3)' 모양).

    주조비·가공비 칸이 아예 없다. 이 모양을 못 읽으면 공정비가 0 이 되고,
    원가가 통째로 비어 보이는데도 합계는 그럴듯하게 나온다.
    """
    from openpyxl import Workbook
    head = ["고객사", "제품명", "모델명", "최종고객사", "차종", "납품 위치",
            "제품중량\n(KG)", "주조톤수\n(TON)", "구분",
            "재료비", "공정비", "관리이윤", "감가상각", "물류",
            "합계\n(원화)", "판가\n(원화)", "불량률 \n(%)", "현재단계", "SOP시점"]
    wb = Workbook(); ws = wb.active; ws.title = "양산주요제품 (3)"
    for i, h in enumerate(head, start=2):
        ws.cell(4, i, h)
    ycol = 2 + len(head) + 1
    ws.cell(4, ycol - 1, "총")
    ws.cell(4, ycol, "2026")
    ws.cell(5, 2, "셰플러")
    vals = ["SVCT", "", "기아자동차", "EV3", "한국\n/이천공장", 5.07, 3500, "비용",
            25830, 17115, 2848, 2951, 1806, 50550, 36531, 0.3, "양산", "2024년 4월"]
    for i, v in enumerate(vals, start=3):
        ws.cell(5, i, v)
    ws.cell(5, ycol, 116755)
    ws.cell(6, ycol, 56.74)
    ws.cell(7, 2, "총 매출")
    return wb


def test_공정비가_한_칸인_시트도_읽는다():
    r = ai.parse(_wb_slim(), None, "양산주요제품 (3)")["rows"][0]
    c = r["auto"]["cost"]
    assert round(c["process"]) == 17115, c
    assert round(c["material"]) == 25830, c
    assert round(sum(c.values())) == 50550, sum(c.values())
    assert r["total_gap"] is None, r["total_gap"]


def test_나뉜_시트는_주조비_가공비를_쓴다():
    """같은 시트에 '공정비' 칸이 또 있어도 겹쳐 더하지 않는다."""
    r = _parse()["rows"][0]
    c = r["auto"]["cost"]
    assert round(c["process"]) == round(10768.11 + 6347.12), c
    assert r["total_gap"] is None, r["total_gap"]


def test_합계가_안_맞으면_표시한다():
    """열 하나를 놓쳐도 숫자는 그럴듯하게 나온다. 엑셀 합계와 대조해서 잡는다."""
    wb = _wb_slim()
    wb["양산주요제품 (3)"].cell(5, 16, 99999)      # 합계(원화) 칸만 딴판으로
    r = ai.parse(wb, None, "양산주요제품 (3)")["rows"][0]
    assert r["total_gap"] is not None and r["total_gap"] < 0, r["total_gap"]


def test_통화는_적지_않는다():
    """엑셀에 통화 열이 없다. 환율만 보고 USD/EUR 를 찍던 걸 뺐다."""
    a = _parse()["rows"][0]["auto"]
    assert "currency" not in a, a
    wb = _wb(price_formula="=8.55*1650")
    a2 = ai.parse(wb, wb, "양산주요제품")["rows"][0]["auto"]
    assert "currency" not in a2 and a2["fx_rate"] == 1650, a2


def test_수식이_아니면_원화로_넣는다():
    wb = _wb(price_formula=45400.5)
    a = ai.parse(wb, wb, "양산주요제품")["rows"][0]["auto"]
    assert a.get("price_krw") == 45400.5 and "price_fx" not in a, a


def test_연도별은_물량과_매출_두_줄을_짝지운다():
    a = _parse()["rows"][0]["auto"]["contract"]
    assert a["2026"] == {"qty": 116755, "revenue": 56.74876775}, a
    assert a["2027"] == {"qty": 140650, "revenue": 68.3629325}, a


def test_현재단계로_구분을_정한다():
    g = _parse()
    assert [r["group"] for r in g["rows"]] == ["양산", "개발"]


def test_고객사는_별칭까지_맞추고_모르면_비워둔다():
    """'고압주조'와 '저압주조'는 한 글자 차이다. 닮았다고 이어 붙이면 안 된다."""
    projs = [{"id": "auto_stellantis", "label": "스텔란티스", "aliases": ["스탈란티스"]},
             {"id": "auto_high_pressure_casting", "label": "고압주조", "aliases": []}]
    rows = [{"customer": "스탈란티스"}, {"customer": "저압주조"}, {"customer": "셰플러"}]
    got = ai.match_projects(rows, projs)
    assert got["스탈란티스"] == "auto_stellantis", got
    assert got["저압주조"] is None, got
    assert got["셰플러"] is None, got


def test_등록된_표기로_바꿔서_보여준다():
    """엑셀은 '스탈란티스', 등록은 '스텔란티스'. 화면에는 등록 표기가 나가야 한다.

    최종고객사 칸에는 기아자동차·CEER 처럼 프로젝트가 아닌 회사도 온다.
    그건 손대면 안 된다.
    """
    projs = [{"id": "auto_stellantis", "label": "스텔란티스",
              "aliases": ["Stellantis", "스탈란티스"]},
             {"id": "auto_valeo", "label": "발레오", "aliases": ["Valeo"]}]
    names = ai.name_labels(projs)
    assert ai.canonical("스탈란티스", names) == "스텔란티스"
    assert ai.canonical("Stellantis", names) == "스텔란티스"
    assert ai.canonical("Valeo", names) == "발레오"
    # 등록에 없는 회사는 그대로
    for n in ("기아자동차", "현대자동차", "CEER", "PSA"):
        assert ai.canonical(n, names) == n, n


if __name__ == "__main__":
    fails = 0
    for name, fn in sorted(globals().items()):
        if not name.startswith("test_"):
            continue
        try:
            fn()
            print(f"  PASS  {name}")
        except AssertionError as e:
            fails += 1
            print(f"  FAIL  {name}\n        {e}")
    print(f"\n{'실패 ' + str(fails) + '건' if fails else '전부 통과'}")
    sys.exit(1 if fails else 0)

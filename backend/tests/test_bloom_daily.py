# 블룸 '대표님 일 보고 자료' 시트 읽기
#   python3 backend/tests/test_bloom_daily.py
#
# 이 시트의 함정 두 가지
#   1) C1:I2 에 걸친 큰 제목이 C~I 모든 열 머리글에 섞여 들어온다
#   2) 품목이 세로 병합이고 그 안에 공정 행이 2~3개 붙는다
import datetime as _dt
import pathlib, sys
import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import bloom_daily_import as bd


def build():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '대표님 일 보고 자료'

    ws['C1'] = '블룸 계획 대비 실적 보고(9/14)'
    ws.merge_cells('C1:I2')                      # 머리글을 오염시키는 제목

    ws['A5'] = 'CODE'
    ws['C3'] = '품목';  ws.merge_cells('C3:C5')
    ws['D3'] = '재고';  ws.merge_cells('D3:E3')
    ws['D4'] = '출하\n대기'; ws['E4'] = '구매품 \n대기'
    ws.merge_cells('D4:D5'); ws.merge_cells('E4:E5')
    ws['F3'] = '공정';  ws.merge_cells('F3:F5')
    ws['G3'] = '공정\n재고'; ws.merge_cells('G3:G5')
    ws['H3'] = '9월 합계'; ws.merge_cells('H3:I4')
    ws['J3'] = '9/9 이전 데이터'; ws.merge_cells('J3:K4')
    ws['L3'] = '37주차'; ws.merge_cells('L3:O3')
    ws['P3'] = '38주차'; ws.merge_cells('P3:S3')
    ws['BB3'] = '비고'; ws.merge_cells('BB3:BB5')

    for col, day in (('L', '2026-09-12'), ('N', '2026-09-13'),
                     ('P', '2026-09-14'), ('R', '2026-09-15')):
        ws[f'{col}4'] = _dt.datetime.fromisoformat(day + ' 00:00:00')
        ws.merge_cells(f'{col}4:{chr(ord(col) + 1)}4')
    for c in range(8, 20):                       # H~S 계획/실적 쌍
        ws.cell(5, c).value = '계획' if c % 2 == 0 else '실적'

    rows = [
        # 품목, code, 출하대기, 구매품대기, [(공정, wip, 월계획, 월실적, 이전계획, 이전실적, 9/12계, 9/12실, 9/13계, 9/13실, 9/14계, 비고)]
        ('Corva \nKPE', '172146\n868000', 80, 228, [
            ('NCT(박닌)',  0,    1050, 233, None, None, 50, 40, 50, 31, 50, ''),
            ('조립(박닌)', 20,   1050, 213, None, None, 50, 42, 50, 42, 50, ''),
            ('출하',       None, 1162, 112, 313,  112,   0, None, 0, None, 0, '퓨즈 대기'),
        ]),
        ('BOP Assy', None, 5, 44, [
            ('조립(박닌)', 0,    1100, 44,  None, None, 55, 16, 55, 28, 55, ''),
            ('출하',       None, 1000, 0,   None, None,  0, None, 0, None, 0, ''),
        ]),
    ]
    r = 6
    for item, code, ws_, wp, steps in rows:
        top = r
        for (step, wip, mp, ma, pp, pa, p12, a12, p13, a13, p14, note) in steps:
            ws.cell(r, 3).value = item if r == top else None
            ws.cell(r, 1).value = code if r == top else None
            ws.cell(r, 4).value = ws_ if r == top else None
            ws.cell(r, 5).value = wp if r == top else None
            ws.cell(r, 6).value = step
            ws.cell(r, 7).value = wip
            ws.cell(r, 8).value = mp;  ws.cell(r, 9).value = ma
            ws.cell(r, 10).value = pp; ws.cell(r, 11).value = pa
            ws.cell(r, 12).value = p12; ws.cell(r, 13).value = a12
            ws.cell(r, 14).value = p13; ws.cell(r, 15).value = a13
            ws.cell(r, 16).value = p14
            ws.cell(r, 54).value = note or None
            r += 1
        for col in (1, 3, 4, 5):                 # 품목/코드/재고는 세로 병합
            if len(steps) > 1:
                ws.merge_cells(start_row=top, start_column=col,
                               end_row=r - 1, end_column=col)

    r += 1
    ws.cell(r, 4).value = '◆KPE- 구매품 1종(퓨즈) ETA:9/16(2000set)'
    ws.cell(r + 1, 4).value = '◆SL7- 터미널 블록 ETA: 9/18'
    ws.cell(r + 2, 4).value = '◆WDM- FTG,PEX 금주중 입고 예정'
    return wb


out = bd.parse_daily(build())
ok = 0

assert out['sheet'] == '대표님 일 보고 자료', out['sheet']
ok += 1
assert out['title'] == '블룸 계획 대비 실적 보고(9/14)', out['title']
assert out['report_date'] == '2026-09-14', out['report_date']
ok += 1
assert out['prior_label'] == '9/9 이전 데이터', out['prior_label']
ok += 1
assert out['dates'] == ['2026-09-12', '2026-09-13', '2026-09-14', '2026-09-15'], out['dates']
ok += 1

items = {i['item']: i for i in out['items']}
assert list(items) == ['Corva KPE', 'BOP Assy'], list(items)
ok += 1

kpe = items['Corva KPE']
assert kpe['code'] == '172146 868000', kpe['code']       # 줄바꿈이 한 칸으로
assert (kpe['wait_ship'], kpe['wait_part']) == (80, 228), kpe
ok += 1

steps = {s['step']: s for s in kpe['steps']}
assert list(steps) == ['NCT(박닌)', '조립(박닌)', '출하'], list(steps)
ok += 1

nct = steps['NCT(박닌)']
assert (nct['month_plan'], nct['month_actual']) == (1050, 233), nct
assert nct['wip'] == 0
assert nct['days']['2026-09-13'] == {'plan': 50, 'actual': 31}, nct['days']
ok += 1

# 실적이 아직 안 들어온 칸은 0 이 아니라 None 이어야 한다 (0 개 했다와 다르다)
assert nct['days']['2026-09-14'] == {'plan': 50, 'actual': None}, nct['days']['2026-09-14']
ok += 1
# 계획 0 은 진짜 0 이다
assert steps['출하']['days']['2026-09-13'] == {'plan': 0, 'actual': None}
ok += 1
# 두 값이 다 비면 그 날은 아예 없다
assert '2026-09-15' not in nct['days'], nct['days']
ok += 1

assert (steps['출하']['prior_plan'], steps['출하']['prior_actual']) == (313, 112)
assert steps['출하']['note'] == '퓨즈 대기', steps['출하']['note']
ok += 1

# 공정이 2개뿐인 품목
assert [s['step'] for s in items['BOP Assy']['steps']] == ['조립(박닌)', '출하']
assert items['BOP Assy']['code'] == ''
ok += 1

etas = [(n['eta'], n['text'][:6]) for n in out['notes']]
assert len(etas) == 3, out['notes']
assert etas[0] == ('9/16', 'KPE- 구'), etas
assert etas[1][0] == '9/18' and etas[2][0] == '', etas
ok += 1

print(f'전부 통과 ({ok}/14)')

# ── 실제 파일로 한 번 더 (있을 때만) ────────────────────────────────
import glob, os
_real = sorted(glob.glob(os.path.join(os.path.dirname(__file__), 'fixtures', '블룸*.xlsx')))
if _real:
    r = bd.parse_daily(openpyxl.load_workbook(_real[0], data_only=True))
    assert len(r['items']) >= 5, f"실제 파일에서 품목을 {len(r['items'])}개만 읽었다"
    assert len(r['dates']) >= 15, r['dates']
    print(f"  실제 파일 {os.path.basename(_real[0])}: "
          f"품목 {len(r['items'])} · 날짜 {len(r['dates'])} · 메모 {len(r['notes'])}")

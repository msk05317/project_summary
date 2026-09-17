# 매출 — 엑셀이 원본이다
#   python3 backend/tests/test_revenue.py
#
# "결론적으로 저 엑셀이 명확한 데이터를 가지고 있어"
#
# 모델 판가 × 수량으로 계산하던 매출을 그만두고 일일보고(실적)와 월
# 계획(Estimate Revenue in Sep)에서 받는다. 지키는 것 셋.
#   1. 실적은 '날짜' 로 저장한다. 엑셀은 9월을 W36~W39 로 보고 OneView 는
#      W36~W40 으로 봐서 한 주가 통째로 어긋났다.
#   2. 같은 파일을 두 번 올려도 더해지지 않는다.
#   3. 모르는 엑셀 행은 조용히 빠뜨리지 않고 세워서 물어본다.
import io, pathlib, sys

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import revenue as R                                    # noqa: E402
import openpyxl                                        # noqa: E402

SRC = (ROOT / 'main.py').read_text(encoding='utf-8')
ok = 0


def daily_book(week, dates, rows, extra_plan_col=False):
    """일일보고 한 장을 만든다. extra_plan_col 이면 W9 처럼 칸이 하나 더 낀다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'W%d' % week
    ws.cell(row=2, column=2, value='구분')
    col = 3
    spans = []
    for d in dates:
        ws.cell(row=2, column=col, value=d)
        ws.cell(row=3, column=col, value='Plan ($)')
        c = col + 1
        if extra_plan_col:
            ws.cell(row=3, column=c, value='Plan ($) update'); c += 1
        ws.cell(row=3, column=c, value="Q'ty")
        ws.cell(row=3, column=c + 1, value='Amount')
        spans.append((col, c, c + 1))
        col = c + 2
    ws.cell(row=2, column=col, value='W%d Total' % week)
    for i, (label, vals) in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=2, value=label)
        for (pc, qc, ac), v in zip(spans, vals):
            if v is None:
                continue
            plan, qty, amt = v
            ws.cell(row=r, column=pc, value=plan)
            ws.cell(row=r, column=qc, value=qty)
            ws.cell(row=r, column=ac, value=amt)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def plan_book(title, items):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sum'
    ws.cell(row=1, column=1, value='Commodity')
    ws.cell(row=1, column=2, value=title)
    for i, (k, v) in enumerate(items):
        ws.cell(row=2 + i, column=1, value=k)
        ws.cell(row=2 + i, column=2, value=v)
    ws.cell(row=2 + len(items), column=2,
            value=sum(v for _, v in items if isinstance(v, (int, float))))
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


# ── 날짜별로 읽는다 ──
raw = daily_book(37, ['09 / 07', '09 / 08'], [
    ('엔클로져 (Encloser)', [(100.0, 1, 120.0), (200.0, 2, 210.0)]),
    ('파워박스 (Powerbox)', [(50.0, 3, 40.0), None]),
    ('총합계 (Grand Total)', [(150.0, 4, 160.0), (200.0, 2, 210.0)]),
])
p = R.parse_daily(raw, 2026)
assert sorted(p['days']) == ['2026-09-07', '2026-09-08'], p['days'].keys()
assert '총합계 (Grand Total)' not in p['rows'], '합계 줄까지 품목으로 세었다'
assert p['days']['2026-09-07']['엔클로져 (Encloser)']['amount'] == 120.0
assert p['days']['2026-09-08']['엔클로져 (Encloser)']['qty'] == 2
ok += 1

# ── 칸이 하나 더 껴 있어도 금액 칸을 제대로 찾는다 (W9) ──
raw2 = daily_book(9, ['02 / 23'], [('메탈 가공 (Metal Machining)',
                                    [(11.0, 5, 999.0)])], extra_plan_col=True)
p2 = R.parse_daily(raw2, 2026)
cell = p2['days']['2026-02-23']['메탈 가공 (Metal Machining)']
assert cell['amount'] == 999.0, f"금액 칸을 잘못 읽었다: {cell}"
assert cell['qty'] == 5 and cell['plan'] == 11.0, cell
ok += 1

# ── 파일의 해를 쓴다 ──
#
# W35 시트에 08/31 이 들어 있는데 그 날의 ISO 주차는 36 이다. 주차로
# 해를 되짚으면 35 에 맞는 2025년으로 밀려난다 (실제로 그랬다).
raw3 = daily_book(35, ['08 / 31'], [('메탈 가공 (Metal Machining)',
                                     [(1.0, 1, 2.0)])])
p3 = R.parse_daily(raw3, 2026)
assert list(p3['days']) == ['2026-08-31'], p3['days'].keys()
ok += 1

# 연말연시만 예외. 2026년 파일의 W01 시트는 2025-12-29 부터라 거기 적힌
# 12/29 는 지난 해다. 반대로 W53 시트는 2027-01-03 까지 가므로 01/02 는 다음 해.
assert R._year_for(1, 12, 29, 2026) == 2025, '12월 W01 이 다음 해로 갔다'
assert R._year_for(53, 1, 2, 2026) == 2027, '1월 W53 이 지난 해로 갔다'
assert R._year_for(37, 9, 7, 2026) == 2026
assert R._year_for(9, 2, 30, 2026) is None, '2월 30일 같은 칸을 날짜로 만들었다'
ok += 1

# ── 지웠던 날은 다시 올리면 사라진다 ──
#
# 날짜마다 덮어쓰기만 하면, 잘못 적은 날을 지우고 다시 올려도 옛 값이
# 그대로 남는다. 빈 칸은 파일에 아예 안 나오기 때문이다.
v1 = daily_book(37, ['09 / 07', '09 / 08'], [
    ('파워박스 (Powerbox)', [(0.0, 1, 100000.0), (0.0, 2, 200000.0)])])
v2 = daily_book(37, ['09 / 07', '09 / 08'], [
    ('파워박스 (Powerbox)', [(0.0, 1, 100000.0), None])])
st = R.blank()
R.apply_daily(st, R.parse_daily(v1, 2026), 'v1.xlsx')
assert R.month_view(st, '2026-09')['actual'] == 300000
got = R.apply_daily(st, R.parse_daily(v2, 2026), 'v2.xlsx')
assert R.month_view(st, '2026-09')['actual'] == 100000, '지운 날이 그대로 남았다'
assert got['dropped'] == ['2026-09-08'], got
ok += 1

# 파일이 안 덮는 기간은 건드리지 않는다
other = daily_book(30, ['07 / 20'], [('파워박스 (Powerbox)', [(0.0, 1, 7.0)])])
R.apply_daily(st, R.parse_daily(other, 2026), 'o.xlsx')
R.apply_daily(st, R.parse_daily(v2, 2026), 'v2.xlsx')
assert R.month_view(st, '2026-07')['actual'] == 7, '다른 달까지 지웠다'
ok += 1

# ── 마지막 날짜 칸이 주 합계를 삼키지 않는다 ──
#
# 'W37 Total' 묶음의 머리글이 날짜 줄 위에 있고 마지막 날짜 칸의
# 아랫줄 이름이 비어 있으면, 그 하루가 그 주 전체 합계로 잡힌다.
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'W37'
ws.cell(row=1, column=8, value='W37 Total')          # 머리글이 한 줄 위에
ws.cell(row=2, column=2, value='구분')
ws.cell(row=2, column=3, value='09 / 07')
ws.cell(row=3, column=3, value='Plan ($)')
ws.cell(row=3, column=4, value="Q'ty"); ws.cell(row=3, column=5, value='Amount')
ws.cell(row=2, column=6, value='09 / 08')            # 아랫줄 이름 없음
ws.cell(row=3, column=9, value='Amount')             # 합계 묶음의 금액
ws.cell(row=4, column=2, value='파워박스 (Powerbox)')
ws.cell(row=4, column=3, value=0); ws.cell(row=4, column=4, value=1)
ws.cell(row=4, column=5, value=100)
ws.cell(row=4, column=6, value=0); ws.cell(row=4, column=7, value=2)
ws.cell(row=4, column=8, value=200)
ws.cell(row=4, column=9, value=300)                  # 주 합계
buf = io.BytesIO(); wb.save(buf)
pb2 = R.parse_daily(buf.getvalue(), 2026)
assert pb2['days']['2026-09-08']['파워박스 (Powerbox)']['amount'] == 200, \
    '마지막 날이 주 합계를 삼켰다: ' + str(pb2['days']['2026-09-08'])
ok += 1

# 아포스트로피 없는 Qty 도 수량이다
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'W9'
ws.cell(row=2, column=2, value='구분'); ws.cell(row=2, column=3, value='02 / 23')
ws.cell(row=3, column=3, value='Plan ($)')
ws.cell(row=3, column=4, value='Plan ($) update')
ws.cell(row=3, column=5, value='Qty'); ws.cell(row=3, column=6, value='Amount')
ws.cell(row=4, column=2, value='파워박스 (Powerbox)')
ws.cell(row=4, column=3, value=11); ws.cell(row=4, column=4, value=77777)
ws.cell(row=4, column=5, value=5); ws.cell(row=4, column=6, value=999)
buf = io.BytesIO(); wb.save(buf)
c9 = R.parse_daily(buf.getvalue(), 2026)['days']['2026-02-23']['파워박스 (Powerbox)']
assert c9 == {'plan': 11.0, 'qty': 5.0, 'amount': 999.0}, c9
ok += 1

# ── 날짜 머리글을 못 찾은 시트는 말해 준다 ──
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'W36'
ws.cell(row=2, column=2, value='구분'); ws.cell(row=4, column=2, value='파워박스 (Powerbox)')
buf = io.BytesIO(); wb.save(buf)
sk = R.parse_daily(buf.getvalue(), 2026)
assert sk['skipped'] == ['W36'], '못 읽은 시트를 조용히 넘겼다'
ok += 1

# ── 같은 파일을 두 번 올려도 더해지지 않는다 ──
st = R.blank()
R.apply_daily(st, R.parse_daily(raw, 2026), 'a.xlsx')
first = R._sum_all(st['days'])
R.apply_daily(st, R.parse_daily(raw, 2026), 'a.xlsx')
assert R._sum_all(st['days']) == first, '두 번 올렸더니 매출이 늘었다'
assert first == 120.0 + 210.0 + 40.0, first
ok += 1

# ── 월 계획 ──
# 글자로 적힌 숫자를 버리면 그 품목이 통째로 사라지고 '모르는 항목' 에도 안 뜬다.
pb_odd = plan_book('Estimate Revenue in Sep',
                   [('Metal', '2,400,000'), ('Frame', 800000), ('Frame', 600000)])
odd = R.parse_plan(pb_odd, 'x.xlsx')
assert odd['items']['Metal'] == 2400000.0, odd['items']
assert odd['items']['Frame'] == 1400000.0, '같은 이름 두 줄 중 하나가 사라졌다'
assert odd['order'].count('Frame') == 1
ok += 1

pb = plan_book('Estimate Revenue in Sep', [('Metal', 300.0), ('PBX', 200.0),
                                           ('Space X', 0.0)])
pp = R.parse_plan(pb, 'Estimate Revenue in Sep.xlsx')
assert pp['month'].endswith('-09'), pp['month']
assert pp['items']['Metal'] == 300.0 and len(pp['items']) == 3, pp['items']
ok += 1

# ── 품목별로 접어서 본다 ──
#
# 엔클로저는 계획 파일에 따로 없고 Metal 안에 들어 있다.
st = R.blank()
R.apply_daily(st, R.parse_daily(raw, 2026), 'a.xlsx')
R.apply_plan(st, R.parse_plan(pb, 'x.xlsx'), 'p.xlsx', '2026-09')
v = R.month_view(st, '2026-09')
by = {r['item']: r for r in v['items']}
assert by['메탈 가공']['actual'] == 330, by['메탈 가공']       # 엔클로져 120+210
assert by['메탈 가공']['plan'] == 300
assert by['파워박스']['actual'] == 40 and by['파워박스']['plan'] == 200
assert v['actual'] == 370 and v['plan'] == 500, (v['actual'], v['plan'])
assert v['rate'] == 74
assert v['items'][0]['item'] == '메탈 가공', '계획 큰 것부터가 아니다'
ok += 1

# 기준일은 실적이 들어온 마지막 날. 계획만 적힌 앞날이 아니다.
assert v['as_of'] == '2026-09-08', v['as_of']
raw4 = daily_book(38, ['09 / 14'], [('엔클로져 (Encloser)', [(900.0, 0, 0.0)])])
R.apply_daily(st, R.parse_daily(raw4, 2026), 'b.xlsx')
assert R.month_view(st, '2026-09')['as_of'] == '2026-09-08', '앞날이 기준일이 됐다'
ok += 1

# 다른 달은 안 섞인다
assert R.month_view(st, '2026-08')['actual'] == 0
# 누적은 보고 있는 달까지. 8월을 보는데 9월이 누적에 들어가면 안 맞는다.
assert R.month_view(st, '2026-08')['ytd'] == 0, '뒤에 오는 달이 누적에 들어갔다'
assert R.month_view(st, '2026-09')['ytd'] == 370
ok += 1

# ── 모르는 행은 합계에서 빠지고, 물어볼 목록에 오른다 ──
raw5 = daily_book(37, ['09 / 07'], [('UCT (SM)', [(0.0, 9, 4706.0)])])
st2 = R.blank()
R.apply_daily(st2, R.parse_daily(raw5, 2026), 'c.xlsx')
assert R.month_view(st2, '2026-09')['actual'] == 0, '연결 안 한 행이 합계에 들어갔다'
un = R.unmapped_in_store(st2)
assert un and un[0]['label'] == 'UCT (SM)' and un[0]['amount'] == 4706, un
assert R.unknown_rows(st2, ['UCT (SM)', '파워박스 (Powerbox)']) == ['UCT (SM)']
ok += 1

# 연결해 주면 바로 합계에 들어온다
st2['map']['UCT (SM)'] = '시트메탈 · 프레임'
assert R.month_view(st2, '2026-09')['actual'] == 4706
assert R.unmapped_in_store(st2) == []
ok += 1

# ── 이름은 줄바꿈·겹공백을 무시하고 붙는다 ──
assert R.norm(' 메탈 가공\n(Metal  Machining) ') == '메탈 가공 (Metal Machining)'
assert R.norm('메탈 가공 (Metal Machining)') in R.DEFAULT_MAP
ok += 1

# ── 저장·불러오기 ──
import tempfile
with tempfile.TemporaryDirectory() as td:
    f = pathlib.Path(td) / 'revenue.json'
    st2['map']['엔클로져 (Encloser)'] = '엔클로저 따로'
    R.save(f, st2)
    got = R.load(f)
    assert got['map']['엔클로져 (Encloser)'] == '엔클로저 따로'
    assert got['map']['파워박스 (Powerbox)'] == '파워박스', '기본 매핑이 사라졌다'
    assert R.month_view(got, '2026-09')['actual'] == 4706
    assert R.load(pathlib.Path(td) / '없는파일.json')['map'], '없는 파일이 빈 매핑'
    ok += 1

    # 기본 매핑을 끊으면 끊긴 채로 있어야 한다.
    # 예전에는 불러올 때 기본값을 다시 깔아서, 끊어도 새로고침하면 살아났다.
    got['map'].pop('엔클로져 (Encloser)')
    R.save(f, got)
    again = R.load(f)
    assert '엔클로져 (Encloser)' not in again['map'], '끊은 매핑이 되살아났다'
    assert again['map']['파워박스 (Powerbox)'] == '파워박스'
    ok += 1

    # 깨진 파일을 빈 값으로 돌려주면 다음 저장이 멀쩡한 파일을 덮어쓴다
    f.write_text('{깨진', encoding='utf-8')
    try:
        R.load(f)
        raise AssertionError('깨진 파일을 그냥 읽었다')
    except R.RevenueFileBroken:
        pass
    ok += 1

# ── 여럿이 동시에 저장해도 파일이 깨지지 않는다 ──
#
# 임시 파일 이름이 하나면 두 사람이 같은 파일을 동시에 자르고 쓴다.
import threading
with tempfile.TemporaryDirectory() as td:
    f = pathlib.Path(td) / 'revenue.json'
    big = R.blank()
    for d in range(1, 29):
        big['days']['2026-08-%02d' % d] = {
            '파워박스 (Powerbox)': {'plan': 1.0, 'qty': 1.0, 'amount': 1000.0}}
    R.save(f, big)
    errs = []

    def hammer():
        for _ in range(12):
            try:
                R.save(f, R.load(f))
            except Exception as e:      # noqa: BLE001
                errs.append(repr(e))

    ts = [threading.Thread(target=hammer) for _ in range(6)]
    for t in ts:
        t.start()
    for t in ts:
        t.join()
    assert not errs, errs[:3]
    assert len(R.load(f)['days']) == 28, '동시에 저장했더니 내용이 날아갔다'
    assert not list(pathlib.Path(td).glob('*.tmp')), '임시 파일이 남았다'
ok += 1

# ── 서버가 받는 곳 ──
for path in ('@app.get("/revenue")', '@app.post("/admin/revenue/import")',
             '@app.put("/admin/revenue/map")', '@app.get("/admin/revenue/state")'):
    assert path in SRC, f'{path} 가 없다'
assert 'import revenue as _rev' in SRC, '서버가 매출 모듈을 안 쓴다'
assert 'mode != "commit"' in SRC, '미리보기가 그냥 저장해버린다'
# 계획 파일은 제 달을 안다. 8월을 보다가 9월 파일을 올려도 8월에 덮으면 안 된다.
assert 'parsed.get("month") or (month or "").strip()' in SRC, \
    '계획 파일의 달보다 화면에서 고른 달이 이긴다'
assert 'REVENUE_FILE' in SRC and 'revenue.json' in SRC
# '2026' 이나 '2026-9' 가 들어오면 그 달이 통째로 0 으로 보인다
assert '_check_month' in SRC and SRC.count('_check_month(') >= 4, \
    '월 형식을 한 곳에서만 본다'
# 깨진 파일을 빈 값으로 돌려주면 다음 저장이 덮어쓴다
assert 'RevenueFileBroken' in SRC, '깨진 파일을 그냥 넘긴다'
ok += 1

# ── 앱: 홈 매출 카드가 이 값을 쓴다 ──
LIB = ROOT.parent / 'mobile' / 'lib'
SVC = (LIB / 'services' / 'revenue_service.dart')
assert SVC.exists(), '앱에 매출 서비스가 없다'
S = SVC.read_text(encoding='utf-8')
assert '/revenue' in S, '앱이 /revenue 를 안 본다'
C = (LIB / 'components' / 'home' / 'exec_revenue_card.dart').read_text(encoding='utf-8')
assert 'RevenueMonth' in C, '매출 카드가 아직 모델 계산값을 쓴다'
ok += 1

# ── admin: 매출 관리 화면 ──
A = (ROOT / 'admin_v2.html').read_text(encoding='utf-8')
assert 'data-page="revenue"' in A, 'admin 메뉴에 매출 관리가 없다'
assert '/admin/revenue/import' in A, 'admin 이 파일을 안 올린다'
assert 'mode' in A and 'preview' in A, '미리보기 없이 바로 저장한다'
assert '모르는 항목' in A, '모르는 행을 안 물어본다'
ok += 1

print(f'전부 통과 · {ok}개 항목')

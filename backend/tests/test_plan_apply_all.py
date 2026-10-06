# -*- coding: utf-8 -*-
"""주차 양식 한 장이 세 가지를 같이 싣는다 — 모델 값 · 주차 숫자 · 미달 사유.

양식을 한 장으로 합쳤으니 읽는 쪽도 한 번에 받아야 한다. 예전에는 주차
숫자만 들어가고, 판가·재료비·PO·실적은 무시되고, 새 모델 줄은 '건너뛰는
행' 으로만 떴다.

여기서 못 박아 두는 것
  · 빈 칸은 '그대로 둬라' — 안 적은 칸이 지금 값을 지우면 안 된다
  · 새 모델은 **만들어진다** — 미리보기에 '새로 등록될 모델' 로 보이고
    저장하면 실제로 들어간다 (보인 것과 저장되는 것이 같아야 한다)
  · 미달 사유는 주차 단위 — 화면의 '계획 미달' 패널과 같은 자리

main.py 는 fastapi 가 있어야 import 되므로, 검사할 함수만 원문에서 떼어
그대로 실행한다. 사본을 만들어 비교하면 언젠가 한쪽만 고쳐진다.
"""
import sys
from io import BytesIO
from pathlib import Path

BACK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACK))

from openpyxl import load_workbook          # noqa: E402

import week_calendar as wcal                # noqa: E402
import weekly_template as wt                # noqa: E402

FAIL = []


def ok(c, m):
    if not c:
        FAIL.append(m)


SRC = (BACK / "main.py").read_text(encoding="utf-8")


def _grab(name):
    i = SRC.index("def %s(" % name)
    j = SRC.index("\ndef ", i + 1)
    return SRC[i:j]


def _const(name, until):
    return "%s = %s" % (name, SRC.split("%s = " % name)[1].split(until)[0])


NS = {"_WEEK_REASON_KINDS": ("문제", "참고")}
for _n in ("_as_money", "_as_int"):
    exec(_grab(_n), NS)
exec(_const("_PLAN_MODEL_COLS", "\n#: 목록이"), NS)
exec(_const("_PLAN_STOP", "\n\n"), NS)
exec(_const("_PLAN_FIELDS", "\n\n"), NS)
for _n in ("_plan_over", "_plan_field_empty",
           "_plan_sheet_models", "_plan_model_changes",
           "_plan_apply_models", "_plan_apply_week_notes"):
    exec(_grab(_n), NS)

read_models = NS["_plan_sheet_models"]
model_changes = NS["_plan_model_changes"]
apply_models = NS["_plan_apply_models"]
apply_notes = NS["_plan_apply_week_notes"]
field_empty = NS["_plan_field_empty"]
plan_over = NS["_plan_over"]


# ── 양식 하나 만들어 사람이 채운 것처럼 ─────────────────────────────
MODELS = [
    {"id": "A-001", "name": "A-001", "group": "양산", "dev_type": "413",
     "price": 1000, "material_cost": 400, "po_qty": 50, "shipped_qty": 20},
    {"id": "A-002", "name": "A-002", "group": "양산", "dev_type": "413",
     "price": 2000, "material_cost": 800, "po_qty": 10, "shipped_qty": 10},
    # PO 가 5 였는데 0 으로 적는다 — 0 은 '비움' 이 아니라 진짜 0이다
    {"id": "A-003", "name": "A-003", "group": "개발", "dev_type": "009",
     "po_qty": 5},
]
MON = "2026-10"
WEEKS = wcal.get_month_weeks(MON)

wb = load_workbook(BytesIO(wt.build_weekly_template(
    "T", sheets=[{"month": MON, "weeks": WEEKS}], models=MODELS)))
ws = wb[MON]

NF = 0
while NF < 14:
    v = ws.cell(wt.HDR_TOP, NF + 1).value
    t = "" if v is None else str(v).strip()
    if not t or t == "월 합계" or t.upper().startswith("W"):
        break
    NF += 1
ok(NF == len(wt.FIXED) - 1 + len(wt.VALUE),
   "왼쪽 칸 수가 예상과 다르다: %d" % NF)

COL = {}
for c in range(1, NF + 1):
    COL[str(ws.cell(wt.HDR_TOP, c).value).strip()] = c

R = wt.FIRST_ROW


def row_of(name):
    """양식은 유형끼리 모아 놓으므로 적은 순서와 줄 순서가 다르다."""
    for r in range(R, R + len(MODELS) + 2):
        if str(ws.cell(r, COL["모델명"]).value or "").strip() == name:
            return r
    raise AssertionError("%s 줄을 못 찾았다" % name)


# A-001: 판가만 고친다 (나머지는 그대로 둬야 한다)
ws.cell(row_of("A-001"), COL["판가($)"], 1250.5)
# A-002: 판가·재료비를 지운다 → '그대로 둬라'
# openpyxl 은 cell(r, c, None) 로는 안 지워진다 — .value 로 비운다
ws.cell(row_of("A-002"), COL["판가($)"]).value = None
ws.cell(row_of("A-002"), COL["재료비($)"]).value = None
# A-003: 판가는 비어 있던 칸을 채우고(900), PO 는 있던 값을 덮는다(5 → 0)
ws.cell(row_of("A-003"), COL["판가($)"], 900)
ws.cell(row_of("A-003"), COL["PO수량"], 0)
# 새 모델 한 줄
NEW = R + len(MODELS)
ws.cell(NEW, COL["모델명"], "A-999")
ws.cell(NEW, COL["유형"], "009")
ws.cell(NEW, COL["구분"], "개발")
ws.cell(NEW, COL["판가($)"], 777.5)
ws.cell(NEW, COL["재료비($)"], 300)
ws.cell(NEW, COL["PO수량"], 12)
# 주차 숫자 + 미달 사유
ws.cell(row_of("A-001"), NF + 1, 20)
ws.cell(row_of("A-001"), NF + 2, 12)
n_rows = len(MODELS) + wt.SPARE_ROWS
r_note = wt.FIRST_ROW + n_rows + 1
ws.cell(r_note, NF + 1, "자재 입고 지연")
ws.cell(r_note, NF + 1 + wt.WEEK_N * 2, "고객 요청으로 보류")

buf = BytesIO()
wb.save(buf)
ws2 = load_workbook(BytesIO(buf.getvalue()), data_only=True)[MON]


# ── 1. 읽히는가 ─────────────────────────────────────────────────────
rows = read_models(ws2)
by = {r["name"]: r for r in rows}
ok(len(rows) == len(MODELS) + 1,
   "모델 줄을 %d개 읽었다 (%d개여야 한다)" % (len(rows), len(MODELS) + 1))
ok("합계" not in by and "미달 사유" not in by,
   "합계·미달 사유 줄을 모델로 읽었다: %s" % list(by))
ok(by["A-001"]["price"] == 1250.5, "고친 판가가 안 읽힌다: %r" % by["A-001"].get("price"))
ok(by["A-002"]["price"] is None and by["A-002"]["material_cost"] is None,
   "지운 칸이 0 으로 읽힌다 — 그러면 지금 값을 덮어쓴다: %r" % by["A-002"])
ok(by["A-003"]["po_qty"] == 0, "0 이 안 읽힌다 (0 은 진짜 0이다): %r" % by["A-003"].get("po_qty"))
ok(by["A-999"]["price"] == 777.5 and by["A-999"]["po_qty"] == 12,
   "새 모델 값이 안 읽힌다: %r" % by["A-999"])


# ── 2. 무엇이 바뀌는지 미리 보이는가 ────────────────────────────────
import copy  # noqa: E402

proj = {"models": copy.deepcopy(MODELS)}
changed, created = model_changes(proj, rows)
ch = {c["name"]: c["fields"] for c in changed}
ok([c["name"] for c in created] == ["A-999"],
   "새로 생길 모델이 A-999 하나가 아니다: %s" % [c["name"] for c in created])
ok("A-001" in ch and ch["A-001"]["price"]["to"] == 1250.5,
   "판가가 바뀐다는 게 안 보인다: %r" % ch.get("A-001"))
ok("A-002" not in ch, "비운 칸을 '바뀐다' 로 봤다 — 그대로 둬야 한다: %r" % ch.get("A-002"))
ok("A-003" in ch and ch["A-003"]["po_qty"]["to"] == 0,
   "PO 0 을 못 봤다: %r" % ch.get("A-003"))

# 칸마다 '비어 있던 것인가' 가 붙어야 화면이 둘로 가를 수 있다
ok(ch["A-003"]["price"].get("fill") is True,
   "비어 있던 판가를 '덮어씀' 으로 봤다: %r" % ch["A-003"]["price"])
ok(ch["A-003"]["po_qty"].get("fill") is False,
   "값이 있던 PO 를 '채움' 으로 봤다: %r" % ch["A-003"]["po_qty"])
ok(ch["A-001"]["price"].get("fill") is False,
   "값이 있던 판가를 '채움' 으로 봤다: %r" % ch["A-001"]["price"])

# 0 은 비어 있는 것으로 본다 — PO 0 은 'PO 가 없다' 지 '0 대를 받았다' 가 아니다
ok(field_empty(0, "po_qty") and field_empty(None, "po_qty")
   and not field_empty(5, "po_qty"),
   "숫자 칸의 '비었다' 판단이 틀렸다")
ok(field_empty("", "dev_type") and field_empty("  ", "dev_type")
   and not field_empty("413", "dev_type"),
   "글자 칸의 '비었다' 판단이 틀렸다")
ok(plan_over("1") and plan_over("on") and not plan_over("") and not plan_over(None),
   "덮어쓰기 체크를 잘못 읽는다")


# ── 3. 저장 — 기본은 '빈 칸만 채움' ────────────────────────────────
#
# 두 달 전에 받아 둔 파일을 다시 올려도 그 사이 손으로 고친 값이
# 옛 값으로 돌아가면 안 된다. 그래서 값이 있는 칸은 그냥 둔다.
res = apply_models(proj, rows)
got = {m["name"]: m for m in proj["models"]}
ok(res["created"] == ["A-999"], "만들어진 모델이 다르다: %s" % res["created"])
ok(len(proj["models"]) == len(MODELS) + 1,
   "모델 수가 %d개다" % len(proj["models"]))
ok(got["A-001"]["price"] == 1000,
   "값이 있는 판가를 기본 저장에서 덮어썼다: %r" % got["A-001"].get("price"))
ok(got["A-003"]["price"] == 900,
   "비어 있던 판가를 안 채웠다: %r" % got["A-003"].get("price"))
ok(got["A-003"]["po_qty"] == 5,
   "값이 있는 PO 를 기본 저장에서 덮어썼다: %r" % got["A-003"].get("po_qty"))
ok(got["A-002"]["price"] == 2000 and got["A-002"]["material_cost"] == 800,
   "비운 칸이 기존 값을 지웠다: %r" % {k: got["A-002"].get(k)
                                      for k in ("price", "material_cost")})
ok(got["A-999"]["group"] == "개발" and got["A-999"]["dev_type"] == "009"
   and got["A-999"]["price"] == 777.5,
   "새 모델이 제대로 안 만들어졌다: %r" % got["A-999"])
ok(res["filled"] >= 1 and res["over"] == 0,
   "센 것이 틀렸다: 채움 %r · 덮어씀 %r" % (res.get("filled"), res.get("over")))


# ── 3-2. 덮어쓰기를 고르면 들어간다 ─────────────────────────────────
proj2 = {"models": copy.deepcopy(MODELS)}
res2 = apply_models(proj2, rows, overwrite=True)
got2 = {m["name"]: m for m in proj2["models"]}
ok(got2["A-001"]["price"] == 1250.5,
   "덮어쓰기인데 판가가 안 들어갔다: %r" % got2["A-001"].get("price"))
ok(got2["A-003"]["po_qty"] == 0,
   "덮어쓰기인데 PO 0 이 안 들어갔다: %r" % got2["A-003"].get("po_qty"))
ok(got2["A-003"]["price"] == 900, "덮어쓰기에서 채움이 빠졌다")
ok(got2["A-002"]["price"] == 2000,
   "덮어쓰기라도 빈 칸은 그대로 둬야 한다: %r" % got2["A-002"].get("price"))
ok(res2["over"] >= 2, "덮어쓴 칸을 안 셌다: %r" % res2.get("over"))

# 미리보기에서 '새로 생긴다' 고 한 것과 실제로 생긴 것이 같아야 한다
ok([c["name"] for c in created] == res["created"],
   "미리보기와 저장이 다르다: %s vs %s" % ([c["name"] for c in created], res["created"]))


# ── 4. 미달 사유가 화면과 같은 자리에 ───────────────────────────────
import pbx_plan_import as pbx               # noqa: E402

g = pbx.parse(load_workbook(BytesIO(buf.getvalue()), data_only=True), month=MON)
ok(g is not None, "주차 파서가 양식을 못 읽는다")
notes = (g or {}).get("week_notes") or {}
ok(notes == {WEEKS[0]: "자재 입고 지연", WEEKS[2]: "고객 요청으로 보류"},
   "주차 사유가 다르게 읽혔다: %r" % (notes,))

n = apply_notes(proj, MON, notes)
ok(n == 2, "사유가 %d주차만 들어갔다" % n)
store = (proj.get("week_reasons") or {}).get(MON) or {}
ok(store.get(WEEKS[0], {}).get("text") == "자재 입고 지연",
   "사유가 안 들어갔다: %r" % store.get(WEEKS[0]))
ok(store.get(WEEKS[0], {}).get("kind") == "문제",
   "구분이 '문제' 가 아니다: %r" % store.get(WEEKS[0]))

# 화면에서 '참고' 로 적어 둔 주는 구분을 안 건드린다
proj["week_reasons"][MON][WEEKS[1]] = {"kind": "참고", "text": "일부러 덜 출하", "at": ""}
apply_notes(proj, MON, {WEEKS[1]: "일부러 덜 출하 (매출 조정)"})
ok(proj["week_reasons"][MON][WEEKS[1]]["kind"] == "참고",
   "'참고' 로 적어 둔 주가 '문제' 로 바뀌었다")


# ── 5. 라우트가 이 함수들을 쓰는가 ──────────────────────────────────
for frag, why in [
    ("mrows = _plan_sheet_models(_ws)", "미리보기가 모델 칸을 안 읽는다"),
    ("mres = _plan_apply_models(proj, mrows,\n"
     "                                          overwrite=_plan_over(overwrite))",
     "저장이 모델 값을 안 넣는다"),
    ("overwrite: str = Form(\"\")", "저장 라우트가 덮어쓰기 여부를 안 받는다"),
    ("_plan_apply_week_notes(proj, mon, p.get(\"week_notes\"))",
     "저장이 미달 사유를 안 넣는다"),
    ('out["unmatched"] = [u for u in out.get("unmatched") or []',
     "새로 생길 모델이 '건너뛰는 행' 으로 남는다"),
]:
    ok(frag in SRC, why)

# 모델을 먼저 만들어야 그 줄의 주차 숫자가 붙는다.
# (저장은 이제 파일에 든 달마다 한 바퀴 돈다 — 그 바퀴보다 앞이어야 한다)
ok(SRC.index("mres = _plan_apply_models(proj, mrows,")
   < SRC.index("    for p in parsed_all:\n        diff = _pbx_diff(proj, p)"),
   "모델보다 주차를 먼저 넣는다 — 새 모델의 주차 숫자가 사라진다")

# 달마다 도는가 — 한 달만 넣으면 1년치 파일에서 열두 달이 조용히 버려진다
ok("parsed_all = _pbx_pick_months(wb, month)" in SRC,
   "저장이 파일에 든 달을 전부 읽지 않는다")
ok(SRC.count("parsed_all = _pbx_pick_months(wb, month)") == 2,
   "미리보기와 저장 중 한쪽만 여러 달을 읽는다")


if FAIL:
    print("실패 %d건" % len(FAIL))
    for f in FAIL:
        print("  -", f)
    raise SystemExit(1)
print("전부 통과 · 검사 5묶음 · 모델 %d종 + 신규 1종" % len(MODELS))

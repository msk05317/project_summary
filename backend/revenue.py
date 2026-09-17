# 사업부 매출 — 실적은 주간보고 엑셀, 예상은 따로 받는다.
#
# 처음엔 주간보고의 '실행계획' 을 그 달 계획으로 썼다. 그런데 그 열은
# 작년에 짜 둔 숫자라 지금 예상과 맞지 않는다. 9월만 봐도 주간보고는
# $1,964만, 실제 예상은 $2,422만이었다. 그래서 나눈다.
#
#   실적  ← 주간보고 '1. 계획 대비 실적 (수정본) Actual FCST' 시트
#           줄마다 '실적' 행, 열마다 W1~W52,
#           맨 아래 [합 계] 블록에 보고서 묶음이 그대로
#
#               매출합계 (Revenue)   반도체 / 데이터 센터 / 우주항공 → 소계
#               내부거래 (Internal)  구미/화성/미국
#               총합 (내부거래 포함) = 소계 + 내부거래
#
#   예상  ← 사람이 넣는다. 금액 하나만 넣어도 되고, Estimate 파일을
#           올리면 Commodity 별 내역까지 들어온다.
#
# 주간보고의 사업계획·실행계획·Open PO 열은 읽지 않는다. 읽어서 어딘가
# 남겨 두면 언젠가 화면에 새어 나온다.
#
# 묶음 합계는 [합 계] 블록을 그대로 쓴다. 세부를 우리가 더해서 만들면
# 파일이 말한 값과 어긋날 수 있고, 어긋나도 아무도 모른다. 세부는 따로
# 읽어서 합이 맞는지 대조만 한다.
import datetime as _dt
import io
import json
import re
import threading as _threading

# [합 계] 블록의 줄 이름 → 묶음 키
GROUPS = [("semi", "반도체 (SEMI)"),
          ("dc", "데이터 센터 (Data Center)"),
          ("space", "우주항공 (Space X)")]
INTERNAL = ("internal", "구미/화성/미국 (Gumi/Hwaseong/USA/YONGIN)")
GROUP_LABEL = dict(GROUPS + [INTERNAL])
ORDER = ["semi", "dc", "space", "internal"]

MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")

# 예상을 적을 때 기본으로 깔리는 줄. Estimate 파일 Sum 시트 순서 그대로다.
# 매달 이름을 다시 치게 하면 오타로 같은 항목이 둘로 갈라진다.
DEFAULT_COMMODITIES = [
    "Plastic", "Major Modules", "Sheet metal", "Frame", "Data Center",
    "Metal", "Hwaseong Sheet metal", "Gumi Sheet Metal", "Cable LAM",
    "Cable Internal", "Gumi MCT", "EMA", "Hwaseong MCT", "PBX", "Space X",
]


def norm(s) -> str:
    """줄바꿈·겹공백을 없앤 이름. 엑셀은 같은 줄을 파일마다 다르게 접는다."""
    return " ".join(str(s or "").split())


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace("$", "").strip())
    except Exception:
        return 0.0


def _open(raw: bytes):
    import openpyxl
    try:
        # 글꼴 family 가 14 를 넘는 파일이 있다. openpyxl 이 그대로 죽는다.
        from openpyxl.styles.fonts import Font
        Font.family.max = 1000
    except Exception:
        pass
    return openpyxl.load_workbook(io.BytesIO(raw), data_only=True)


def _sheet(wb):
    for ws in wb.worksheets:
        t = norm(ws.title)
        if "계획 대비 실적" in t and "FCST" in t.upper():
            return ws
    for ws in wb.worksheets:
        if "계획 대비 실적" in norm(ws.title):
            return ws
    return wb.worksheets[0]


def _group_of(div: str, cust: str) -> str:
    """세부 한 줄이 어느 묶음인지.

    엑셀은 사업부 칸을 한 번만 적고 아래로 이어 쓴다. 그래서 텍슨
    사이트 줄들이 '우주항공' 아래에 붙어 있는데, 실제로는 내부거래다.
    """
    d, c = norm(div), norm(cust)
    if "Texon" in c or "Seojin" in c:
        return "internal"
    if d.startswith("반도체"):
        return "semi"
    if "네트워크" in d or "데이터" in c:
        return "dc"
    if "우주항공" in d or "스페이스" in c:
        return "space"
    return "semi"


def parse_weekly(raw: bytes, year_hint: int = None) -> dict:
    """주간보고 시트 한 장에서 주차별 '실적' 만 꺼낸다."""
    wb = _open(raw)
    ws = _sheet(wb)
    grid = list(ws.iter_rows(values_only=True))

    # 주차 머리 줄 찾기 (W1 · W2 … 가 여럿 있는 줄)
    hi = -1
    for i, r in enumerate(grid[:12]):
        got = sum(1 for v in r if re.fullmatch(r"W\d{1,2}", norm(v)))
        if got >= 2:
            hi = i
            break
    if hi < 0:
        raise ValueError("W1 · W2 … 주차 머리 줄을 못 찾았습니다.")

    # 월은 한 줄 위에 한 번만 적혀 있고 아래로 이어진다
    mrow = grid[hi - 1] if hi > 0 else ()
    wcol, month_of, cur = {}, {}, ""
    for j, v in enumerate(grid[hi]):
        lab = norm(mrow[j]) if j < len(mrow) else ""
        if lab and "월" in lab:
            cur = lab
        w = norm(v)
        if re.fullmatch(r"W\d{1,2}", w):
            wcol[w] = j
            month_of[w] = cur
    weeks = sorted(wcol, key=lambda w: int(w[1:]))

    year = year_hint or _dt.date.today().year
    for r in grid[:6]:
        for v in r:
            m = re.search(r"(20\d\d)\s*년", norm(v))
            if m:
                year = int(m.group(1))
                break

    def months():
        out = {}
        for w in weeks:
            lab = month_of.get(w) or ""
            m = re.search(r"(\d{1,2})\s*월", lab)
            if not m:
                continue
            out.setdefault("%04d-%02d" % (year, int(m.group(1))), []).append(w)
        return out

    # ── 줄 읽기 ───────────────────────────────────────────
    def actual_row(i):
        """i 줄에서 시작하는 한 덩이의 '실적' 행. 없으면 None.

        덩이는 '사업계획' 에서 시작해 다음 '사업계획' 전까지다. 그 열은
        읽지 않지만 덩이를 가르는 표시라서 찾기는 해야 한다.
        """
        for k in range(i, min(i + 6, len(grid))):
            lab = norm(grid[k][5] if len(grid[k]) > 5 else "")
            if k > i and lab == "사업계획":
                return None
            if lab == "실적":
                return grid[k]
        return None

    def series(row):
        return {w: _num(row[wcol[w]]) if wcol[w] < len(row) else 0.0
                for w in weeks}

    lines, totals = [], {}
    div = cust = ""
    head_at = -1
    for i, r in enumerate(grid):
        c2 = norm(r[2] if len(r) > 2 else "")
        c3 = norm(r[3] if len(r) > 3 else "")
        c4 = norm(r[4] if len(r) > 4 else "")
        c5 = norm(r[5] if len(r) > 5 else "")
        if c2 and "합" in c2 and "계" in c2 and len(c2) <= 6:
            head_at = i                      # [합 계] 부터는 묶음 합계
        if c2:
            div, cust = c2, ""
        if c3:
            cust = c3
        if c5 != "사업계획":
            continue
        row = actual_row(i)
        if row is None:
            continue
        rec = {"label": c4 or c3 or c2, "actual": series(row)}
        if head_at >= 0 and i > head_at:
            key = _total_key(c2, c3)
            if key:
                totals[key] = rec
            continue
        rec["div"] = div
        rec["cust"] = cust
        rec["group"] = _group_of(div, cust)
        lines.append(rec)

    if not totals:
        raise ValueError("[합 계] 블록을 못 찾았습니다.")
    return {"year": year, "weeks": weeks, "months": months(),
            "totals": totals, "lines": lines}


def _total_key(c2: str, c3: str) -> str:
    """[합 계] 블록의 줄 이름 → 키."""
    t = norm(c3) or norm(c2)
    if t.startswith("반도체"):
        return "semi"
    if "데이터" in t:
        return "dc"
    if "우주항공" in t or "Space X" in t:
        return "space"
    if t.lower().startswith("sub-total") or t == "소계":
        return "subtotal"
    if "구미" in t and "화성" in t:
        return "internal"
    if norm(c2).startswith("총합"):
        return "grand"
    return ""


# ── 예상 매출 ─────────────────────────────────────────────
#
# Estimate 파일은 Commodity 로 적혀 있고 보고서 묶음은 사업부로 나뉜다.
# 이름으로 이어 붙인다. 사이트 이름(구미·화성·용인·USA)이나 Internal 이
# 들어간 것은 내부거래다 — 주간보고에서 텍슨 줄을 내부거래로 보내는 것과
# 같은 규칙이다.
_EST_RULES = (
    ("internal", ("internal", "gumi", "구미", "hwaseong", "화성",
                  "yongin", "용인", "usa", "texon", "텍슨")),
    ("dc", ("data center", "datacenter", "데이터")),
    ("space", ("space x", "spacex", "starlink", "starship", "우주", "스페이스")),
)


def estimate_group(name: str) -> str:
    """Commodity 이름이 어느 묶음인지. 모르면 반도체로 본다."""
    t = norm(name).lower()
    for key, words in _EST_RULES:
        if any(w in t for w in words):
            return key
    return "semi"


def parse_estimate(raw: bytes) -> dict:
    """Estimate 파일에서 그 달 예상 매출을 꺼낸다.

    'Sum' 시트가 `Commodity | 금액` 두 칸이다. 이름이 빈 채로 숫자만
    있는 마지막 줄은 그 시트가 스스로 낸 합계라서 항목으로 세지 않고,
    우리가 더한 값과 맞는지 대조만 한다.
    """
    wb = _open(raw)
    ws = None
    for w in wb.worksheets:
        if norm(w.title).lower() in ("sum", "합계", "총계"):
            ws = w
            break
    if ws is None:
        ws = wb.worksheets[0]

    items, stated = [], None
    for r in ws.iter_rows(values_only=True):
        if not r:
            continue
        name = norm(r[0] if len(r) > 0 else "")
        amt = None
        for v in r[1:]:
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                amt = float(v)
                break
        if amt is None:
            continue
        if name:
            items.append({"item": name, "amount": round(amt),
                          "group": estimate_group(name)})
        else:
            stated = round(amt)             # 시트가 적어 둔 합계
    if not items:
        raise ValueError("Commodity 와 금액이 있는 줄을 못 찾았습니다.")
    total = sum(x["amount"] for x in items)
    return {"items": items, "total": total, "stated": stated,
            "groups": _by_group(items),
            "gap": None if stated is None else stated - total}


def _by_group(items) -> dict:
    out = {k: 0 for k in ORDER}
    for x in (items or []):
        g = x.get("group") or "semi"
        out[g] = out.get(g, 0) + round(_num(x.get("amount")))
    return out


def set_estimate(store: dict, month: str, items=None, source: str = "",
                 total=None) -> dict:
    """그 달 예상 매출을 넣는다. 다 더한 값이 0 이면 그 달을 지운다.

    합계는 항목을 더해서 낸다 — 따로 받으면 둘이 어긋났을 때 어느 쪽이
    맞는지 알 수 없다. 항목이 아예 없을 때만 total 을 쓴다 (묶음별
    예상이 없는 달이 되고, 그때는 부서별 달성률을 지어내지 않는다).
    """
    if not MONTH_RE.match(str(month or "")):
        raise ValueError("달은 2026-09 모양이어야 합니다.")
    est = store.setdefault("estimates", {})
    rows = []
    for x in (items or []):
        name = norm(x.get("item"))
        amt = round(_num(x.get("amount")))
        if not name or amt <= 0:
            continue                       # 이름만 있고 안 채운 줄은 버린다
        rows.append({"item": name, "amount": amt,
                     "group": x.get("group") or estimate_group(name)})
    amount = sum(r["amount"] for r in rows) if rows else round(_num(total))
    if amount <= 0:
        est.pop(month, None)
        return {"month": month, "total": 0, "items": 0, "removed": True}
    rec = {"total": amount,
           "items": rows,
           "groups": _by_group(rows) if rows else {},
           "source": source or "직접 입력",
           "at": _dt.datetime.now().isoformat(timespec="seconds")}
    est[month] = rec
    return {"month": month, "total": amount, "items": len(rows),
            "groups": rec["groups"], "removed": False}


def _estimate_of(store: dict, month: str) -> dict:
    return ((store.get("estimates") or {}).get(month) or {})


# ── 저장 ──────────────────────────────────────────────────
def blank() -> dict:
    return {"version": 3, "updated_at": None, "years": {},
            "estimates": {}, "sources": {}}


class RevenueFileBroken(Exception):
    """파일이 있는데 못 읽는다. 빈 값으로 갈아엎으면 안 된다."""


def load(path) -> dict:
    """없으면 빈 것, 깨졌으면 예외.

    깨진 걸 빈 것으로 돌려주면 다음 저장이 멀쩡한 파일을 빈 값으로
    덮어쓴다. 한 번 그러면 되돌릴 방법이 없다.
    """
    import os
    if not os.path.exists(path):
        return blank()
    try:
        with open(path, "r", encoding="utf-8") as f:
            d = json.load(f)
    except Exception as e:
        raise RevenueFileBroken(str(e))
    if not isinstance(d, dict):
        raise RevenueFileBroken("최상위가 객체가 아니다")
    base = blank()
    for k, v in base.items():
        d.setdefault(k, v)
    _drop_plan_columns(d)
    return d


def _drop_plan_columns(store: dict) -> None:
    """version 2 까지는 사업계획·실행계획도 같이 담겨 있었다.

    남겨 두면 언젠가 화면에 새어 나온다. 읽을 때 털어낸다.
    """
    if store.get("version") == 3:
        return
    for y in (store.get("years") or {}).values():
        for rec in (y.get("totals") or {}).values():
            rec.pop("plan", None)
            rec.pop("budget", None)
        for rec in (y.get("lines") or []):
            rec.pop("plan", None)
            rec.pop("budget", None)
    store["version"] = 3


_SAVE_LOCK = _threading.Lock()


def save(path, store: dict) -> None:
    """같은 이름의 임시 파일을 여럿이 동시에 쓰면 서로를 덮어쓴다."""
    import os
    store["updated_at"] = _dt.datetime.now().isoformat(timespec="seconds")
    tmp = "%s.%d.%d.tmp" % (path, os.getpid(), _threading.get_ident())
    with _SAVE_LOCK:
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(store, f, ensure_ascii=False, indent=1)
        os.replace(tmp, path)


def apply_weekly(store: dict, parsed: dict, filename: str = "") -> dict:
    """그 해를 통째로 갈아 끼운다. 파일 한 장이 한 해 전체다.

    예상은 이 파일에서 오지 않으므로 건드리지 않는다.
    """
    y = str(parsed.get("year"))
    before = _grand(store, y)
    store.setdefault("years", {})[y] = {
        "weeks": parsed.get("weeks") or [],
        "months": parsed.get("months") or {},
        "totals": parsed.get("totals") or {},
        "lines": parsed.get("lines") or [],
    }
    store.setdefault("sources", {})["weekly"] = {
        "file": filename,
        "at": _dt.datetime.now().isoformat(timespec="seconds"),
        "year": y,
        "weeks": len(parsed.get("weeks") or []),
        "lines": len(parsed.get("lines") or []),
    }
    return {"year": y, "before": before, "after": _grand(store, y)}


def _grand(store: dict, year: str) -> int:
    got = ((store.get("years") or {}).get(str(year)) or {})
    rec = (got.get("totals") or {}).get("grand") or {}
    return round(sum((rec.get("actual") or {}).values()))


def _months_of(y: dict, month: str) -> list:
    return list((y.get("months") or {}).get(month) or [])


def _sum(rec: dict, key: str, weeks) -> int:
    got = (rec or {}).get(key) or {}
    return round(sum(float(got.get(w) or 0) for w in weeks))


def month_view(store: dict, month: str) -> dict:
    """그 달의 묶음별 실적 + 사람이 넣은 예상. 홈·상세·admin 이 같이 쓴다."""
    year = str(month)[:4]
    y = (store.get("years") or {}).get(year) or {}
    weeks = _months_of(y, month)
    totals = y.get("totals") or {}
    lines = y.get("lines") or []

    est = _estimate_of(store, month)
    est_g = est.get("groups") or {}

    def box(key, label):
        rec = totals.get(key) or {}
        a = _sum(rec, "actual", weeks)
        e = round(_num(est_g.get(key)))
        items = []
        for ln in lines:
            if ln.get("group") != key:
                continue
            items.append({"item": ln.get("label") or "",
                          "cust": ln.get("cust") or "",
                          "group": key,
                          "actual": _sum(ln, "actual", weeks)})
        items.sort(key=lambda r: (-r["actual"], r["item"]))
        return {"key": key, "label": label, "actual": a,
                "estimate": e,
                "rate": round(a * 100 / e) if e > 0 else None,
                "items": items,
                # 세부를 더한 값. 합계와 다르면 화면에서 말해 준다.
                "items_actual": sum(r["actual"] for r in items)}

    boxes = [box(k, lb) for k, lb in GROUPS]
    internal = box(*INTERNAL)

    def pack(rec, label, fallback, est_keys):
        a = _sum(rec, "actual", weeks)
        if not a and fallback:
            a = fallback
        e = sum(round(_num(est_g.get(k))) for k in est_keys)
        return {"label": label, "actual": a, "estimate": e,
                "rate": round(a * 100 / e) if e > 0 else None}

    subtotal = pack(totals.get("subtotal") or {}, "소계 (Sub-total)",
                    sum(x["actual"] for x in boxes),
                    ("semi", "dc", "space"))
    grand = totals.get("grand") or {}
    grand_v = pack(grand, "총합 (내부거래 포함)",
                   subtotal["actual"] + internal["actual"], ORDER)
    actual = grand_v["actual"]

    # 총합 예상은 사람이 넣은 값이 정본이다. 묶음별은 그 안을 나눈 것뿐.
    estimate = round(_num(est.get("total")))
    grand_v["estimate"] = estimate
    grand_v["rate"] = round(actual * 100 / estimate) if estimate > 0 else None
    rate = grand_v["rate"]

    # 누적은 보고 있는 달까지 (뒤에 오는 달이 들어가면 1~8월 합과 안 맞는다)
    upto = []
    for m, ws in sorted((y.get("months") or {}).items()):
        if m <= month:
            upto.extend(ws)
    ytd = _sum(grand, "actual", upto) or sum(
        _sum(totals.get(k) or {}, "actual", upto) for k in ORDER)

    last = ""
    for w in weeks:
        if _sum(grand or (totals.get("semi") or {}), "actual", [w]) > 0:
            last = w
    return {
        "month": month, "weeks": weeks,
        "groups": boxes, "internal": internal,
        "subtotal": subtotal, "grand": grand_v,
        "actual": actual,
        "estimate": estimate,
        "estimate_items": est.get("items") or [],
        # 아직 안 넣은 달이면 화면이 이 이름들로 빈 줄을 깔아 준다
        "commodities": DEFAULT_COMMODITIES,
        "estimate_groups": est_g,
        # 묶음별로 나뉘어 들어왔는지. 금액만 넣은 달은 꺼진다.
        "has_estimate_groups": bool(est_g),
        "estimate_at": est.get("at") or "",
        "estimate_source": est.get("source") or "",
        "rate": rate,
        "left": max(0, estimate - actual) if estimate > 0 else 0,
        "ytd": ytd, "as_of": last,
        "has_data": bool(weeks) and bool(totals),
        "has_estimate": estimate > 0,
        "sources": store.get("sources") or {},
    }

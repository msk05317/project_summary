# -*- coding: utf-8 -*-
"""작성용 양식 — 한 해 한 장, 적는 달만 펼쳐 둔다.

### 왜 한 장인가

달마다 시트 한 장짜리 양식(weekly_template)은 이번 달만 보인다. 그런데
이번 주 계획을 적는 사람은 보통 "지난 달까지 얼마나 나갔더라" 를 보면서
적는다. 그걸 보려고 시트를 왔다 갔다 하거나 보드를 따로 띄워야 했다.

그래서 한 해를 한 장에 놓는다. 열두 달 주차가 **전부 들어 있고**, 적는
달만 펼쳐져 있다. 나머지는 접혀서 월 합계 두 칸만 보인다 — 지워진 게
아니라 접힌 것이라, 엑셀에서 열 머리 위 [+] 를 누르면 그 자리에서
펼쳐지고 거기 적어 올리면 그 달도 들어간다.

### 접힌 달도 올라간다

올리는 쪽은 원래 시트 하나를 달 하나로 읽었다. 시트 이름이 '2026-10'
이면 10월 주차만 읽고 나머지는 버렸다. 그대로 두면 9월을 펴서 적어
올렸을 때 조용히 사라진다 — 적을 수 있게 생겼는데 안 먹는 칸이 제일
나쁘다. pbx_plan_import.sheet_months_of() 가 '이 시트는 한 해를 담고
있다' 를 알아보고 열두 달을 따로 읽게 고쳤다.

### 적는 달 표시

빨간 **테두리**다. 칸을 빨갛게 칠하면 '여기가 문제다' 로 읽힌다 —
보드에서 빨강은 이미 '계획에 못 미쳤다' 라서.

### 보고용(plan_export)과 다르다

보고용은 진행현황에 붙이는 것이라 유형별 묶음 합계가 있고 판가·재료비가
없다. 이건 올리는 양식이라 반대다. 묶음 합계 줄도 없다 — 읽는 쪽이
'합계' 글자를 만나면 표가 끝난 줄 알아서, 중간에 두면 거기서 읽기가
끊긴다. 합계는 맨 아래 한 줄뿐이다.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import week_calendar as _wc
import weekly_template as _wt

NAVY = _wt.NAVY
HEAD2 = "1E3A5F"
BLUE = _wt.BLUE        # 계획
WARM = _wt.WARM        # 실적
GREY = _wt.GREY        # 수식
LOCK = _wt.LOCK        # 누구인지 — 건드리지 않는 칸
NOTEBG = _wt.NOTEBG
SUMBG = _wt.SUMBG
KFONT = _wt.KFONT
BORDER = _wt.BORDER
RED = "C00000"         # 적는 달 테두리
DASH = "-"

ROW_TITLE = 1
HDR_TOP = 2            # 월 이름
HDR_WEEK = 3           # 주차 · 월 합계
HDR_SUB = 4            # 계획/실적
FIRST_ROW = 5
SPARE_ROWS = 8         # 새 모델 적을 빈 줄
SUB = ("계획", "실적")


def layout(year: int, month: str) -> list:
    """달마다 자리 → [{'month','weeks','open'}, ...]

    열두 달이 **전부** 주차를 갖는다. open 은 펼쳐 둘 달 하나다.
    """
    out = []
    for i in range(1, 13):
        ym = "%04d-%02d" % (int(year), i)
        out.append({"month": ym,
                    "weeks": list(_wc.get_month_weeks(ym) or []),
                    "open": ym == str(month)})
    return out


def _head(ws, r, c, text, fill=NAVY, size=10):
    x = ws.cell(r, c, text)
    x.font = Font(name=KFONT, bold=True, size=size, color="FFFFFF")
    x.fill = PatternFill("solid", fgColor=fill)
    x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    x.border = BORDER
    return x


def _box(ws, r0, c0, r1, c1, color=RED):
    """바깥 테두리만 굵게. 안쪽 줄은 그대로 둔다."""
    thick = Side(style="medium", color=color)
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            cell = ws.cell(r, c)
            cur = cell.border
            try:
                cell.border = Border(
                    left=thick if c == c0 else cur.left,
                    right=thick if c == c1 else cur.right,
                    top=thick if r == r0 else cur.top,
                    bottom=thick if r == r1 else cur.bottom)
            except AttributeError:
                pass           # 병합된 칸의 자투리 — 왼쪽 위 칸이 대표한다


def build_year_form(project_label: str, year: int, month: str, models: list,
                    notes: dict = None) -> bytes:
    """작성용 한 해 양식 바이트.

    month  펼쳐 둘 달 ('2026-10'). 시트 이름도 이것이다.
    notes  {'2026-10': {'W41': '자재 지연'}} — 달마다 미달 사유.
    """
    year = int(year)
    month = str(month).strip()
    if not _wc.get_month_weeks(month):
        raise ValueError("%s 는 주차가 없는 달입니다" % month)
    if not month.startswith("%04d-" % year):
        raise ValueError("%s 는 %d년이 아닙니다" % (month, year))
    notes = notes or {}

    models = _wt.order_models(models or [])
    cols = _wt.fixed_cols(models)
    nid = len(cols)
    col_left = nid + len(_wt.VALUE) + 1        # 잔량
    nlead = col_left

    blocks = layout(year, month)
    c = nlead + 1
    for b in blocks:
        b["cols"] = [c + i * 2 for i in range(len(b["weeks"]))]
        b["total"] = c + len(b["weeks"]) * 2
        c = b["total"] + 2
    last = c - 1
    cur = next(b for b in blocks if b["open"])

    wb = Workbook()
    ws = wb.active
    ws.title = month

    body = Font(name=KFONT, size=10)
    lock_f = PatternFill("solid", fgColor=LOCK)
    plan_f = PatternFill("solid", fgColor=BLUE)
    act_f = PatternFill("solid", fgColor=WARM)
    calc_f = PatternFill("solid", fgColor=GREY)

    # ── 제목 ─────────────────────────────────────────────────────────
    ws.merge_cells(start_row=ROW_TITLE, start_column=1,
                   end_row=ROW_TITLE, end_column=last)
    t = ws.cell(ROW_TITLE, 1,
                (("%s · " % project_label) if project_label else "")
                + "%d년 주차별 계획 · 실적    " % year
                + "[ 빨간 테두리가 %s 입니다. 다른 달은 접혀 있고, "
                  "열 머리 위 [+] 를 누르면 그 달 주차가 펼쳐집니다 — "
                  "거기 적어 올려도 들어갑니다 ]" % month)
    t.font = Font(name=KFONT, bold=True, size=11, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[ROW_TITLE].height = 26

    # ── 왼쪽 머리글 ──────────────────────────────────────────────────
    ci = 1
    for head, width, _key in cols:
        ws.merge_cells(start_row=HDR_TOP, start_column=ci,
                       end_row=HDR_SUB, end_column=ci)
        _head(ws, HDR_TOP, ci, head)
        ws.column_dimensions[get_column_letter(ci)].width = width
        ci += 1
    for head, width, _key, _fmt in _wt.VALUE:
        ws.merge_cells(start_row=HDR_TOP, start_column=ci,
                       end_row=HDR_SUB, end_column=ci)
        _head(ws, HDR_TOP, ci, head, fill=HEAD2)
        ws.column_dimensions[get_column_letter(ci)].width = width
        ci += 1
    ws.merge_cells(start_row=HDR_TOP, start_column=col_left,
                   end_row=HDR_SUB, end_column=col_left)
    _head(ws, HDR_TOP, col_left, "잔량", fill=HEAD2)
    ws.column_dimensions[get_column_letter(col_left)].width = 10

    # ── 달 머리글 ────────────────────────────────────────────────────
    for b in blocks:
        c0 = b["cols"][0] if b["cols"] else b["total"]
        ws.merge_cells(start_row=HDR_TOP, start_column=c0,
                       end_row=HDR_TOP, end_column=b["total"] + 1)
        _head(ws, HDR_TOP, c0,
              "%d월%s" % (int(b["month"][5:7]), "   ← 적는 달" if b["open"] else ""))
        for wk, wc in zip(b["weeks"], b["cols"]):
            ws.merge_cells(start_row=HDR_WEEK, start_column=wc,
                           end_row=HDR_WEEK, end_column=wc + 1)
            _head(ws, HDR_WEEK, wc, wk, fill=HEAD2)
            for j, s in enumerate(SUB):
                _head(ws, HDR_SUB, wc + j, s, fill=HEAD2)
                ws.column_dimensions[get_column_letter(wc + j)].width = 8
        ws.merge_cells(start_row=HDR_WEEK, start_column=b["total"],
                       end_row=HDR_WEEK, end_column=b["total"] + 1)
        _head(ws, HDR_WEEK, b["total"], "월 합계", fill=HEAD2)
        for j, s in enumerate(SUB):
            _head(ws, HDR_SUB, b["total"] + j, s, fill=HEAD2)
            ws.column_dimensions[get_column_letter(b["total"] + j)].width = 9

    ws.row_dimensions[HDR_TOP].height = 24
    ws.row_dimensions[HDR_WEEK].height = 18
    ws.row_dimensions[HDR_SUB].height = 18

    # ── 줄 ───────────────────────────────────────────────────────────
    L = get_column_letter
    r = FIRST_ROW
    for m in models:
        ci = 1
        for v in _wt._model_row(m, cols):
            x = ws.cell(r, ci, v)
            x.font = body
            x.fill = lock_f
            x.border = BORDER
            x.alignment = Alignment(horizontal="left", vertical="center")
            ci += 1
        for v, (_h, _w, _k, fmt) in zip(_wt._value_row(m), _wt.VALUE):
            x = ws.cell(r, ci, v)
            x.font = body
            x.border = BORDER
            x.number_format = fmt
            x.alignment = Alignment(horizontal="right", vertical="center")
            ci += 1

        # 잔량 — 한쪽이 비어 있어도 안 깨지게 SUM 으로 뺀다
        cpo, cshp = L(nid + 3), L(nid + 4)
        x = ws.cell(r, col_left,
                    '=IF(COUNT(%s%d:%s%d)=0,"",SUM(%s%d)-SUM(%s%d))'
                    % (cpo, r, cshp, r, cpo, r, cshp, r))
        x.font = Font(name=KFONT, size=10, bold=True, color=NAVY)
        x.fill = calc_f
        x.border = BORDER
        x.number_format = "#,##0"
        x.alignment = Alignment(horizontal="right", vertical="center")

        for b in blocks:
            wk_vals = _wt.week_cells(m, b["month"])
            for wk, wc in zip(b["weeks"], b["cols"]):
                got = wk_vals.get(wk) or {}
                for j, key in enumerate(("plan", "actual")):
                    x = ws.cell(r, wc + j, got.get(key))
                    x.font = body
                    x.fill = plan_f if j == 0 else act_f
                    x.border = BORDER
                    x.number_format = "#,##0"
                    x.alignment = Alignment(horizontal="right", vertical="center")
            for j in range(2):
                refs = ",".join("%s%d" % (L(wc + j), r) for wc in b["cols"])
                x = ws.cell(r, b["total"] + j,
                            '=IF(COUNT(%s)=0,"",SUM(%s))' % (refs, refs))
                x.font = Font(name=KFONT, size=10, bold=True, color=NAVY)
                x.fill = calc_f
                x.border = BORDER
                x.number_format = "#,##0"
                x.alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[r].height = 17
        r += 1

    # 새 모델 적을 빈 줄
    for _ in range(SPARE_ROWS):
        for ci2 in range(1, last + 1):
            x = ws.cell(r, ci2)
            x.font = body
            x.border = BORDER
            x.alignment = Alignment(
                horizontal="left" if ci2 <= nid else "right", vertical="center")
        ws.row_dimensions[r].height = 17
        r += 1
    r_end = r - 1

    # ── 합계 · 미달 사유 ─────────────────────────────────────────────
    #
    # 묶음(유형별) 합계는 두지 않는다. 읽는 쪽이 '합계' 글자를 만나면 표가
    # 끝난 줄 알아서, 중간에 두면 그 아래 줄이 통째로 안 읽힌다.
    r_sum = r
    sum_f = PatternFill("solid", fgColor=SUMBG)
    sum_font = Font(name=KFONT, size=10, bold=True, color="FFFFFF")
    for ci2 in range(1, last + 1):
        x = ws.cell(r_sum, ci2)
        x.font = sum_font
        x.fill = sum_f
        x.border = BORDER
        x.number_format = "#,##0"
        x.alignment = Alignment(horizontal="right", vertical="center")
    lab = ws.cell(r_sum, 1, "합계")
    lab.alignment = Alignment(horizontal="left", vertical="center")
    for ci2 in range(nid + 3, last + 1):      # 판가·재료비는 더하지 않는다
        c2 = L(ci2)
        ws.cell(r_sum, ci2, '=IF(COUNT(%s%d:%s%d)=0,"",SUM(%s%d:%s%d))'
                % (c2, FIRST_ROW, c2, r_end, c2, FIRST_ROW, c2, r_end))
    ws.row_dimensions[r_sum].height = 20

    r_note = r_sum + 1
    note_f = PatternFill("solid", fgColor=NOTEBG)
    for ci2 in range(1, last + 1):
        x = ws.cell(r_note, ci2)
        x.font = body
        x.fill = note_f
        x.border = BORDER
        x.alignment = Alignment(horizontal="left", vertical="center",
                                wrap_text=True)
    nl = ws.cell(r_note, 1, "미달 사유")
    nl.font = Font(name=KFONT, size=10, bold=True, color=NAVY)
    if nlead >= 2:
        ws.merge_cells(start_row=r_note, start_column=1,
                       end_row=r_note, end_column=nlead)
    for b in blocks:
        mn = notes.get(b["month"]) or {}
        for wk, wc in zip(b["weeks"], b["cols"]):
            ws.merge_cells(start_row=r_note, start_column=wc,
                           end_row=r_note, end_column=wc + 1)
            txt = str(mn.get(wk) or "").strip()
            if txt:
                ws.cell(r_note, wc, txt)
        ws.merge_cells(start_row=r_note, start_column=b["total"],
                       end_row=r_note, end_column=b["total"] + 1)
    x = ws.cell(r_note, cur["total"], "— 계획을 못 채운 주에만")
    x.font = Font(name=KFONT, size=9, color="94A3B8")
    ws.row_dimensions[r_note].height = 32

    # ── 적는 달 테두리 ───────────────────────────────────────────────
    #
    # 칸을 빨갛게 칠하지 않는다. 보드에서 빨강은 '계획에 못 미쳤다' 라
    # 색을 칠하면 '여기가 문제다' 로 읽힌다.
    _box(ws, HDR_TOP, cur["cols"][0], r_note, cur["total"] + 1)

    # ── 안 적는 달은 접는다 ──────────────────────────────────────────
    #
    # 월 합계 두 칸은 남긴다. 접힌 달도 '그 달 얼마' 는 보이고, 궁금하면
    # [+] 로 펼쳐 거기 적어 올려도 들어간다.
    for b in blocks:
        if b["open"] or not b["cols"]:
            continue
        ws.column_dimensions.group(L(b["cols"][0]), L(b["cols"][-1] + 1),
                                   hidden=True, outline_level=1)

    # 구분은 새 줄을 적을 때 고르는 칸이다
    gi = next((i for i, c2 in enumerate(cols, start=1) if c2[2] == "group"), None)
    if gi:
        dv = DataValidation(type="list", formula1='"양산,개발"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("%s%d:%s%d" % (L(gi), FIRST_ROW, L(gi), r_end))

    ws.freeze_panes = ws.cell(FIRST_ROW, nlead + 1)
    ws.sheet_view.showGridLines = False
    # 인쇄 — 세로로는 한 장, 가로로는 필요한 만큼. 가로까지 한 장에
    # 맞추면 달을 다 펼쳤을 때 백서른 칸이 뭉개진다.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToHeight = 1
    ws.page_setup.fitToWidth = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "%d:%d" % (HDR_TOP, HDR_SUB)
    ws.print_title_cols = "%s:%s" % (L(1), L(nlead))

    _guide(wb, project_label, year, month)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _guide(wb, label, year, month):
    g = wb.create_sheet("작성 안내")
    g.column_dimensions["A"].width = 20
    g.column_dimensions["B"].width = 86
    g.merge_cells("A1:B1")
    t = g.cell(1, 1, "%s%d년 작성용 — 어디에 적나"
               % (("%s · " % label) if label else "", year))
    t.font = Font(name=KFONT, bold=True, size=12, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    g.row_dimensions[1].height = 26

    rows = [
        ("시트", "'%s' 한 장입니다. 한 해 열두 달이 가로로 늘어서 있습니다." % month),
        ("빨간 테두리", "%s 입니다. 보통 여기만 적으시면 됩니다." % month),
        ("다른 달", "접혀 있고 월 합계 두 칸만 보입니다. 열 머리 위 [+] 를 누르면 "
                 "그 달 주차가 펼쳐지고, 거기 적어 올려도 들어갑니다."),
        ("적는 칸", "주차 칸(계획·실적)과 왼쪽의 판가·재료비·PO수량·실적수량입니다."),
        ("빈 칸", "한 줄에 숫자가 하나라도 있으면 그 줄의 **그 달** 주차 칸은 전부 "
                "파일대로 들어갑니다. 비운 칸은 0 이 됩니다. 한 줄을 통째로 "
                "비워 두면 그 줄은 손대지 않습니다."),
        ("판가·재료비·PO", "비어 있던 칸만 채웁니다. 이미 값이 있는 칸은 그대로 둡니다 — "
                       "덮어쓰려면 올릴 때 미리보기에서 고르세요."),
        ("새 모델", "맨 아래 빈 줄에 적으면 새로 만들어집니다. 모델명과 구분(양산/개발)은 "
                 "꼭 적어 주세요."),
        ("미달 사유", "맨 아래 줄입니다. 계획을 못 채운 주에만 그 주 칸에 적으시면 됩니다."),
        ("합계 줄 · 잔량", "자동 계산입니다. 건드리지 마세요."),
    ]
    r = 3
    for k, v in rows:
        a = g.cell(r, 1, k)
        a.font = Font(name=KFONT, bold=True, size=10, color=NAVY)
        a.alignment = Alignment(horizontal="left", vertical="top")
        b = g.cell(r, 2, v)
        b.font = Font(name=KFONT, size=10)
        b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        g.row_dimensions[r].height = 30
        r += 1
    g.sheet_view.showGridLines = False
    return g

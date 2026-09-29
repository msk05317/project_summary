"""'월별 매입 현황 분석' 엑셀 읽기.

시트 모양이 달마다 조금씩 달라서 좌표를 박지 않는다. 대신
  1) 월이 적힌 머리줄을 찾아 '어느 열이 몇 월인지' 를 만들고
  2) 왼쪽 라벨을 보며 줄을 해석한다.

라벨은 보통 세 칸이다 — [매입] [원소재] [매입액]. 마지막 칸이 '무슨
값인지'(매입액·지급액·잔액), 그 앞 칸이 '어느 항목인지'(원소재·소모품…).
항목 이름은 미리 정해 두지 않는다. 실제 파일에는 '조립자재' 처럼 우리가
모르는 항목이 있고, 모르는 항목을 앞 항목에 밀어 넣으면 원소재 숫자가
0 으로 덮인다.

'합계' 줄은 값을 읽지 않는다 — 우리가 더해서 확인하는 쪽이 낫다.
못 알아본 줄은 버리지 않고 warnings 로 올려 보낸다. 조용히 빠진 줄이
제일 위험하다.

월 오른쪽의 '합계 · 해외 · 국내' 세 열은 누적이다. 해외/국내는 달별로
안 나눠져 있고 거기에만 있다. 그래서 달에 섞지 않고 summary 로 따로
낸다. 그 오른쪽 '비고' 는 줄마다의 비고로, 어느 줄에 달린 것인지까지
같이 낸다 (revenue · item:raw:buy · stock:wip …). 비고보다 더 오른쪽에
적어 둔 메모('추가 분석 필요' 목록 같은 것)는 누적 비고에 모아 둔다 —
버리면 적어 둔 사람만 손해다.

재고·외상매출·총매입액 줄은 이름으로 알아본다. 그 블록의 '구분' 칸이
'합계' 라서, 구분 칸만 보고 건너뛰면 통째로 빠진다.
"""
from __future__ import annotations

import datetime as _dt
import io
import os
import re
import tempfile

from openpyxl import load_workbook

# 이름이 이렇게 적혀 있으면 이 키로 맞춰 준다. 여기 없는 항목도 그대로
# 받는다 (키는 g1, g2 … 로 만든다).
ITEM_KEYS = [
    ("raw", ("원소재", "원자재")),
    ("supply", ("소모품", "소모자재")),
    ("etc", ("기타", "약품", "포장재")),
    ("outsourcing", ("외주", "외주가공")),
]
METRICS = {
    "buy": ("매입액", "매입", "구매액", "구매"),
    "paid": ("지급액", "지급", "결제액"),
    "balance": ("잔액", "미지급", "잔여"),
}
PAYMENT = {
    "prepaid": ("선급금", "선급"),
    "actual_paid": ("실지급액", "실지급"),
    "invest": ("설비투자", "실비투자", "설비", "실비"),
}
REVENUE = ("매출액", "매출")
# 재고·미수금·총매입액 줄. 이름이 또렷해서 '구분' 칸을 안 봐도 알아본다 —
# 아래쪽 블록 이름이 '합계' 라서 구분 칸만 보면 통째로 건너뛰게 된다.
STOCK_KEYS = (
    ("available", ("가용원재고", "가용원자재", "가용재고")),
    ("dead", ("불용원자재", "불용원재고", "불용재고")),
    ("wip", ("재공재고", "재공품재고", "재공")),
    ("finished", ("완제품재고", "완제품")),
)
SIDE_KEYS = (
    ("receivable", ("외상매출", "미수금")),
    ("total_buy", ("총매입액", "총매입")),
    ("ratio", ("매출대비매입비율", "매입비율", "매출대비")),
)
REGION_WORDS = {"total": ("합계", "총계"), "overseas": ("해외",), "domestic": ("국내",)}
NOTE_WORDS = ("비고", "특이사항", "note")
MAX_EXTRA = 24
SKIP_IN = ("합계", "총계", "소계")      # 어디에 끼어 있어도 건너뛴다
SKIP_EQ = ("계", "total")                # 딱 그 글자일 때만

_MONTH_TXT = re.compile(r"(?:^|\D)(1[0-2]|0?[1-9])\s*월")


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return "%d월" % v.month
    return re.sub(r"\s+", " ", str(v).replace(" ", " ")).strip()


def _norm(s) -> str:
    return re.sub(r"[\s()（）\[\]·.,/]+", "", _txt(s)).lower()


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = _txt(v).replace(",", "").replace("$", "").replace("₩", "")
    if not t or t in ("-", "—", "–"):
        return None
    neg = False
    if t.startswith("(") and t.endswith(")"):
        neg, t = True, t[1:-1].strip()
    elif t.startswith("-"):
        neg, t = True, t[1:].strip()
    try:
        f = float(t)
    except ValueError:
        return None
    return -f if neg else f


def _is_skip(s) -> bool:
    n = _norm(s)
    if not n:
        return False
    if any(_norm(w) in n for w in SKIP_IN):
        return True
    return n in {_norm(k) for k in SKIP_EQ}


def _month_of(v, year: int):
    if isinstance(v, (_dt.datetime, _dt.date)):
        return "%04d-%02d" % (v.year, v.month)
    t = _txt(v)
    if not t or _is_skip(t):
        return None
    m = re.match(r"^(20\d\d)[-./\s]+(1[0-2]|0?[1-9])$", t)
    if m:
        return "%04d-%02d" % (int(m.group(1)), int(m.group(2)))
    m = _MONTH_TXT.search(t)
    if m:
        return "%04d-%02d" % (year, int(m.group(1)))
    return None


def _known_key(label):
    n = _norm(label)
    if not n:
        return None
    for key, words in ITEM_KEYS:
        if any(_norm(w) in n for w in words):
            return key
    return None


def _metric(label):
    n = _norm(label)
    if not n:
        return None
    for key, words in METRICS.items():
        if any(_norm(w) == n for w in words):
            return key
    for key, words in METRICS.items():
        if any(_norm(w) in n for w in words):
            return key
    return None


def _match(label, words):
    n = _norm(label)
    return bool(n) and any(_norm(w) in n for w in words)


def _special(label):
    """재고·미수금·총매입액·비율 줄인지. ('stock', 키) 또는 (키, None)."""
    n = _norm(label)
    if not n:
        return None
    for key, words in STOCK_KEYS:
        if any(_norm(w) in n for w in words):
            return ("stock", key)
    for key, words in SIDE_KEYS:
        if any(_norm(w) in n for w in words):
            return (key, None)
    return None


def _note_txt(v) -> str:
    """비고는 줄바꿈을 지킨다 — 몇 건인지 눈으로 세야 하니까."""
    if v is None:
        return ""
    t = str(v).replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    lines = [ln.rstrip() for ln in t.split("\n")]
    while lines and not lines[0]:
        lines.pop(0)
    while lines and not lines[-1]:
        lines.pop()
    return "\n".join(lines)


def _open_book(data):
    """bytes · 파일객체 · 경로를 다 받는다.

    서버에서 BytesIO 로 열다가 TypeError 가 난 적이 있어(환경 차이),
    실패하면 임시 파일로 한 번 더 해 본다. 다른 화면들이 경로로 여는
    방식은 그 서버에서 잘 돈다.
    """
    if isinstance(data, (bytes, bytearray, memoryview)):
        raw = bytes(data)
    elif hasattr(data, "read"):
        raw = data.read()
        if not isinstance(raw, (bytes, bytearray)):
            raise ValueError("파일을 바이트로 읽지 못했습니다.")
        raw = bytes(raw)
    else:
        return load_workbook(data, data_only=True)

    try:
        return load_workbook(io.BytesIO(raw), data_only=True)
    except Exception:
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        try:
            with os.fdopen(fd, "wb") as f:
                f.write(raw)
            return load_workbook(path, data_only=True)
        finally:
            try:
                os.unlink(path)
            except OSError:
                pass


def _grid(ws, max_row=400, max_col=40):
    out = []
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, max_row),
                          max_col=min(ws.max_column or 1, max_col)):
        out.append([c.value for c in r])
    return out


def _header(rows, year: int):
    """월이 2개 이상 있는 줄을 머리줄로 본다."""
    best, best_i = {}, -1
    for i, row in enumerate(rows[:40]):
        cols, seen = {}, set()
        for j, v in enumerate(row):
            m = _month_of(v, year)
            if m and m not in seen:
                cols[j] = m
                seen.add(m)
        if len(cols) >= 2 and len(cols) > len(best):
            best, best_i = cols, i
    return best_i, best


def _side_cols(row, after: int):
    """머리줄에서 월 오른쪽에 있는 합계·해외·국내·비고 열을 찾는다.

    이 엑셀은 해외/국내를 달마다 안 나누고 4~8월 누적에만 적는다. 그 세
    열이 '누적' 한 칸이 되고, 그 오른쪽 비고가 줄마다의 비고다.
    """
    regions, note = {}, None
    want = {r: {_norm(w) for w in words} for r, words in REGION_WORDS.items()}
    for j in range(after + 1, len(row)):
        n = _norm(row[j])
        if not n:
            continue
        if note is None and any(_norm(w) in n for w in NOTE_WORDS):
            note = j
            continue
        if note is not None:
            continue
        for r, ns in want.items():
            if n in ns and r not in regions:
                regions[r] = j
    return regions, note


def parse_workbook(data, year=None, sheet=None) -> dict:
    year = int(year or _dt.date.today().year)
    wb = _open_book(data)
    ws = wb[sheet] if (sheet and sheet in wb.sheetnames) else wb[wb.sheetnames[0]]

    rows = _grid(ws)
    hi, cols = _header(rows, year)
    if not cols:
        raise ValueError("월(4월·5월…)이 적힌 머리줄을 찾지 못했습니다.")

    regions, note_col = _side_cols(rows[hi], max(cols))
    months_keys = sorted(set(cols.values()))

    blank = lambda: {"revenue": None, "items": {}, "prepaid": None,      # noqa: E731
                     "actual_paid": None, "invest": None, "stock": {},
                     "receivable": None, "total_buy": None}
    acc = {m: blank() for m in months_keys}
    # 누적 — 달을 더해서 만들지 않는다. 표의 합계/해외/국내 열을 그대로 담는다.
    sacc = blank()
    notes, extra = {}, []
    warnings, read = [], []
    label_cols = list(range(0, min(cols))) or [0]

    order, labels, keymap = [], {}, {}

    def item_key(label):
        """항목 이름 → 키. 모르는 이름도 그대로 받는다."""
        n = _norm(label)
        if n in keymap:
            return keymap[n]
        key = _known_key(label)
        if key is None or key in order:
            key = "g%d" % (len(order) + 1)
        keymap[n] = key
        order.append(key)
        labels[key] = _txt(label)
        return key

    def note_path(target):
        t, arg = target
        if t == "item":
            return "item:%s:%s" % arg
        if t == "stock":
            return "stock:%s" % arg
        if t == "sum":
            return "sum:%s" % arg
        return t

    def put(bucket, target, v, region=None):
        """한 칸 넣기. region=None 이면 그 달, 아니면 누적의 그 지역."""
        t, arg = target
        if t == "sum":
            return                      # 합계 줄은 우리가 더해서 확인한다
        if t == "item":
            ikey, mkey = arg
            box = bucket["items"].setdefault(ikey, {})
            if region is None:
                box[mkey] = v
            else:
                box.setdefault(mkey, {})[region] = v
            return
        if t == "stock":
            box = bucket["stock"].setdefault(arg, {})
            if region is None:
                box["value"] = v
            else:
                box.setdefault("value", {})[region] = v
            return
        if region is None:
            bucket[t] = v
        else:
            cur = bucket.get(t)
            cur = cur if isinstance(cur, dict) else {}
            cur[region] = v
            bucket[t] = cur

    cur_item = None
    skip_group = False

    for ri in range(hi + 1, len(rows)):
        row = rows[ri]
        parts = [t for t in
                 (_txt(row[j]) if j < len(row) else "" for j in label_cols) if t]
        vals = {m: (_num(row[j]) if j < len(row) else None) for j, m in cols.items()}
        regs = {r: (_num(row[j]) if j < len(row) else None)
                for r, j in regions.items()}
        note = (_note_txt(row[note_col])
                if (note_col is not None and note_col < len(row)) else "")
        # 비고 오른쪽에 적어 둔 메모(‘추가 분석 필요’ 목록 같은 것)도 챙긴다
        side = []
        if note_col is not None:
            for j in range(note_col + 1, len(row)):
                t = _note_txt(row[j])
                if t:
                    side.append(t)

        group = parts[-2] if len(parts) >= 2 else None
        metric = parts[-1] if parts else None
        has_val = (any(v is not None for v in vals.values())
                   or any(v is not None for v in regs.values()))

        def stash(*texts):
            for t in texts:
                for ln in [x for x in t.split("\n") if x.strip()] if t else []:
                    if ln not in extra and len(extra) < MAX_EXTRA:
                        extra.append(ln)

        stash(*side)
        if metric is None:
            stash(note)
            continue

        target = None
        sp = _special(metric)
        if sp:
            cur_item, skip_group = None, False
            target = sp if sp[0] == "stock" else (sp[0], None)
        else:
            if group is not None:
                skip_group = _is_skip(group)
                if skip_group:
                    cur_item = None
            if _is_skip(metric):
                stash(note)              # 재고 합계·지급 합계 줄 — 비고만
                continue
            if skip_group:
                mk = _metric(metric)
                if mk:
                    target = ("sum", mk)
                else:
                    stash(note)
                    continue
            else:
                if _match(metric, REVENUE):
                    target = ("revenue", None)
                if target is None:
                    for key, words in PAYMENT.items():
                        if _match(metric, words):
                            target = (key, None)
                            break
                if target is None:
                    mk = _metric(metric)
                    if mk:
                        if group is not None:
                            cur_item = item_key(group)
                        if cur_item:
                            target = ("item", (cur_item, mk))

        if target is None:
            if has_val:
                warnings.append("알 수 없는 줄: %s" % " ".join(parts)[:40])
            stash(note)
            continue
        if not has_val and not note:
            continue

        if target[0] == "ratio":
            # 비율은 저장하지 않는다 — 총매입액÷매출액으로 우리가 낸다.
            if note:
                notes.setdefault("ratio", note)
            continue

        for m, v in vals.items():
            if v is not None:
                put(acc[m], target, v)
        for r, v in regs.items():
            if v is not None:
                put(sacc, target, v, r)
        if note:
            notes.setdefault(note_path(target), note)
        read.append(" ".join(parts)[:40])

    def items_of(box):
        out = []
        for key in order:
            got = box.get(key)
            if not got:
                continue
            row = {"key": key, "label": labels.get(key, key)}
            for mk in ("buy", "paid", "balance"):
                row[mk] = got.get(mk)
            out.append(row)
        return out

    def stock_of(box):
        return [{"key": k, "value": v.get("value")}
                for k, v in box.items() if v.get("value") is not None]

    out = {}
    for m, mv in acc.items():
        items = items_of(mv["items"])
        stock = stock_of(mv["stock"])
        if not items and mv["revenue"] is None and not stock:
            continue
        one = {"revenue": mv["revenue"], "items": items,
               "prepaid": mv["prepaid"], "actual_paid": mv["actual_paid"],
               "invest": mv["invest"], "receivable": mv["receivable"],
               "total_buy": mv["total_buy"]}
        if stock:
            one["stock"] = stock
        out[m] = one

    if not out:
        raise ValueError("읽을 수 있는 숫자가 없습니다. 시트를 확인해 주세요.")

    summary = None
    s_items, s_stock = items_of(sacc["items"]), stock_of(sacc["stock"])
    if regions and (s_items or s_stock or any(
            sacc.get(k) for k in ("revenue", "prepaid", "actual_paid",
                                  "invest", "receivable", "total_buy"))):
        summary = {"revenue": sacc["revenue"], "items": s_items,
                   "prepaid": sacc["prepaid"], "actual_paid": sacc["actual_paid"],
                   "invest": sacc["invest"], "stock": s_stock,
                   "receivable": sacc["receivable"], "total_buy": sacc["total_buy"],
                   "from": months_keys[0], "to": months_keys[-1],
                   "notes": dict(notes)}
        if extra:
            summary["notes"].setdefault("month", "\n".join(extra))

    return {"sheet": ws.title, "sheets": wb.sheetnames, "year": year,
            "months": out, "summary": summary,
            "regions": [r for r in ("total", "overseas", "domestic") if r in regions],
            "read": read, "warnings": warnings,
            "items": [{"key": k, "label": labels.get(k, k)} for k in order]}


def merge(old: dict, new: dict) -> dict:
    """엑셀에 없는 달은 그대로 둔다. 있는 달은 숫자를 통째로 바꾼다.

    비고는 엑셀에 없고 화면에서만 적어 둔 것이 있어서, 새로 읽은 달에
    비고가 없으면 있던 비고를 지킨다.
    """
    out = dict(old or {})
    for m, mv in (new or {}).items():
        mv = dict(mv or {})
        prev = out.get(m)
        if not mv.get("notes") and isinstance(prev, dict) and prev.get("notes"):
            mv["notes"] = prev["notes"]
        out[m] = mv
    return out


def merge_summary(old, new):
    """누적도 숫자는 새 걸로, 비고는 합친다 (같은 줄이면 새 것)."""
    if not isinstance(new, dict):
        return old if isinstance(old, dict) else None
    out = dict(new)
    notes = dict((old or {}).get("notes") or {}) if isinstance(old, dict) else {}
    notes.update(out.get("notes") or {})
    if notes:
        out["notes"] = notes
    return out

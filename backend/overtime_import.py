"""'전사 사업부별 주간 잔특근 취합' 엑셀 읽기.

파일 한 개에 사업부 시트가 16장 있고, 각 장이 월~일 7줄이다. 요약
시트(전체_주간현황 · W39_주간현황 · W38_W39_비교)와 확인사항 시트가
따로 있다.

머리줄이 두 층이고 가로로 병합돼 있다.

    [총 보유인원      ][지원인력 ][실제 가용인원        ][잔업 인원  ]…
    [전체][주간][야간 ][ (+)][(-)][전체][주간][야간][야간비율][전체][직접][간접]

병합은 왼쪽 위 칸에만 값이 있어서, 그대로 읽으면 두 번째 칸부터
빈칸이 된다. 그래서 병합을 먼저 펴 놓고 (위 라벨, 아래 라벨) 짝으로
열을 찾는다. 좌표를 박지 않는 이유는 늘 같다 — 다음 주에 열이 하나
끼면 전부 밀린다.

'X' 는 0 이 아니다. 일요일 잔업, 평일 특근이 'X' 다. 0 으로 바꾸면
주 평균이 6/7 로 줄어든다. num() 이 None 으로 돌려주고, 저장도 None 이다.

비고와 일치확인은 사업부 시트가 아니라 [전체_주간현황] 에 있다.
확인사항 시트는 주차 전체에 붙는 메모로 담는다.
"""
from __future__ import annotations

import datetime as _dt
import io
import os
import re
import tempfile

from openpyxl import load_workbook

DAYS = ("mon", "tue", "wed", "thu", "fri", "sat", "sun")
_DAY_OF = {"월": "mon", "화": "tue", "수": "wed", "목": "thu",
           "금": "fri", "토": "sat", "일": "sun"}

# (위 머리줄, 아래 머리줄) → 어디에 담을지
GROUPS = [
    ("총보유인원", "headcount"),
    ("지원인력", "support"),
    ("실제가용인원", "available"),
    ("직접가용인원", "direct"),
    ("간접가용인원", "indirect"),
    ("잔업인원", ("overtime", "people")),
    ("특근인원", ("special", "people")),
]
def _sub_key(label) -> str:
    """아래 머리줄 → 어느 칸인지.

    라벨이 한국어·영어·베트남어가 줄바꿈으로 붙어 있고, 첫 줄만 보면
    '야간(가용인원)' 과 '야간비율(%)' 이 둘 다 '야간' 이 된다 — 비율이
    인원을 덮어썼다. 그래서 한글만 남겨 놓고 본다.
    """
    ko = re.sub(r"[^가-힣]", "", _txt(label))
    if not ko:
        return ""
    if "비율" in ko:
        return ""                    # 야간비율 등 — 우리가 인원으로 다시 낸다
    if ko.startswith("전체"):
        return "total"
    if ko.startswith("주간"):
        return "day"
    if ko.startswith("야간"):
        return "night"
    if ko.startswith("직접"):
        return "direct"
    if ko.startswith("간접"):
        return "indirect"
    return ""
SKIP_SHEETS_RE = re.compile(r"(주간현황|비교|확인사항|취합|안내|guide)", re.I)
_W_RE = re.compile(r"\bW(\d{1,2})\b", re.I)
_RANGE_RE = re.compile(r"(\d{1,2})\s*/\s*(\d{1,2})\s*~\s*(\d{1,2})\s*/\s*(\d{1,2})")


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.isoformat()
    return str(v).replace(" ", " ").strip()


def _norm(v) -> str:
    """머리줄 비교용 — 줄바꿈·괄호·공백·기호를 다 턴다."""
    t = _txt(v)
    t = re.split(r"[\n\r]", t)[0] if "\n" in t else t
    return re.sub(r"[\s()（）\[\]·.,/+\-%]+", "", t).lower()


def _norm_all(v) -> str:
    return re.sub(r"[\s()（）\[\]·.,/+\-%]+", "", _txt(v)).lower()


def num(v):
    """빈 칸과 'X' 는 None — 0 과 다르다."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = _txt(v).replace(",", "").replace("%", "")
    if not t or t in ("-", "—", "–") or t.upper() == "X":
        return None
    neg = False
    if t.startswith("(") and t.endswith(")"):
        neg, t = True, t[1:-1].strip()
    try:
        f = float(t)
    except ValueError:
        return None
    return -f if neg else f


def _open_book(data):
    if isinstance(data, (bytes, bytearray, memoryview)):
        raw = bytes(data)
    elif hasattr(data, "read"):
        raw = data.read()
        if not isinstance(raw, (bytes, bytearray)):
            raise ValueError("파일을 바이트로 읽지 못했습니다.")
        raw = bytes(raw)
    else:
        return load_workbook(data, data_only=True)
    try:
        return load_workbook(io.BytesIO(raw), data_only=True)
    except Exception:
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        try:
            with os.fdopen(fd, "wb") as f:
                f.write(raw)
            return load_workbook(path, data_only=True)
        finally:
            try:
                os.unlink(path)
            except OSError:
                pass


def _grid(ws, max_row=80, max_col=60):
    """병합을 펴서 읽는다 — 병합 칸은 왼쪽 위에만 값이 있다."""
    rows = min(ws.max_row or 1, max_row)
    cols = min(ws.max_column or 1, max_col)
    g = [[None] * cols for _ in range(rows)]
    for r in range(rows):
        for c in range(cols):
            g[r][c] = ws.cell(row=r + 1, column=c + 1).value
    for rng in ws.merged_cells.ranges:
        r0, c0 = rng.min_row - 1, rng.min_col - 1
        if r0 >= rows or c0 >= cols:
            continue
        v = g[r0][c0]
        for r in range(r0, min(rng.max_row, rows)):
            for c in range(c0, min(rng.max_col, cols)):
                g[r][c] = v
    return g


def _header(grid):
    """'요일' 이 있는 줄을 위 머리줄로 본다. 그 다음 줄이 아래 머리줄."""
    for i, row in enumerate(grid[:12]):
        for j, v in enumerate(row):
            if _norm(v) == "요일":
                return i, j
    return -1, -1


def _columns(grid, h1: int):
    """(위, 아래) 라벨을 보고 어느 열이 무엇인지 만든다."""
    top, sub = grid[h1], grid[h1 + 1] if h1 + 1 < len(grid) else []
    cols, extra = {}, {}
    for j in range(len(top)):
        t = _norm(top[j])
        if t in ("비고", "일치확인"):
            extra[t] = j
            continue
        for word, dest in GROUPS:
            if t != word:
                continue
            if dest == "support":
                # 지원인력은 아래 라벨이 '지원인력(+)' / '지원인력(-)' 이다
                raw = _txt(sub[j]) if j < len(sub) else ""
                if "+" in raw:
                    cols[j] = ("support", "in")
                elif "-" in raw or "−" in raw:
                    cols[j] = ("support", "out")
                break
            key = _sub_key(sub[j] if j < len(sub) else "")
            if not key:
                break                      # 야간비율 등 — 우리가 다시 계산한다
            cols[j] = (dest, key) if isinstance(dest, str) else dest + (key,)
            break
    return cols, extra


def _put(day: dict, path, value):
    cur = day
    for step in path[:-1]:
        cur = cur.setdefault(step, {})
    cur[path[-1]] = value


def _sheet_days(grid, h1: int, dayj: int, cols: dict):
    """월~일 줄을 읽는다. '합계/평균' 줄에서 멈춘다."""
    days, warnings = {}, []
    for ri in range(h1 + 2, len(grid)):
        raw = _txt(grid[ri][dayj]) if dayj < len(grid[ri]) else ""
        head = raw.split("\n")[0].strip()
        if not head:
            continue
        if any(w in _norm_all(raw) for w in ("합계", "평균", "total")):
            break
        d = _DAY_OF.get(head[:1])
        if d is None:
            warnings.append("모르는 요일 줄: %s" % head[:12])
            continue
        one = {}
        for j, path in cols.items():
            v = num(grid[ri][j]) if j < len(grid[ri]) else None
            _put(one, path if isinstance(path, tuple) else (path,), v)
        days[d] = one
    return days, warnings


def _week_of(text: str, year: int):
    m = _W_RE.search(text or "")
    if not m:
        return "", ""
    wk = "%04d-W%02d" % (year, int(m.group(1)))
    rg = _RANGE_RE.search(text or "")
    span = ("%s/%s~%s/%s" % rg.groups()) if rg else ""
    return wk, span


def _is_division_sheet(grid, h1: int, dayj: int) -> bool:
    for ri in range(h1 + 2, min(h1 + 12, len(grid))):
        head = _txt(grid[ri][dayj]).split("\n")[0].strip() if dayj < len(grid[ri]) else ""
        if head[:1] in _DAY_OF:
            return True
    return False


def _summary_notes(grid, h1: int, dayj: int, extra: dict):
    """[전체_주간현황] — 사업부별 비고 · 일치확인. 이름 → (비고, 확인)."""
    out = {}
    namej = 0
    for ri in range(h1 + 2, len(grid)):
        name = _txt(grid[ri][namej])
        if not name or any(w in _norm_all(name) for w in ("합계", "평균")):
            continue
        note = _txt(grid[ri][extra["비고"]]) if "비고" in extra else ""
        chk = _txt(grid[ri][extra["일치확인"]]) if "일치확인" in extra else ""
        chk = chk.split("\n")[0].strip()
        out[_norm_all(name)] = (note, chk)
    return out


def _issues(ws):
    """확인사항 시트 — 주차 전체에 붙는 메모."""
    out = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 40), max_col=4):
        vals = [_txt(c.value) for c in row]
        if not any(vals):
            continue
        if _norm_all(vals[0]) in ("사업부", ""):
            continue
        if "확인사항" in vals[0]:
            continue
        parts = [v for v in vals[1:3] if v]
        out.append(("%s — %s" % (vals[0], " / ".join(parts))).strip(" —/"))
    return out


def parse_workbook(data, year=None, week=None) -> dict:
    year = int(year or _dt.date.today().year)
    wb = _open_book(data)

    wk, span = "", ""
    divisions, order, warnings, read = {}, [], [], []
    notes, issues = {}, []

    for name in wb.sheetnames:
        ws = wb[name]
        grid = _grid(ws)
        if not grid:
            continue
        title = _txt(grid[0][0]) if grid[0] else ""
        w2, s2 = _week_of(title, year)
        if w2 and not wk:
            wk = w2
        if s2 and not span:
            span = s2
        h1, dayj = _header(grid)
        if h1 < 0:
            if _norm_all(name).startswith("확인사항"):
                issues = _issues(ws)
            continue
        cols, extra = _columns(grid, h1)
        if _is_division_sheet(grid, h1, dayj):
            if SKIP_SHEETS_RE.search(name):
                continue
            days, w = _sheet_days(grid, h1, dayj, cols)
            warnings += ["[%s] %s" % (name, x) for x in w]
            if not days:
                continue
            key = _norm_all(name)[:32] or ("d%d" % (len(order) + 1))
            divisions[key] = {"label": name, "days": days}
            order.append(key)
            read.append(name)
        elif "비고" in extra or "일치확인" in extra:
            notes.update(_summary_notes(grid, h1, dayj, extra))

    if not divisions:
        raise ValueError("사업부 시트를 찾지 못했습니다. 파일을 확인해 주세요.")
    if not wk:
        wk = str(week or "").strip()
        if not wk:
            raise ValueError("주차(W39)를 찾지 못했습니다. 주차를 직접 적어 주세요.")

    for key, body in divisions.items():
        note, chk = notes.get(key, ("", ""))
        if not note:
            note, chk = notes.get(_norm_all(body["label"]), (note, chk))
        body["note"] = note
        body["check"] = chk

    return {"week": wk, "range": span, "year": year,
            "divisions": divisions, "order": order,
            "issues": issues, "read": read, "warnings": warnings,
            "sheets": wb.sheetnames}

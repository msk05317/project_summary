"""모델 등록용 빈 엑셀 양식.

프로젝트마다 필요한 모델 정보는 결국 같다. 이 여덟 칸이면 등록이 끝난다.

    파트넘버 · 모델명 · 구분 · 유형 · 판가 · 재료비 · PO수량 · 실적수량

파트넘버를 첫 칸에 두는 이유가 있다. 임포터는 파트넘버로 모델을 찾는다.
큐리 파일은 '버스바' 5줄, '시트메탈' 6줄이 파트넘버로만 갈리는데,
예전처럼 모델명으로 맞추면 16줄이 7종으로 뭉개진다.

이 양식이 담지 않는 것 두 가지는 따로 올린다.
  - 주차별 계획/실적 → 주차 현황 보드의 '엑셀 올리기'
  - 개발 모델의 공정 일정 → 개발 공정 탭
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "0F2C59"
GREY = "F1F5F9"
LINE = "D6DCE5"
KFONT = "맑은 고딕"

_thin = Side(style="thin", color=LINE)
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

# (머리글, 너비, 숫자서식) — 순서가 곧 열 순서다
COLS = [
    ("파트넘버", 18, None),
    ("모델명", 22, None),
    ("구분", 9, None),
    ("유형", 14, None),
    ("판가($)", 13, '#,##0.00'),
    ("재료비($)", 13, '#,##0.00'),
    ("재료비율", 10, '0.0%'),
    ("PO수량", 11, '#,##0'),
    ("실적수량", 11, '#,##0'),
    ("비고", 34, None),
]
RATIO_COL = 7          # 재료비율 = 수식
ROWS = 60              # 빈 줄


def build_model_template(project_label="", types=None):
    """빈 모델 등록 양식 바이트를 돌려준다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "모델 목록"
    last = len(COLS)

    # ── 안내 줄
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    t = ws.cell(1, 1)
    t.value = ((f"{project_label} · " if project_label else "")
               + "모델 등록 양식    "
               + "[ 파트넘버가 모델을 구분하는 열쇠입니다. 이름이 같아도 파트넘버가 다르면 다른 모델로 들어갑니다 ]    "
               + "회색 칸은 수식이니 비워 두세요.")
    t.font = Font(name=KFONT, bold=True, size=11, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=NAVY)
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 26

    # ── 머리글
    for i, (name, width, _fmt) in enumerate(COLS, start=1):
        c = ws.cell(2, i, name)
        c.font = Font(name=KFONT, bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[2].height = 20

    # ── 빈 줄
    body = Font(name=KFONT, size=10)
    calc = PatternFill("solid", fgColor=GREY)
    for r in range(3, 3 + ROWS):
        for i, (_name, _w, fmt) in enumerate(COLS, start=1):
            c = ws.cell(r, i)
            c.font = body
            c.border = BORDER
            c.alignment = Alignment(
                horizontal="left" if i in (1, 2, 3, 4, last) else "right",
                vertical="center")
            if fmt:
                c.number_format = fmt
            if i == RATIO_COL:
                # 판가가 비어 있으면 아무것도 안 보이게 한다
                c.value = f'=IF(E{r}>0,F{r}/E{r},"")'
                c.fill = calc
        ws.row_dimensions[r].height = 18

    # ── 입력 편의
    dv_g = DataValidation(type="list", formula1='"양산,개발"', allow_blank=True)
    ws.add_data_validation(dv_g)
    dv_g.add(f"C3:C{2 + ROWS}")

    safe = ",".join(str(x).replace('"', "").replace(",", " ") for x in (types or []))[:250]
    if safe:
        dv_t = DataValidation(type="list", formula1=f'"{safe}"', allow_blank=True)
        ws.add_data_validation(dv_t)
        dv_t.add(f"D3:D{2 + ROWS}")

    ws.freeze_panes = ws.cell(3, 3)          # 파트넘버·모델명 고정
    ws.sheet_view.showGridLines = False

    # ── 설명 시트
    guide = wb.create_sheet("작성 안내")
    guide.column_dimensions["A"].width = 14
    guide.column_dimensions["B"].width = 92
    lines = [
        ("파트넘버", "모델을 구분하는 열쇠입니다. 이 값으로 기존 모델을 찾아 갱신합니다. "
                     "이름이 같아도(예: 버스바) 파트넘버가 다르면 각각 다른 모델로 등록됩니다."),
        ("모델명", "사람이 읽는 이름입니다. 같은 이름이 여러 줄 있어도 됩니다."),
        ("구분", "양산 또는 개발. 비워 두면 양산으로 들어갑니다. "
                 "개발로 넣으면 공정 12단계가 자동으로 붙습니다."),
        ("유형", "같은 프로젝트 안에서 묶어 볼 이름입니다 (예: 413, 009, EMA). "
                 "주차 현황 보드의 행을 이 값으로 묶습니다. 없으면 비워 두세요."),
        ("판가 / 재료비", "달러 기준입니다. 소수점 그대로 넣으셔도 됩니다. "
                          "판가는 매출 계산에 그대로 쓰입니다. 비어 있으면 그 모델 매출은 0 이 됩니다."),
        ("재료비율", "수식입니다. 손대지 마세요."),
        ("PO수량 / 실적수량", "지금까지 받은 주문 전체와 지금까지 나간 전체(누적)입니다. "
                              "그 달치가 아닙니다. 열을 통째로 비워 두면 기존 값을 그대로 둡니다."),
        ("비고", "메모입니다. 모델 목록의 비고 칸에 들어갑니다."),
        ("", ""),
        ("이 양식에 없는 것", "주차별 계획·실적은 주차 현황 보드의 '엑셀 올리기'로, "
                              "개발 모델의 공정 일정은 개발 공정 탭에서 따로 넣습니다."),
    ]
    r = 1
    for k, v in lines:
        a = guide.cell(r, 1, k)
        a.font = Font(name=KFONT, bold=True, size=10, color=NAVY)
        a.alignment = Alignment(vertical="top")
        b = guide.cell(r, 2, v)
        b.font = Font(name=KFONT, size=10)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        guide.row_dimensions[r].height = 30 if v else 10
        r += 1
    guide.sheet_view.showGridLines = False

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
